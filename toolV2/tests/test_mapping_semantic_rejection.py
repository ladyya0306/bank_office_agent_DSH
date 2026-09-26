"""Invalid semantic mappings are rejected before they can be silently retired."""
from __future__ import annotations

from docx import Document

from workflow.runner import dispatch


def test_later_typo_reports_actual_failed_id_without_changing_correct_positions(tmp_path):
    source = Document(); source.add_paragraph('借款人：合成企业'); source.save(tmp_path / 'source.docx')
    target = Document()
    for text in ('说明甲：____', '说明乙：____'): target.add_paragraph(text)
    target.save(tmp_path / 'target.docx')
    start = answer_all(dispatch({'action': 'start', 'work': str(tmp_path), 'source': ['source.docx'], 'targets': ['target.docx']}))
    slots = start['mapping_requests'][0]['positions']
    assert len(slots) == 2
    typo = slots[1]['id'][:3] + slots[1]['id'][4:]
    updates = [{'template': 'target.docx', 'slot_id': slots[0]['id'], 'leave_blank': True, 'reason': '来源缺少补充说明'},
               {'template': 'target.docx', 'slot_id': typo, 'leave_blank': True, 'reason': '来源缺少补充说明'}]
    query = {'action': 'update_positions', 'work': str(tmp_path), 'task_id': start['task_id'], 'updates': updates}
    rejected = dispatch(query)
    detail = rejected['rejected_updates'][0]
    assert detail['slot_id'] == typo and detail['slot_id'] != slots[0]['id']
    assert typo in detail['reason']
    assert detail['similar_positions'][0]['id'] == slots[1]['id']
    assert rejected['counters'] == start['counters']
    updates[1]['slot_id'] = slots[1]['id']
    assert dispatch(query)['status'] == 'completed'


def answer_all(result):
    while result['status'] in ('awaiting_source', 'awaiting_fill'):
        result = dispatch({'action': 'resume', 'work': result['work'], 'task_id': result['task_id'],
                           'answers': [{'id': q['id'], 'selected': [q['options'][0]['label']], 'custom': ''}
                                       for q in result['questions']]})
    return result


def test_product_slot_rejects_contract_but_nearby_contract_number_remains_fillable(tmp_path):
    source = Document()
    source.add_paragraph('\u501f\u6b3e\u4eba\uff1a\u5408\u6210\u4f01\u4e1a')
    source.add_paragraph('\u7efc\u5408\u6388\u4fe1\u5408\u540c\uff1aSYNTH-001')
    source.save(tmp_path / 'source.docx')
    target = Document()
    target.add_paragraph('\u5408\u540c\u7f16\u53f7\u4e3a\uff1a____\uff1b\u5408\u540c\u54c1\u79cd\uff1a____')
    target.save(tmp_path / 'target.docx')
    result = answer_all(dispatch({'action': 'start', 'work': str(tmp_path), 'source': ['source.docx'],
                                  'targets': ['target.docx'], 'batch': '20260926-96'}))
    product = next(item for item in result['mapping_requests'][0]['positions']
                   if '\u54c1\u79cd' in item['label'])
    counters, task_id = result['counters'].copy(), result['task_id']
    update = {'template': 'target.docx', 'slot_id': product['id'],
              'field': '\u501f\u6b3e\u4eba\u7efc\u5408\u6388\u4fe1\u5408\u540c'}
    rejected = dispatch({'action': 'update_positions', 'work': str(tmp_path), 'task_id': task_id,
                         'updates': [update]})
    detail = rejected['rejected_updates'][0]
    assert detail['slot_id'] == product['id'] and detail['label'] == product['label']
    assert '\u5408\u540c\u54c1\u79cd' in detail['reason']
    repeated = dispatch({'action': 'update_positions', 'work': str(tmp_path), 'task_id': task_id,
                         'updates': [update]})
    assert repeated['counters'] == counters
    assert repeated['rejected_updates'][0]['reason'] == detail['reason']

    valid = next(item for item in rejected['mapping_requests'][0]['positions']
                 if '\u7f16\u53f7' in item['label'])
    completed = dispatch({'action': 'update_positions', 'work': str(tmp_path), 'task_id': task_id,
                          'updates': [{'template': 'target.docx', 'slot_id': valid['id'],
                                       'field': '\u501f\u6b3e\u4eba\u7efc\u5408\u6388\u4fe1\u5408\u540c'},
                                      {'template': 'target.docx', 'slot_id': product['id'], 'leave_blank': True,
                                       'reason': '\u6765\u6e90\u672a\u63d0\u4f9b\u54c1\u79cd'}]})
    assert completed['status'] == 'completed', completed
    assert 'SYNTH-001' in Document(completed['results'][0]['output']).paragraphs[0].text
