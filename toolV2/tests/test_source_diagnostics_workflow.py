"""Source diagnostics survive process changes without erasing stable answers."""
from pathlib import Path
import json
import sqlite3
import subprocess
import sys
import tempfile
import unittest
from contextlib import closing
from docx import Document

ROOT = Path(__file__).resolve().parents[1]


def request(payload):
    run = subprocess.run([sys.executable, str(ROOT / 'office.py')],
                         input=json.dumps(payload), capture_output=True,
                         encoding='utf-8', timeout=90)
    return json.loads(run.stdout)


class DiagnosticWorkflowTests(unittest.TestCase):
    def test_unparsed_fact_stays_visible_after_answers_and_new_process(self):
        with tempfile.TemporaryDirectory() as tmp:
            work = Path(tmp)
            src, target = work / '来源.docx', work / '模板.docx'
            doc = Document()
            doc.add_paragraph('借款人：合成甲公司')
            doc.add_paragraph('联系电话：13800001234')
            quote = '借款金额超过300000元时，另以200000元作为保证金额上限。'
            doc.add_paragraph(quote)
            doc.save(src)
            doc = Document()
            doc.add_paragraph('联系电话：')
            doc.save(target)
            payload = dict(action='start', work=str(work), source=[str(src)], targets=[str(target)], batch='20260925-71')
            r = request(payload)
            for _ in range(4):
                if not r.get('questions'):
                    break
                self.assertTrue(any(x.get('quote') == quote for x in r['issues']), r)
                r = request(dict(action='resume', work=str(work), task_id=r['task_id'],
                                 answers=[dict(id=q['id'], selected=[q['options'][0]['label']], custom='')
                                          for q in r['questions']]))
            self.assertEqual('partial', r['status'], r)
            self.assertEqual('completed', r['results'][0]['status'])
            self.assertTrue(any(x.get('quote') == quote for x in r['issues']))
            before = r['counters']['fill_processes']
            r = request(payload)
            self.assertEqual('partial', r['status'])
            self.assertFalse(r['questions'])
            self.assertEqual(before, r['counters']['fill_processes'])
            self.assertTrue(any(x.get('quote') == quote for x in r['issues']))

    def test_parser_refresh_keeps_unchanged_answer(self):
        with tempfile.TemporaryDirectory() as tmp:
            work = Path(tmp)
            src, target = work / '来源.docx', work / '模板.docx'
            doc = Document()
            doc.add_paragraph('借款人：合成甲公司')
            doc.add_paragraph('联系电话：13800001234')
            doc.save(src)
            doc = Document()
            doc.add_paragraph('联系电话：')
            doc.save(target)
            payload = dict(action='start', work=str(work), source=[str(src)], targets=[str(target)], batch='20260925-72')
            r = request(payload)
            for _ in range(4):
                if not r.get('questions'):
                    break
                r = request(dict(action='resume', work=str(work), task_id=r['task_id'],
                                 answers=[dict(id=q['id'], selected=[q['options'][0]['label']], custom='')
                                          for q in r['questions']]))
            self.assertEqual('completed', r['status'], r)
            before = r['counters']['fill_processes']
            with closing(sqlite3.connect(work / 'db/workflow.db')) as conn:
                task = json.loads(conn.execute('SELECT payload FROM office_v2_task WHERE id=?', (r['task_id'],)).fetchone()[0])
                task['source_signature'] = 'older-parser-version'
                task.pop('source_issues', None)
                conn.execute('UPDATE office_v2_task SET payload=? WHERE id=?', (json.dumps(task), r['task_id']))
                conn.commit()
            r = request(payload)
            self.assertEqual('completed', r['status'], r)
            self.assertFalse(r['questions'])
            self.assertEqual(before, r['counters']['fill_processes'])


if __name__ == '__main__':
    unittest.main()
