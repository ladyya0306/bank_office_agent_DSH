"""First-run semantics with renamed templates and reordered source roles."""
from pathlib import Path
from docx import Document
from office_kit.store_v2 import StoreV2
from office_kit.fact_catalog import qualified_facts, target_subject_context
from office_kit.template_slots import discover_slots
from workflow.mapping import scoped_candidates
from workflow.mapping_view import page
from workflow.runner import dispatch


def test_precise_certificate_span_maps_on_first_run(tmp_path):
    source=Document()
    for line in ['借款人：合成甲企业','法定代表人：李合成','身份证号码：SYNTH-123']:
        source.add_paragraph(line)
    source.save(tmp_path/'information.docx')
    template=Document();template.add_paragraph('____同志（证件号：____），系我单位法定代表人。')
    template.save(tmp_path/'任意前缀-法定代表人身份证明（借款人）.docx')
    result=dispatch({'action':'start','work':str(tmp_path),'source':['information.docx'],
                     'targets':['任意前缀-法定代表人身份证明（借款人）.docx'],'batch':'20260926-65'})
    assert result['status']=='completed',result
    text=Document(result['results'][0]['output']).paragraphs[0].text
    assert '李合成同志' in text and '证件号：SYNTH-123' in text
    before=result['counters'].copy()
    again=dispatch({'action':'status','work':str(tmp_path),'task_id':result['task_id']})
    assert again['counters']==before


def test_spouse_never_gets_guarantor_compound_identity(tmp_path):
    with StoreV2(tmp_path/'data.db') as s:
        s.batch_scope='20260926-66'
        eid,_=s.ensure_entity('合成个人',personal=True)
        s.set_role(eid,'保证人',evidence='保证人1',batch_no=s.batch_scope)
        s.put_fact(s.batch_scope,'保证人名称','合成个人',entity_id=eid)
        s.put_fact(s.batch_scope,'证件号码','SYNTH-G',entity_id=eid)
        d=Document();d.add_paragraph('保证人1');d.add_paragraph('担保人配偶姓名及身份证号码：____');p=tmp_path/'任意材料.docx';d.save(p)
        slot=discover_slots(p)[0]
        scope=target_subject_context(s,p,slot['target'])
        assert scope['field_hint']=='配偶姓名及证件号码'
        assert all(c['score']<.85 for c in scoped_candidates(s,p,slot,qualified_facts(s)))


def test_source_page_returns_cached_quotes_and_bounded_long_rows():
    task={'rows':[{'_source':'a.docx','line':1,'quote':'借款人：合成企业','entity_name':'合成企业'},
                  {'_source':'a.docx','line':2,'quote':'长证据'*3000,'entity_name':'合成企业'}]}
    query={'section':'source'}; collected=[]
    while query:
        result=page(task,query);collected.extend(result['items']);query=result['next']
        assert len(str(result['items']))<6000
    assert collected[0]['original']=='借款人：合成企业'
    assert any('长证据' in str(item) for item in collected[1:])


def test_underlined_short_separators_do_not_become_questions(tmp_path):
    doc=Document()
    for left,right in [('向合成银行股份有限公司','合成分行申请授信'),
                       ('申请授信','人民币（币种）'),('合同编号为','2030合成字第')]:
        p=doc.add_paragraph(left);p.add_run(' ').underline=True;p.add_run(right)
    p=doc.add_paragraph('联系电话：');p.add_run(' ').underline=True
    path=tmp_path/'format-variant.docx';doc.save(path)
    slots=discover_slots(path)
    assert len(slots)==1 and slots[0]['label']=='联系电话'


def test_available_credit_uses_same_owner_and_exact_currency_units():
    from office_kit.fact_catalog import _available_credit
    fields={'借款人授信金额':{'value':'900万元','entity_id':1},
            '借款人已用额度':{'value':'400000元','entity_id':1}}
    _available_credit(fields)
    assert fields['借款人可用授信额度']['value']=='860万元'
    assert fields['借款人可用授信额度']['source_kind']=='computed'
    mismatch={'借款人授信金额':{'value':'900万元','entity_id':1},
              '借款人已用额度':{'value':'40万元','entity_id':2}}
    _available_credit(mismatch)
    assert '借款人可用授信额度' not in mismatch


def test_table_currency_label_requires_immediate_neighbor(tmp_path):
    doc=Document(); table=doc.add_table(rows=2, cols=3)
    table.cell(0,0).text='授信金额'
    table.cell(0,1).text=' 人民币    万元'
    table.cell(1,0).text='已用额度'
    table.cell(1,2).text=' 人民币    万元'
    path=tmp_path/'variant.docx';doc.save(path)
    slots=discover_slots(path)
    labelled=[s for s in slots if s['label']=='授信金额']
    assert len(labelled)==1,slots
    assert labelled[0]['target']['expected_text']==' 人民币    万元'
    assert not any(s['label']=='已用额度' and
                   s['target'].get('expected_text')==' 人民币    万元' for s in slots)


def test_event_place_does_not_inherit_company_address(tmp_path):
    from office_kit.harness import _field_for_target
    with StoreV2(tmp_path/'data.db') as store:
        store.batch_scope='20260926-68'
        eid,_=store.ensure_entity('合成甲企业')
        store.set_role(eid,'借款人',evidence='借款人',batch_no=store.batch_scope)
        store.put_fact(store.batch_scope,'借款人住所','合成路1号',entity_id=eid)
        doc=Document();doc.add_paragraph('核实地点：____')
        path=tmp_path/'签约核实书.docx';doc.save(path)
        slot=discover_slots(path)[0]
        scope=target_subject_context(store,path,slot['target'])
        assert scope['field_hint']=='核实地点'
        assert _field_for_target('借款人住所',scope,qualified_facts(store))=='借款人核实地点'


def test_numbered_guarantee_contract_follows_local_owner(tmp_path):
    from office_kit.harness import build_fill_plan
    with StoreV2(tmp_path/'data.db') as store:
        store.batch_scope='20260926-69'
        for number, name in [(2,'合成公司'),(1,'合成人员')]:
            eid,_=store.ensure_entity(name,personal=number==1)
            store.set_role(eid,'保证人',evidence=f'保证人{number}',batch_no=store.batch_scope)
            store.put_fact(store.batch_scope,'最高额保证合同',f'SYNTH-{number}',entity_id=eid)
        doc=Document();doc.add_paragraph('保证人1');doc.add_paragraph('《最高额保证合同》（编号为：____）')
        path=tmp_path/'任意核保.docx';doc.save(path)
        slot=discover_slots(path)[0]
        tid=store.register_template(path,store.batch_scope)
        store.add_rule(tid,'保证人2最高额保证合同','合同编号',slot['target'],confidence=1,batch_no=store.batch_scope)
        plan=build_fill_plan(store,[path],batch_no=store.batch_scope,run_id='synthetic-contract')
        assert plan['rows'][0]['value']=='SYNTH-1',plan


def test_contract_blank_does_not_take_representative_named_later(tmp_path):
    from office_kit.harness import _field_for_target
    with StoreV2(tmp_path/'contract.db') as store:
        eid,_=store.ensure_entity('合成企业')
        store.set_role(eid,'借款人',evidence='借款人')
        doc=Document();doc.add_paragraph('经核实，合成企业（合同编号为 2030合成贷字第 ____ 号）的借据上的法定代表人签字真实有效。')
        path=tmp_path/'签章核实书.docx';doc.save(path)
        slot=discover_slots(path)[0]
        scope=target_subject_context(store,path,slot['target'])
        assert scope['field_hint']=='合同编号',scope
        assert _field_for_target('借款人法定代表人',scope,{})=='借款人合同编号'
        assert _field_for_target('借款人流动资金贷款合同',scope,{})=='借款人流动资金贷款合同'


def test_contract_context_can_fill_a_new_complete_contract(tmp_path):
    source=Document()
    for line in ('借款人：合成企业','法定代表人：合成姓名','流动资金贷款合同：2030合成贷字第ABC号'):
        source.add_paragraph(line)
    source.save(tmp_path/'source.docx')
    doc=Document();doc.add_paragraph('经核实，合成企业（合同编号为 2030合成贷字第 ____ 号）的借据上的法定代表人签字真实有效。')
    doc.save(tmp_path/'签章核实书.docx')
    result=dispatch({'action':'start','work':str(tmp_path),'source':['source.docx'],
                     'targets':['签章核实书.docx'],'batch':'20260926-71'})
    if result['status']=='needs_mapping':
        slot=result['mapping_requests'][0]['positions'][0]
        result=dispatch({'action':'update_positions','work':str(tmp_path),'task_id':result['task_id'],
                         'updates':[{'template':'签章核实书.docx','slot_id':slot['id'],
                                     'field':'借款人流动资金贷款合同'}]})
    assert result['status']=='completed',result
    text=Document(result['results'][0]['output']).paragraphs[0].text
    assert '合成姓名' not in text
    assert text.count('2030合成贷字第')==1 and 'ABC' in text,text


def test_excel_value_wrapper_does_not_label_another_empty_cell(tmp_path):
    from openpyxl import Workbook
    from openpyxl.styles import Border, Side
    book=Workbook();sheet=book.active;sheet.title='不同表名'
    sheet['A1']='已使用额度：';sheet['B1']='人民币   万元'
    sheet.merge_cells('C1:E1')
    sheet['C1'].border=Border(bottom=Side(style='thin'))
    path=tmp_path/'pairs.xlsx';book.save(path);book.close()
    slots=discover_slots(path)
    assert [(s['target']['cell'],s['label']) for s in slots]==[('B1','已使用额度')],slots
    from workflow.mapping import _retire_obsolete_slots
    with StoreV2(tmp_path/'history.db') as store:
        tid=store.register_template(path,'20260926-70')
        old={'kind':'xlsx_cell','sheet':'不同表名','cell':'C1','label_cell':'B1','slot_id':'legacy-placeholder'}
        store.add_rule(tid,'合成金额','人民币',old,confidence=1,batch_no='20260926-70')
        record={'slots':slots,'slot_migration_version':6}
        _retire_obsolete_slots(store,path,record,'20260926-70')
        assert store.rules_for(path)==[]


def test_company_stamp_is_not_another_company_name_slot(tmp_path):
    doc=Document()
    for text in ('单位公章：____','借款人（公章）：____','法定代表人或授权代理人：____'):
        doc.add_paragraph(text)
    path=tmp_path/'正文与盖章.docx';doc.save(path)
    slots=discover_slots(path)
    assert len(slots)==3
    assert [s['protected'] for s in slots]==[True,True,False]


def test_business_product_blank_does_not_take_representative(tmp_path):
    from office_kit.harness import _field_for_target
    with StoreV2(tmp_path/'product.db') as store:
        eid,_=store.ensure_entity('合成公司');store.set_role(eid,'借款人',evidence='借款人')
        doc=Document();doc.add_paragraph('综合授信 □专项授信 单笔用信业务：____')
        path=tmp_path/'授信申请.docx';doc.save(path)
        slot=discover_slots(path)[0];scope=target_subject_context(store,path,slot['target'])
        assert scope['field_hint']=='用信业务品种'
        assert _field_for_target('借款人法定代表人',scope,{})=='借款人用信业务品种'


def test_printed_bank_addressee_is_not_an_account_bank_blank(tmp_path):
    doc=Document();doc.add_paragraph('合成银行股份有限公司合成分行：')
    doc.add_paragraph('开户银行：')
    path=tmp_path/'不同申请书.docx';doc.save(path)
    slots=discover_slots(path)
    assert len(slots)==1 and slots[0]['label']=='开户银行',slots


def test_guarantee_resolution_credit_amount_refers_to_borrower_credit(tmp_path):
    source = Document()
    for text in ('借款人：合成借款公司', '授信金额：900万元',
                 '保证人2：合成保证公司', '保证金额：350万元'):
        source.add_paragraph(text)
    source.save(tmp_path/'source.docx')
    target = Document()
    target.add_paragraph('为合成借款公司向合成银行申请的人民币（币种）____万元授信提供连带责任保证担保。')
    target.save(tmp_path/'保证人决议.docx')
    result = dispatch({'action':'start', 'work':str(tmp_path), 'source':['source.docx'],
                       'targets':['保证人决议.docx'], 'batch':'20260926-73'})
    assert result['status']=='completed', result
    assert result['questions']==[]
    text = Document(result['results'][0]['output']).paragraphs[0].text
    assert '900万元授信' in text and '350' not in text, text
