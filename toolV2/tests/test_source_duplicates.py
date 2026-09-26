"""Cross-process regressions for duplicate source rows and recovery.

Every input is generated in a temporary workspace.  The workflow is driven
through office.py in separate Python processes so SQLite/task snapshots and
the real Word writer are exercised together.
"""
from __future__ import annotations

import json
import sqlite3
import subprocess
import sys
import tempfile
import time
import unittest
from pathlib import Path

from docx import Document


TOOL = Path(__file__).resolve().parents[1]


def request(payload: dict) -> dict:
    completed = subprocess.run(
        [sys.executable, str(TOOL / "office.py")],
        input=json.dumps(payload, ensure_ascii=False),
        text=True,
        encoding="utf-8",
        capture_output=True,
        timeout=60,
    )
    if completed.returncode:
        raise AssertionError(f"office.py failed: {completed.stderr}\n{completed.stdout}")
    return json.loads(completed.stdout)


def write_docx(path: Path, *paragraphs: str) -> None:
    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    document.save(path)


def answer_all(result: dict, work: Path) -> dict:
    """Answer one public question page in a fresh process until completion."""
    for _ in range(6):
        if result["status"] not in ("awaiting_source", "awaiting_fill"):
            return result
        questions = result["questions"]
        if not questions:
            raise AssertionError(result)
        answers = []
        for question in questions:
            if question["id"].startswith("source-"):
                answers.append({"id": question["id"], "selected": [], "custom": "合成甲"})
            else:
                answers.append({"id": question["id"],
                                "selected": [question["options"][0]["label"]],
                                "custom": ""})
        result = request({"action": "resume", "work": str(work),
                          "task_id": result["task_id"], "answers": answers})
    raise AssertionError(f"workflow did not finish: {result}")


class SourceDuplicateTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.work = Path(self.tmp.name) / "synthetic-work"
        self.work.mkdir()
        self.source = self.work / "source.docx"
        self.target = self.work / "target.docx"

    def tearDown(self) -> None:
        for attempt in range(20):
            try:
                self.tmp.cleanup()
                return
            except PermissionError:
                if attempt == 19:
                    raise
                time.sleep(0.25)

    def payload(self, batch: str = "20260925-01") -> dict:
        return {"action": "start", "work": str(self.work),
                "source": [str(self.source)], "targets": [str(self.target)],
                "batch": batch}

    def test_duplicate_unowned_row_asks_once_then_cross_process_word_output_is_reused(self) -> None:
        write_docx(self.source, "收款账号：123456789", "收款账号：123456789")
        write_docx(self.target, "收款账号：")

        first = request(self.payload())
        self.assertEqual("awaiting_source", first["status"], first)
        self.assertEqual(1, len(first["questions"]), first["questions"])
        source_question = first["questions"][0]

        answered_source = request({"action": "resume", "work": str(self.work),
                                   "task_id": first["task_id"],
                                   "answers": [{"id": source_question["id"],
                                                "selected": [], "custom": "合成甲"}]})
        self.assertEqual("completed", answered_source["status"], answered_source)
        self.assertEqual([], answered_source["questions"])

        completed = answer_all(answered_source, self.work)
        self.assertEqual("completed", completed["status"], completed)
        output = Path(completed["results"][0]["output"])
        self.assertTrue(output.is_file())
        paragraphs = [p.text for p in Document(output).paragraphs]
        self.assertIn("收款账号：123456789", paragraphs)

        rerun = request(self.payload())
        self.assertEqual("completed", rerun["status"], rerun)
        self.assertEqual([], rerun["questions"])
        self.assertEqual(completed["counters"]["fill_processes"],
                         rerun["counters"]["fill_processes"])
        self.assertEqual(str(output), rerun["results"][0]["output"])

    def test_same_value_for_different_subjects_keeps_two_facts(self) -> None:
        write_docx(self.source, "借款人：合成甲", "联系电话：13800000000",
                   "保证人：合成乙", "联系电话：13800000000")
        write_docx(self.target, "联系电话：")

        started = request(self.payload())
        # Source sections are explicit, but the anonymous destination still
        # has two owners even though their phone values happen to be equal.
        self.assertEqual("awaiting_fill", started["status"], started)
        self.assertEqual(1, len(started["questions"]))
        self.assertEqual(3, len(started['questions'][0]['options']))
        conn = sqlite3.connect(self.work / "db" / "workflow.db")
        try:
            facts = conn.execute(
                "SELECT key, value, entity_id FROM fact "
                "WHERE batch_no=? AND source_kind='source' AND superseded_by IS NULL "
                "AND key='联系电话' ORDER BY entity_id", ("20260925-01",)
            ).fetchall()
        finally:
            conn.close()
        self.assertEqual(2, len(facts), facts)
        self.assertEqual({"13800000000"}, {row[1] for row in facts})
        self.assertEqual(2, len({row[2] for row in facts}))

    def test_changed_value_and_batch_do_not_reuse_saved_source_answer(self) -> None:
        write_docx(self.source, "收款账号：123456789")
        write_docx(self.target, "收款账号：")
        first = answer_all(request(self.payload()), self.work)
        self.assertEqual("completed", first["status"], first)

        write_docx(self.source, "收款账号：987654321")
        changed_value = request(self.payload())
        self.assertEqual("awaiting_source", changed_value["status"], changed_value)
        self.assertEqual(1, len(changed_value["questions"]))
        self.assertIn("987654321", changed_value["questions"][0]["question"])

        unchanged_source_new_batch = request(self.payload("20260925-02"))
        self.assertEqual("awaiting_source", unchanged_source_new_batch["status"],
                         unchanged_source_new_batch)
        self.assertEqual(1, len(unchanged_source_new_batch["questions"]))

    def test_existing_underdeduped_question_snapshot_is_normalized_and_resumable(self) -> None:
        write_docx(self.source, "收款账号：123456789", "收款账号：123456789")
        write_docx(self.target, "收款账号：")
        first = request(self.payload())
        self.assertEqual("awaiting_source", first["status"], first)
        self.assertEqual(1, len(first["questions"]))

        # Reproduce a task persisted by the old implementation, which exposed
        # the same semantic row twice with the same question ID.
        database = self.work / "db" / "workflow.db"
        conn = sqlite3.connect(database)
        try:
            payload = json.loads(conn.execute(
                "SELECT payload FROM office_v2_task WHERE id=?", (first["task_id"],)
            ).fetchone()[0])
            legacy = dict(payload["questions"][0])
            payload["questions"] = [payload["questions"][0], legacy]
            conn.execute("UPDATE office_v2_task SET payload=? WHERE id=?",
                         (json.dumps(payload, ensure_ascii=False), first["task_id"]))
            conn.commit()
        finally:
            conn.close()

        normalized = request({"action": "status", "work": str(self.work),
                              "task_id": first["task_id"]})
        self.assertEqual("awaiting_source", normalized["status"], normalized)
        self.assertEqual(1, len(normalized["questions"]))
        resumed = request({"action": "resume", "work": str(self.work),
                           "task_id": first["task_id"],
                           "answers": [{"id": normalized["questions"][0]["id"],
                                        "selected": [], "custom": "合成甲"}]})
        self.assertIn(resumed["status"], ("awaiting_fill", "completed"), resumed)


if __name__ == "__main__":
    unittest.main()
