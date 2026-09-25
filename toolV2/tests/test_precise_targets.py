from pathlib import Path
import zipfile

import pytest
from docx import Document
from openpyxl import Workbook

from office_kit.common import OfficeKitError
from office_kit.doc_fill import XmlEngine
from office_kit.target_validation import validate_target


def _template(tmp_path: Path) -> Path:
    doc = Document()
    doc.add_paragraph("借款人法定代表人：________")
    doc.add_paragraph("借款人法定代表人：________")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "地址"
    table.cell(0, 1).text = "________"
    path = tmp_path / "重复目标.docx"
    doc.save(path)
    return path


def test_precise_paragraph_and_span_only_hits_requested_duplicate(tmp_path):
    path = _template(tmp_path)
    engine = XmlEngine(path)
    target = {
        "kind": "anchor",
        "anchor": "借款人法定代表人：",
        "paragraph_index": 1,
        "expected_text": "借款人法定代表人：________",
        "span_start": len("借款人法定代表人："),
        "span_end": len("借款人法定代表人：________"),
    }
    hits = engine.find_anchor_hits(target)
    assert len(hits) == 1
    assert engine.fill_span(hits[0][0], hits[0][1], hits[0][2], "乙方代表")
    out = tmp_path / "out.docx"
    engine.save(out)
    doc = Document(out)
    assert doc.paragraphs[0].text == "借款人法定代表人：________"
    assert doc.paragraphs[1].text == "借款人法定代表人：乙方代表"


def test_precise_span_rejects_existing_text(tmp_path):
    path = _template(tmp_path)
    engine = XmlEngine(path)
    target = {
        "kind": "anchor", "anchor": "借款人法定代表人：",
        "paragraph_index": 0,
        "expected_text": "借款人法定代表人：________",
        "span_start": 0, "span_end": len("借款人法定代表人：") + 2,
    }
    with pytest.raises(OfficeKitError, match="已有文字"):
        engine.find_anchor_hits(target)


def test_target_validation_rejects_block_and_validates_cell(tmp_path):
    path = _template(tmp_path)
    with pytest.raises(OfficeKitError, match="不支持的 Word 目标 kind"):
        validate_target(path, {"kind": "block", "block_index": 1})
    result = validate_target(path, {"kind": "cell", "table": 0, "row": 0, "col": 1})
    assert result["kind"] == "cell"


def test_precise_target_needs_no_anchor(tmp_path):
    path = _template(tmp_path)
    target = {
        "kind": "anchor", "paragraph_index": 0,
        "expected_text": "借款人法定代表人：________",
        "span_start": len("借款人法定代表人："),
        "span_end": len("借款人法定代表人：________"),
    }
    assert validate_target(path, target)["kind"] == "anchor"


def test_validation_rejects_duplicate_hits_and_negative_cell(tmp_path):
    path = _template(tmp_path)
    with pytest.raises(OfficeKitError, match="未指定 occurrence"):
        validate_target(path, {"kind": "anchor", "anchor": "借款人法定代表人："})
    with pytest.raises(OfficeKitError, match="不能为负数"):
        validate_target(path, {"kind": "cell", "table": 0, "row": -1, "col": 1})


def test_cell_checks_all_paragraphs_and_xlsx_multi(tmp_path):
    path = _template(tmp_path)
    doc = Document(path)
    cell = doc.tables[0].cell(0, 1)
    cell.add_paragraph("已有说明")
    doc.save(path)
    with pytest.raises(OfficeKitError, match="已有文字"):
        validate_target(path, {"kind": "cell", "table": 0, "row": 0, "col": 1})

    book = Workbook()
    ws = book.active
    ws.title = "一"
    ws["A1"] = "姓名："
    ws2 = book.create_sheet("二")
    ws2["A1"] = "地址："
    xlsx = tmp_path / "multi.xlsx"
    book.save(xlsx)
    book.close()
    multi = {"kind": "multi", "targets": [
        {"kind": "xlsx_cell", "sheet": "一", "cell": "A1", "anchor": "姓名："},
        {"kind": "xlsx_cell", "sheet": "二", "cell": "A1", "anchor": "地址："},
    ]}
    assert validate_target(xlsx, multi)["kind"] == "multi"
    with pytest.raises(OfficeKitError, match="嵌套 multi"):
        validate_target(xlsx, {"kind": "multi", "targets": [multi]})


def test_xlsx_precise_spans_use_original_offsets_and_reject_overlap(tmp_path):
    book = Workbook()
    ws = book.active
    ws["A1"] = "姓名：____ 地址：________"
    xlsx = tmp_path / "precise.xlsx"
    book.save(xlsx)
    book.close()
    from office_kit.doc_fill import XlsxEngine
    original = "姓名：____ 地址：________"
    name_s, name_e = original.index("____"), original.index("____") + 4
    addr_s = original.index("________")
    engine = XlsxEngine(xlsx)
    base = {"kind": "xlsx_cell", "sheet": "Sheet", "cell": "A1",
            "expected_text": original}
    engine.fill_xlsx_cell({**base, "span_start": name_s, "span_end": name_e}, "很长的姓名")
    engine.fill_xlsx_cell({**base, "span_start": addr_s, "span_end": len(original)}, "短址")
    with pytest.raises(OfficeKitError, match="重叠"):
        engine.fill_xlsx_cell({**base, "span_start": name_s + 1, "span_end": name_e}, "重叠")
    out = tmp_path / "precise-out.xlsx"
    engine.save(out)
    from openpyxl import load_workbook
    wb = load_workbook(out)
    assert wb["Sheet"]["A1"].value == "姓名：很长的姓名 地址：短址"
    wb.close()


def test_xlsx_precise_validation_rejects_wrong_kind_and_formula(tmp_path):
    book = Workbook()
    ws = book.active
    ws["A1"] = "=1+2"
    xlsx = tmp_path / "formula.xlsx"
    book.save(xlsx)
    book.close()
    with pytest.raises(OfficeKitError, match="公式"):
        validate_target(xlsx, {"kind": "xlsx_cell", "cell": "A1",
                               "expected_text": "=1+2", "span_start": 0, "span_end": 4})
    with pytest.raises(OfficeKitError, match="Excel 目标 kind"):
        validate_target(xlsx, {"kind": "anchor", "cell": "A1"})


def test_precise_paragraph_index_survives_unpaired_field_code_paragraph(tmp_path):
    """A non-paired middle paragraph must not compress later XML part indexes."""
    path = tmp_path / "field-code.docx"
    doc = Document()
    doc.add_paragraph("前置段落")
    doc.add_paragraph("域代码可见占位")
    doc.add_paragraph("后续字段：________")
    doc.save(path)

    # Keep the paragraph visible to python-docx while making the raw scanner
    # see a different text-bearing element. This reproduces field/hyperlink
    # material found in real approval templates without changing the fixture
    # through the production engine.
    tmp = tmp_path / "repacked.docx"
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            payload = zin.read(item.filename)
            if item.filename == "word/document.xml":
                payload = payload.replace("<w:t>域代码可见占位</w:t>".encode(),
                                          "<w:instrText>域代码可见占位</w:instrText>".encode())
            zout.writestr(item, payload)
    tmp.replace(path)

    engine = XmlEngine(path)
    expected = "后续字段：________"
    target = {"kind": "anchor", "paragraph_index": 2,
              "expected_text": expected,
              "span_start": len("后续字段："), "span_end": len(expected)}
    hits = engine.find_anchor_hits(target)
    assert len(hits) == 1
    assert engine.fill_span(hits[0][0], hits[0][1], hits[0][2], "通过")
    out = tmp_path / "field-code-out.docx"
    engine.save(out)
    assert Document(out).paragraphs[2].text == "后续字段：通过"
