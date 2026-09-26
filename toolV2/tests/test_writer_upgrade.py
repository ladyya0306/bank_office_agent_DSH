"""A spreadsheet writer upgrade refreshes spreadsheets, retaining Word results."""
from unittest.mock import patch
from hashlib import sha256
from pathlib import Path
from docx import Document
from openpyxl import Workbook
from workflow import runner


def test_writer_upgrade_only_repeats_affected_file_type(tmp_path):
    source = Document(); source.add_paragraph('借款人：合成企业'); source.save(tmp_path / 'source.docx')
    word = Document(); word.add_paragraph('借款人名称：____'); word.save(tmp_path / 'word.docx')
    book = Workbook(); book.active['A1'] = '借款人名称：____'; book.save(tmp_path / 'sheet.xlsx')
    before = runner.dispatch({'action': 'start', 'work': str(tmp_path), 'source': ['source.docx'], 'targets': ['word.docx', 'sheet.xlsx']})
    assert before['status'] == 'completed'
    original = next(r['output'] for r in before['results'] if r['template'].endswith('.docx'))
    original_hash = sha256(Path(original).read_bytes()).hexdigest()
    query = {'action': 'status', 'work': str(tmp_path), 'task_id': before['task_id']}
    with patch.object(runner, 'XLSX_WRITER_VERSION', runner.XLSX_WRITER_VERSION + 1):
        after = runner.dispatch(query)
        assert after['status'] == 'completed' and after['questions'] == []
        assert after['counters']['fill_processes'] - before['counters']['fill_processes'] == 1
        assert after['counters']['previews'] - before['counters']['previews'] == 1
        assert after['counters']['source_reads'] == before['counters']['source_reads'] == 1
        assert runner.dispatch(query)['counters'] == after['counters']
    unchanged = next(r['output'] for r in after['results'] if r['template'].endswith('.docx'))
    assert unchanged == original and sha256(Path(unchanged).read_bytes()).hexdigest() == original_hash
