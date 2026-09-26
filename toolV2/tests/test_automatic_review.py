"""Certain source facts do not become questions or forged user answers."""
import sys
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from office_kit.fill_decisions import needs_answer, saved_choices, fingerprint, reuse_unchanged_choices
from office_kit.store_v2 import StoreV2
from office_kit.target_validation import template_read_cache, read_word_template
from docx import Document


class AutomaticReviewTests(unittest.TestCase):
    def test_only_uncertain_or_required_missing_values_need_answers(self):
        self.assertFalse(needs_answer({'kind': 'slot', 'decision': 'auto'}))
        self.assertTrue(needs_answer({'kind': 'slot', 'decision': 'ask'}))
        self.assertFalse(needs_answer({'kind': 'slot', 'decision': 'empty'}))
        self.assertTrue(needs_answer({'kind': 'slot', 'decision': 'empty', 'is_required': True}))

    def test_business_reference_does_not_borrow_other_subjects_identity_or_guarantee_amount(self):
        from office_kit.harness import _target_fact
        for field in ('证件编号', '保证金额', '保证期限', '收款账号'):
            facts = {field:{'id':1,'entity_id':9,'value':'source-value'}}
            self.assertIsNone(_target_fact(facts,field,[3])[1])
        facts = {'授信期限':{'id':2,'entity_id':9,'value':'2026-01-01至2027-01-01'}}
        self.assertEqual(facts['授信期限'],_target_fact(facts,'授信期限',[3])[1])

    def test_auto_has_no_new_confirmation_but_retains_prior_user_blank(self):
        with tempfile.TemporaryDirectory() as tmp:
            db = Path(tmp) / 'workflow.db'
            row = {'n': 1, 'kind': 'slot', 'decision': 'auto', 'template_sha256': 'synthetic', 'field': '电话'}
            plan = {'db_path': str(db), 'batch_no': '20260926-01', 'rows': [row]}
            with StoreV2(db) as store:
                pass
            choice, missing = saved_choices(plan)
            self.assertEqual([], missing)
            self.assertEqual('', choice['select'])
            with StoreV2(db) as store:
                store.conn.execute('INSERT INTO fill_decision(fingerprint,batch_no,template_sha256,field,action,decided_at,run_id) VALUES(?,?,?,?,?,?,?)',
                                   (fingerprint(plan,row),'20260926-01','synthetic','电话','blank','2026-09-26','synthetic-run'))
                store.conn.commit()
            choice, missing = saved_choices(plan)
            self.assertEqual(['1'], choice['blank'])
            self.assertEqual([], missing)

    def test_policy_change_preserves_prior_blank_but_changed_value_does_not(self):
        with tempfile.TemporaryDirectory() as tmp:
            db = Path(tmp) / 'workflow.db'
            row = {'n':1,'kind':'slot','decision':'ask','ask_reason':'旧类别确认',
                   'template_sha256':'synthetic','field':'金额','value':'100元'}
            before = {'db_path':str(db),'batch_no':'20260926-01','rows':[row]}
            after = {**before,'rows':[{**row,'decision':'auto','ask_reason':''}]}
            with StoreV2(db) as store:
                store.conn.execute('INSERT INTO fill_decision(fingerprint,batch_no,template_sha256,field,action,decided_at,run_id) VALUES(?,?,?,?,?,?,?)',
                                   (fingerprint(before,row),'20260926-01','synthetic','金额','blank','2026-09-26','prior-user-run'))
                store.conn.commit()
            self.assertEqual(1,reuse_unchanged_choices(before,after))
            self.assertEqual(['1'],saved_choices(after)[0]['blank'])
            changed = {**after,'rows':[{**after['rows'][0],'value':'200元'}]}
            self.assertEqual(0,reuse_unchanged_choices(before,changed))
            self.assertEqual([],saved_choices(changed)[0]['blank'])

    def test_composite_values_retain_owner_and_do_not_substitute_representative(self):
        from office_kit.fact_catalog import _composite_facts
        from office_kit.harness import _field_for_target
        facts = {'保证人2名称': {'value':'合成企业','entity_id':2,'_relation_owner_eid':2},
                 '保证人2法定代表人': {'value':'合成代表','entity_id':3,'_relation_owner_eid':2},
                 '保证人2法定代表人证件号码': {'value':'SYNTH-ID','entity_id':3,'_relation_owner_eid':2}}
        _composite_facts(facts)
        self.assertNotIn('保证人2名称及证件号码',facts)
        self.assertEqual('合成代表 SYNTH-ID',facts['保证人2法定代表人姓名及证件号码']['value'])
        self.assertEqual('保证人1法定代表人姓名及证件号码', _field_for_target('保证人2法定代表人',
                         {'role':'保证人','number':1,'field_hint':'法定代表人姓名及证件号码'}, facts))

    def test_source_contract_placeholder_is_missing_not_a_number(self):
        from office_kit.harness import _missing_source_value_reason
        self.assertIsNotNone(_missing_source_value_reason('贷款合同','2026合成字第    号'))
        self.assertIsNone(_missing_source_value_reason('贷款合同','2026合成字第123号'))
        self.assertIsNone(_missing_source_value_reason('地址','第    号'))

    def test_quoted_contract_title_resolves_number_without_another_question(self):
        from workflow.runner import dispatch
        with tempfile.TemporaryDirectory() as tmp:
            work=Path(tmp)
            d=Document(); d.add_paragraph('借款人：合成企业'); d.add_paragraph('综合授信合同：合成字第123号'); d.save(work/'source.docx')
            d=Document(); d.add_paragraph('借款人与贵行签署编号为【        】的《【综合授信合同】》'); d.save(work/'target.docx')
            result=dispatch({'action':'start','work':tmp,'source':['source.docx'],'targets':['target.docx'],'batch':'20260926-81'})
            self.assertEqual('completed',result['status'])
            self.assertIn('【合成字第123号】',Document(result['results'][0]['output']).paragraphs[0].text)

    def test_guarantee_recipient_is_borrower_even_in_guarantor_document(self):
        from office_kit.fact_catalog import target_subject_context
        with tempfile.TemporaryDirectory() as tmp:
            with StoreV2(Path(tmp)/'workflow.db') as store:
                store.batch_scope='20260926-01'
                borrower,_=store.ensure_entity('合成借款企业')
                guarantor,_=store.ensure_entity('合成保证企业')
                store.set_role(borrower,'借款人',batch_no=store.batch_scope)
                store.set_role(guarantor,'保证人',batch_no=store.batch_scope)
                for raw in ('关于本公司为    公司在某银行的授信提供保证担保。',
                            '同意为    公司向某银行申请授信提供保证担保。'):
                    start=raw.index('    ')
                    scope=target_subject_context(store,Path(tmp)/'保证人董事会.docx',
                        {'kind':'anchor','expected_text':raw,'span_start':start,'span_end':start+4})
                    self.assertEqual('借款人',scope['role'])
                    self.assertEqual([borrower],scope['entity_ids'])
                    self.assertEqual('名称',scope['field_hint'])

    def test_read_cache_reuses_parse_but_refreshes_changed_template(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'template.docx'
            doc = Document(); doc.add_paragraph('电话：'); doc.save(path)
            with template_read_cache():
                first = read_word_template(path)
                self.assertIs(first, read_word_template(path))
                doc.add_paragraph('地址：'); doc.save(path)
                self.assertIsNot(first, read_word_template(path))
            self.assertIsNot(first, read_word_template(path))

    def test_parser_upgrade_keeps_blanks_when_source_bytes_are_unchanged(self):
        from workflow import source
        from workflow.runner import dispatch
        with tempfile.TemporaryDirectory() as tmp:
            work = Path(tmp)
            doc = Document(); doc.add_paragraph('借款人：合成甲公司'); doc.save(work/'source.docx')
            doc = Document(); doc.add_paragraph('未知补充项目：____'); doc.save(work/'target.docx')
            result = dispatch({'action':'start','work':tmp,'source':['source.docx'],
                               'targets':['target.docx'],'batch':'20260926-78'})
            self.assertEqual('needs_mapping', result['status'])
            slot = result['mapping_requests'][0]['positions'][0]
            result = dispatch({'action':'update_positions','work':tmp,'task_id':result['task_id'],
                               'updates':[{'template':'target.docx','slot_id':slot['id'],
                                           'leave_blank':True,'reason':'来源未提供此项'}]})
            self.assertEqual('completed',result['status'])
            with patch.object(source, 'SOURCE_PARSE_VERSION', source.SOURCE_PARSE_VERSION + 1):
                upgraded=dispatch({'action':'status','work':tmp,'task_id':result['task_id']})
            self.assertEqual('completed',upgraded['status'])
            self.assertEqual(1,upgraded['results'][0]['coverage']['left_blank'])


if __name__ == '__main__':
    unittest.main()
