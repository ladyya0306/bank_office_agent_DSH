"""Regression tests for the public toolV2 JSON workflow contract.

All files are synthetic.  These tests deliberately use an actual Word source
and Excel template so a future workflow change cannot replace document input
with an in-memory mock.
"""
from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import time
import unittest
from pathlib import Path

from docx import Document
from openpyxl import Workbook, load_workbook


TOOL = Path(__file__).resolve().parents[1]


def request(payload: dict) -> dict:
    run = subprocess.run([sys.executable, str(TOOL / "office.py")], input=json.dumps(payload),
                         text=True, encoding="utf-8", capture_output=True, timeout=30)
    result = json.loads(run.stdout)
    return result


class WorkflowContractTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.work = Path(self.tmp.name) / "synthetic-work"
        self.work.mkdir()
        self.source = self.work / "source.docx"
        source = Document()
        source.add_paragraph("借款人：合成甲公司")
        source.add_paragraph("联系电话：13800000000")
        source.save(self.source)
        self.target = self.work / "target.xlsx"
        book = Workbook()
        sheet = book.active
        sheet["A1"] = "借款人名称："
        sheet["A2"] = "联系电话："
        book.save(self.target)

    def tearDown(self) -> None:
        # Windows can retain a just-exited child process's SQLite handle briefly.
        for attempt in range(5):
            try:
                self.tmp.cleanup()
                return
            except PermissionError:
                if attempt == 4:
                    raise
                time.sleep(0.2)

    def start(self, *, batch: str = "20260925-01") -> dict:
        return request({"action": "start", "work": str(self.work),
                        "source": [str(self.source)], "targets": [str(self.target)], "batch": batch})

    def accept_all(self, result: dict) -> dict:
        """Drive only the public question protocol; each answer accepts its shown value."""
        for _ in range(5):
            if result["status"] not in ("awaiting_source", "awaiting_fill"):
                return result
            self.assertTrue(result["questions"], result)
            result = request({"action": "resume", "work": str(self.work),
                              "task_id": result["task_id"],
                              "answers": [{"id": q["id"], "selected": [q["options"][0]["label"]],
                                           "custom": ""} for q in result["questions"]]})
        self.fail("流程在五次公开回答后仍未完成：%s" % result)

    @staticmethod
    def sheet_text(path: str) -> str:
        book = load_workbook(path, data_only=True)
        return "\n".join(str(cell.value or "") for row in book.active.iter_rows() for cell in row)

    def test_json_entry_creates_task_and_returns_source_question(self) -> None:
        result = self.start()
        self.assertTrue(result["ok"])
        self.assertEqual("awaiting_source", result["status"])
        self.assertTrue(result["task_id"])
        self.assertEqual("20260925-01", result["batch"])
        self.assertTrue(any(q["id"].startswith("source-") for q in result["questions"]))
        self.assertTrue((self.work / "db" / "workflow.db").is_file())

    def test_unchanged_request_reuses_public_task_state(self) -> None:
        first = self.start()
        second = self.start()
        self.assertEqual(first["task_id"], second["task_id"])
        self.assertEqual(first["status"], second["status"])
        self.assertEqual(first["questions"], second["questions"])

    def test_changed_source_reasks_only_changed_evidence_without_consuming_old_answer(self) -> None:
        original = self.start()
        old_question = original["questions"][0]
        answered = request({"action": "resume", "work": str(self.work), "task_id": original["task_id"],
                            "answers": [{"id": old_question["id"],
                                         "selected": [old_question["options"][0]["label"]], "custom": ""}]})
        self.assertNotEqual("awaiting_source", answered["status"])
        self.source.unlink()
        changed = Document()
        changed.add_paragraph("借款人：合成甲公司")
        changed.add_paragraph("联系电话：13900000000")
        changed.save(self.source)
        source_changed = self.start()
        self.assertEqual(original["task_id"], source_changed["task_id"])
        self.assertEqual("awaiting_source", source_changed["status"])
        self.assertIn("13900000000", source_changed["questions"][0]["question"])
        no_answer = request({"action": "resume", "work": str(self.work),
                             "task_id": original["task_id"], "answers": []})
        self.assertEqual("awaiting_source", no_answer["status"])

    def test_changed_batch_is_a_different_public_task(self) -> None:
        self.assertNotEqual(self.start()["task_id"], self.start(batch="20260925-02")["task_id"])

    def test_empty_resume_does_not_execute_fill(self) -> None:
        task = self.start()
        result = request({"action": "resume", "work": str(self.work),
                          "task_id": task["task_id"], "answers": []})
        self.assertEqual("awaiting_source", result["status"])
        self.assertFalse(list((self.work / "out").rglob("*_已填写.xlsx")) if (self.work / "out").exists() else [])

    def test_changed_source_fills_new_value_and_batch_does_not_cross_contaminate(self) -> None:
        first = self.accept_all(self.start(batch="20260925-01"))
        self.assertEqual("completed", first["status"])
        old_output = next(r["output"] for r in first["results"] if r["template"] == "target.xlsx")
        self.assertIn("13800000000", self.sheet_text(old_output))

        changed = Document()
        changed.add_paragraph("借款人：合成甲公司")
        changed.add_paragraph("联系电话：13900000000")
        changed.save(self.source)
        second = self.accept_all(self.start(batch="20260925-02"))
        self.assertEqual("completed", second["status"])
        new_output = next(r["output"] for r in second["results"] if r["template"] == "target.xlsx")
        self.assertIn("13900000000", self.sheet_text(new_output))
        self.assertIn("13800000000", self.sheet_text(old_output))
        self.assertNotEqual(old_output, new_output)

    def test_position_repair_only_fills_failed_template(self) -> None:
        bad = self.work / "needs-position.docx"
        doc = Document()
        doc.add_paragraph("未知标签：")
        doc.save(bad)
        started = request({"action": "start", "work": str(self.work), "source": [str(self.source)],
                           "targets": [str(self.target), str(bad)], "batch": "20260925-01"})
        partial = self.accept_all(started)
        self.assertEqual("needs_mapping", partial["status"])
        good = next(r for r in partial["results"] if r["template"] == "target.xlsx")
        self.assertEqual("pending", good["status"])
        before = partial["counters"]["fill_processes"]

        repaired = request({"action": "update_positions", "work": str(self.work),
                            "task_id": partial["task_id"], "updates": [{
                                "template": "needs-position.docx", "field": "联系电话",
                                "label": "未知标签", "target": {"kind": "anchor", "anchor": "未知标签：", "max_blank": 200},
                            }]})
        completed = self.accept_all(repaired)
        self.assertEqual("completed", completed["status"])
        good_after = next(r for r in completed["results"] if r["template"] == "target.xlsx")
        self.assertTrue(Path(good_after["output"]).is_file())
        self.assertEqual(before + 2, completed["counters"]["fill_processes"])
        repeated = request({"action": "status", "work": str(self.work), "task_id": completed['task_id']})
        self.assertEqual(completed['counters']['fill_processes'], repeated['counters']['fill_processes'])


if __name__ == "__main__":
    unittest.main()
