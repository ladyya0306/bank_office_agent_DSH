"""Synthetic regression checks for slot structure and local fit failures."""
from __future__ import annotations

import hashlib
from pathlib import Path
from tempfile import TemporaryDirectory

import pytest
from docx import Document
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Border, Side

from office_kit.common import OfficeKitError
from office_kit.doc_fill import XlsxEngine, XmlEngine
from office_kit.harness import _guard_fill, _refresh_xlsx_value_cells, extract_labels, parse_fill_selection
from office_kit.target_validation import target_structure_issue, value_target_issue
from office_kit.template_slots import discover_slots
from office_kit.value_fit import fit_value


def _styled_blank(cell) -> None:
    cell.border = Border(bottom=Side(style="thin"))


def test_xlsx_prefers_adjacent_value_cell_but_keeps_real_sentence_blank():
    with TemporaryDirectory() as tmp:
        template = Path(tmp) / "结构.xlsx"
        book = Workbook(); sheet = book.active
        sheet["A1"] = "借款人："; _styled_blank(sheet["B1"])
        sheet["A2"] = "期限____年"; _styled_blank(sheet["B2"])
        book.save(template); book.close()
        original_hash = hashlib.sha256(template.read_bytes()).hexdigest()

        slots = discover_slots(template)
        assert [s["target"]["cell"] for s in slots if s["label"] == "借款人"] == ["B1"]
        duration = next(s for s in slots if s["target"].get("cell") == "A2")
        assert duration["target"]["span_start"] < duration["target"]["span_end"]

        output = Path(tmp) / "out.xlsx"
        engine = XlsxEngine(template)
        engine.fill_xlsx_cell({"kind": "xlsx_cell", "sheet": sheet.title, "cell": "B1"}, "合成公司")
        engine.fill_xlsx_cell(duration["target"], "3年")
        engine.save(output)
        check = load_workbook(output)
        assert check.active["A1"].value == "借款人："
        assert check.active["B1"].value == "合成公司"
        assert check.active["A2"].value == "期限3年"
        check.close()
        assert hashlib.sha256(template.read_bytes()).hexdigest() == original_hash


def test_duration_rejects_date_range_but_allows_normal_date_and_duration():
    with TemporaryDirectory() as tmp:
        template = Path(tmp) / "期限.docx"
        doc = Document(); doc.add_paragraph("期限____年"); doc.add_paragraph("日期：____")
        doc.save(template)
        duration, date = discover_slots(template)
        assert value_target_issue(template, duration["target"], "2026年1月1日至2027年1月1日")
        with pytest.raises(OfficeKitError):
            fit_value("期限____年", 2, 6, "2026年1月1日至2027年1月1日")
        for interval in ('2026-01-01 - 2027-01-01', '2026年到2027年', '2026/1/1~2027/1/1'):
            with pytest.raises(OfficeKitError):
                fit_value("期限____年", 2, 6, interval)
        assert fit_value("期限____年", 2, 6, "3年") == "3"

        output = Path(tmp) / "日期输出.docx"
        engine = XmlEngine(template)
        hits = engine.find_anchor_hits(date["target"])
        assert len(hits) == 1
        par, start, end = hits[0]
        assert engine.fill_span(par, start, end, "2026年9月25日")
        engine.save(output)
        assert "日期：2026年9月25日" in "\n".join(p.text for p in Document(output).paragraphs)


def test_normal_year_month_day_date_slots_remain_fillable():
    with TemporaryDirectory() as tmp:
        template = Path(tmp) / "日期.docx"
        doc = Document(); doc.add_paragraph("日期：____年____月____日")
        doc.save(template)
        slots = discover_slots(template)
        assert len(slots) == 3
        output = Path(tmp) / "日期填写.docx"
        engine = XmlEngine(template)
        for slot, value in zip(slots, ("2026", "9", "25")):
            par, start, end = engine.find_anchor_hits(slot["target"])[0]
            assert engine.fill_span(par, start, end, value)
        engine.save(output)
        assert Document(output).paragraphs[0].text == "日期：2026年9月25日"


def test_old_xlsx_inline_rule_is_a_local_issue_and_other_target_can_proceed():
    with TemporaryDirectory() as tmp:
        template = Path(tmp) / "旧规则.xlsx"
        book = Workbook(); sheet = book.active
        sheet["A1"] = "借款人："; _styled_blank(sheet["B1"])
        book.save(template); book.close()
        old = {"kind": "xlsx_cell", "sheet": sheet.title, "cell": "A1",
               "expected_text": "借款人：", "span_start": 4, "span_end": 4}
        issue = target_structure_issue(template, old)
        assert issue and "B1" in issue
        plan = {"rows": [
            {"kind": "slot", "n": 1, "decision": "ask", "local_issue": issue,
             "ask_reason": issue, "value": "合成公司"},
            {"kind": "slot", "n": 2, "decision": "auto", "local_issue": None,
             "ask_reason": "", "value": "13800000000"},
        ]}
        accepted, problems = _guard_fill(plan, parse_fill_selection("", apply_all=True, blank=["1"]))
        assert not problems
        assert [(row["n"], row["action"]) for row in accepted] == [(1, "blank"), (2, "accept")]


def test_proposal_finds_right_value_cell_and_refreshes_only_stale_target():
    with TemporaryDirectory() as tmp:
        template = Path(tmp) / "迁移.xlsx"
        book = Workbook(); sheet = book.active
        sheet["A1"] = "借款人："; _styled_blank(sheet["B1"])
        book.save(template); book.close()
        discovered = extract_labels(template, known_labels={"借款人名称"})
        target = discovered[0]["target"]
        assert target["cell"] == "B1" and target["label_cell"] == "A1"
        old = {"kind": "multi", "targets": [
            {"kind": "xlsx_cell", "sheet": sheet.title, "cell": "A1", "anchor": "借款人："},
            {"kind": "xlsx_cell", "sheet": sheet.title, "cell": "D5", "anchor": "另一处："},
        ]}
        refreshed, count = _refresh_xlsx_value_cells(old, [target])
        assert count == 1
        assert {item["cell"] for item in refreshed} == {"B1", "D5"}


def test_existing_amount_fit_rules_remain_strict():
    assert fit_value("金额____万元", 2, 6, "800万元") == "800"
    with pytest.raises(OfficeKitError):
        fit_value("金额____元", 2, 6, "800美元")
