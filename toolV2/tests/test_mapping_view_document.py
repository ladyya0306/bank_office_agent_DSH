"""On-demand document pages expose only templates owned by the task."""
from __future__ import annotations

import pytest
from docx import Document
from openpyxl import Workbook

from workflow.mapping_view import page


def task(tmp_path):
    word = Document()
    word.add_paragraph('\u8bf4\u660e\u6bb5\u843d\uff1a\u6b64\u5904\u6ca1\u6709\u7a7a\u4f4d\u4e5f\u9700\u8981\u663e\u793a\u3002')
    word.add_paragraph('\u586b\u5199\u9879\uff1a____')
    word.add_table(rows=1, cols=1).cell(0, 0).text = '表格中的完整说明'
    word.save(tmp_path / 'full.docx')
    book = Workbook()
    book.active.title = 'Sheet1'
    book.active['A1'] = '\u7ecf\u529e\u673a\u6784\uff1a\u5408\u6210\u652f\u884c'
    book.active['B1'] = None
    other = book.create_sheet('Sheet2')
    other['C3'] = 7
    book.save(tmp_path / 'full.xlsx')
    return {'last_work': str(tmp_path), 'source_signature': 'synthetic', 'available_fields': [],
            'documents': {'full.docx': {'status': 'needs_mapping', 'slot_hash': 'w', 'slot_version': 1,
                                        'current_positions': [], 'blank_slots': {}},
                          'full.xlsx': {'status': 'needs_mapping', 'slot_hash': 'x', 'slot_version': 1,
                                        'current_positions': [], 'blank_slots': {}}}}


def test_document_page_reads_full_word_and_nonempty_excel_cells(tmp_path):
    data = task(tmp_path)
    word = page(data, {'section': 'document', 'template': 'full.docx'})
    assert any(item['part'] == 'word/document.xml' and item['paragraph_index'] == 0
               and '\u6ca1\u6709\u7a7a\u4f4d' in item['text'] for item in word['items']), word
    assert any('\u586b\u5199\u9879' in item['text'] for item in word['items'])
    assert any(item['text'] == '表格中的完整说明' for item in word['items'])
    excel = page(data, {'section': 'document', 'template': 'full.xlsx'})
    assert {(item['sheet'], item['cell'], item['text']) for item in excel['items']} == {
        ('Sheet1', 'A1', '\u7ecf\u529e\u673a\u6784\uff1a\u5408\u6210\u652f\u884c'), ('Sheet2', 'C3', '7')}


def test_document_page_rejects_non_task_template(tmp_path):
    data = task(tmp_path)
    with pytest.raises(ValueError, match='本任务'):
        page(data, {'section': 'document', 'template': 'outside.docx'})


def test_output_page_reads_only_current_completed_artifact_inside_work(tmp_path):
    data = task(tmp_path)
    output = Document(); output.add_paragraph('\u5df2\u586b\u5199\u4ea7\u7269\u4e0e\u6a21\u677f\u4e0d\u540c')
    (tmp_path / 'out').mkdir()
    output.save(tmp_path / 'out' / 'full-filled.docx')
    data['documents']['full.docx'].update(status='completed', output='out/full-filled.docx')
    result = page(data, {'section': 'output', 'template': 'full.docx'})
    assert result['template'] == 'full.docx' and result['output'].endswith('full-filled.docx')
    assert any('\u5df2\u586b\u5199\u4ea7\u7269' in item['text'] for item in result['items'])
    data['documents']['full.docx']['output'] = '../outside.docx'
    blocked = page(data, {'section': 'output', 'template': 'full.docx'})
    assert blocked['items'][0]['reason'] == '\u5f53\u524d\u4ea7\u7269\u4e0d\u5b58\u5728\u6216\u4e0d\u5728\u672c\u5de5\u4f5c\u533a\u5185'
    data['documents']['full.xlsx']['status'] = 'needs_mapping'
    absent = page(data, {'section': 'output', 'template': 'full.xlsx'})
    assert absent['items'][0]['reason'] == '\u8be5\u6a21\u677f\u5f53\u524d\u6ca1\u6709\u5df2\u5b8c\u6210\u4ea7\u7269'
