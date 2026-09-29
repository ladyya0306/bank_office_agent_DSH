"""Synthetic regressions for validating dynamic-script source evidence."""
from __future__ import annotations

from hashlib import sha256

import pytest
from docx import Document
from openpyxl import Workbook

from workflow.source_evidence import document_items, validate_updates


def digest(path) -> str:
    return sha256(path.read_bytes()).hexdigest()


def source_task(path) -> dict:
    return {"source": [path.name]}


def make_xlsx(path) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "来源"
    sheet["A1"] = "借款人：合成甲企业"
    sheet["B2"] = "联系电话：13800000000"
    book.save(path)


def make_docx(path) -> None:
    document = Document()
    document.add_paragraph("借款人：合成乙企业")
    document.add_paragraph("用途：流动资金")
    document.save(path)


def test_validates_xlsx_and_docx_evidence_and_returns_absorb_rows(tmp_path):
    book = tmp_path / "source.xlsx"
    word = tmp_path / "source.docx"
    make_xlsx(book)
    make_docx(word)
    task = {"source": [book.name, word.name]}
    updates = [
        {"source": book.name, "source_sha256": digest(book), "rows": [
            {"key": "联系电话", "value": "13800000000", "entity_name": "合成甲企业",
             "role": "借款人", "evidence": {"sheet": "来源", "cell": "B2"}},
        ]},
        {"source": word.name, "source_sha256": digest(word), "rows": [
            {"key": "用途", "value": "流动资金", "entity_name": "合成乙企业",
             "evidence": {"paragraph_index": 1}},
        ]},
    ]

    rows = validate_updates(tmp_path, task, updates)
    assert [(row["key"], row["value"]) for row in rows] == [("联系电话", "13800000000"), ("用途", "流动资金")]
    assert rows[0]["quote"] == "联系电话：13800000000"
    assert rows[0]["line"] == "来源!B2"
    assert rows[1]["line"] == "word/document.xml#1"
    assert rows[0]["evidence"] == {"sheet": "来源", "cell": "B2"}
    assert all(row["needs"] == "收" and row["assumed"] is False and row["confidence"] == 1.0
               for row in rows)
    assert {(item["sheet"], item["cell"]) for item in document_items(tmp_path, task, book.name)} == {
        ("来源", "A1"), ("来源", "B2")}


def test_rejects_claim_not_present_at_the_evidence_position(tmp_path):
    book = tmp_path / "source.xlsx"
    make_xlsx(book)
    with pytest.raises(ValueError, match="字面子串"):
        validate_updates(tmp_path, source_task(book), [{
            "source": book.name, "source_sha256": digest(book), "rows": [{
                # This text exists in A1, but the declared B2 location must
                # stand on its own; source-wide presence is not enough.
                "key": "任意字段", "value": "合成甲企业", "evidence": {"sheet": "来源", "cell": "B2"},
            }],
        }])


def test_rejects_changed_file_hash_before_returning_any_partial_rows(tmp_path):
    book = tmp_path / "source.xlsx"
    make_xlsx(book)
    old_hash = digest(book)
    book.write_bytes(book.read_bytes() + b"changed")
    with pytest.raises(ValueError, match="source_sha256"):
        validate_updates(tmp_path, source_task(book), [{
            "source": book.name, "source_sha256": old_hash, "rows": [{
                "key": "联系电话", "value": "13800000000", "evidence": {"sheet": "来源", "cell": "B2"},
            }],
        }])


def test_rejects_unknown_source_and_workspace_escape(tmp_path):
    book = tmp_path / "source.xlsx"
    make_xlsx(book)
    outside = tmp_path.parent / (tmp_path.name + "-outside.xlsx")
    make_xlsx(outside)
    task = source_task(book)
    for source in ("missing.xlsx", "../" + outside.name):
        with pytest.raises(ValueError, match="当前任务|工作区"):
            validate_updates(tmp_path, task, [{
                "source": source, "source_sha256": digest(book), "rows": [{
                    "key": "联系电话", "value": "13800000000", "evidence": {"sheet": "来源", "cell": "B2"},
                }],
            }])


def test_derived_value_requires_an_explicit_explanation_and_retains_raw_evidence(tmp_path):
    path = tmp_path / "derived.xlsx"
    book = Workbook()
    book.active["A1"] = "额度：12万元"
    book.save(path)
    proposal = {"source": path.name, "source_sha256": digest(path), "rows": [
        {"key": "金额（元）", "value": "120000", "evidence": {"sheet": "Sheet", "cell": "A1"}}]}
    with pytest.raises(ValueError, match="字面子串"):
        validate_updates(tmp_path, source_task(path), [proposal])
    proposal["rows"][0]["derivation"] = "12万元乘以10000，转换为元；须用户核对"
    result = validate_updates(tmp_path, source_task(path), [proposal])
    assert result[0]["quote"] == "额度：12万元"
    assert result[0]["derivation"] == proposal["rows"][0]["derivation"]
