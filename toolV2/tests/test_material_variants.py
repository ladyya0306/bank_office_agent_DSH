"""Synthetic material-variant checks for the toolV2 source reader.

The fixtures deliberately resemble common office material fragments, but contain
only invented names and identifiers.  The last test records the expected
non-silent handling of a narrative paragraph mixed with recognised fields.
"""
from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import time
import unittest
from pathlib import Path

from docx import Document

TOOL = Path(__file__).resolve().parents[1]
if str(TOOL) not in sys.path:
    sys.path.insert(0, str(TOOL))
from office_kit.absorb import absorb, source_questions  # noqa: E402


def write_docx(path: Path, *paragraphs: str) -> None:
    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    document.save(path)


def request(payload: dict) -> dict:
    completed = subprocess.run(
        [sys.executable, str(TOOL / "office.py")],
        input=json.dumps(payload, ensure_ascii=False), text=True, encoding="utf-8",
        capture_output=True, timeout=60,
    )
    if completed.returncode:
        raise AssertionError(f"office.py failed: {completed.stderr}\n{completed.stdout}")
    return json.loads(completed.stdout)


class MaterialVariantTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.work = Path(self.tmp.name) / "synthetic-material-work"
        self.work.mkdir()

    def tearDown(self) -> None:
        for attempt in range(20):
            try:
                self.tmp.cleanup()
                return
            except PermissionError:
                if attempt == 19:
                    raise
                time.sleep(0.25)

    def doc(self, name: str, *paragraphs: str) -> Path:
        path = self.work / name
        write_docx(path, *paragraphs)
        return path

    @staticmethod
    def row(rows: list[dict], key: str, value: str) -> dict:
        return next(item for item in rows if item["key"] == key and item["value"] == value)

    def start(self, sources: list[Path], target: Path, batch: str = "20260925-31") -> dict:
        return request({"action": "start", "work": str(self.work),
                        "source": [str(path) for path in sources], "targets": [str(target)],
                        "batch": batch})

    def test_arabic_and_chinese_numbered_guarantors_keep_separate_names(self) -> None:
        source = self.doc("numbered-guarantors.docx", "保证人1：合成保证公司甲",
                          "统一社会信用代码：CREDIT-A", "保证人（二）：合成自然人乙",
                          "身份证号码：ID-B")
        rows = absorb(source)
        self.assertEqual("合成保证公司甲", self.row(rows, "统一社会信用代码", "CREDIT-A")["entity_name"])
        self.assertEqual("合成自然人乙", self.row(rows, "证件号码", "ID-B")["entity_name"])

    def test_legal_representative_switches_when_company_switches(self) -> None:
        source = self.doc("legal-representatives.docx", "借款人：合成借款公司",
                          "法定代表人：借款方法人", "身份证号码：BORROWER-ID",
                          "保证人：合成保证公司", "法人代表：保证方法人", "护照号：GUARANTOR-PASS")
        rows = absorb(source)
        self.assertEqual("借款方法人", self.row(rows, "证件号码", "BORROWER-ID")["entity_name"])
        self.assertEqual("保证方法人", self.row(rows, "证件号码", "GUARANTOR-PASS")["entity_name"])

    def test_numbered_guarantor_sections_do_not_overwrite_current_owner(self) -> None:
        source = self.doc("two-guarantors.docx", "借款人：合成借款公司",
                          "地址：借款地址", "保证人1：合成个人保证人",
                          "联系电话：PERSON-PHONE", "保证人2：合成保证公司",
                          "地址：保证公司地址", "法人：保证公司法人",
                          "开户行：中国合成银行 账号：622200000001")
        rows = absorb(source)
        self.assertEqual("合成个人保证人", self.row(rows, "联系电话", "PERSON-PHONE")["entity_name"])
        self.assertEqual("合成保证公司", self.row(rows, "地址", "保证公司地址")["entity_name"])
        self.assertEqual("合成保证公司", self.row(rows, "保证人法定代表人", "保证公司法人")["entity_name"])
        self.assertEqual("合成保证公司", self.row(rows, "开户行", "中国合成银行")["entity_name"])
        self.assertEqual("合成保证公司", self.row(rows, "收款账号", "622200000001")["entity_name"])

    def test_guarantee_amount_and_term_on_one_line_are_separate_facts(self) -> None:
        source = self.doc("guarantee-amount-term.docx", "最高额保证合同：CONTRACT-2",
                          "保证人2：合成保证公司",
                          "保证金额：800万元  期限：2026-09-17至2027-09-17")
        rows = absorb(source)
        amount = self.row(rows, "保证金额", "800万元")
        term = self.row(rows, "保证期限", "2026-09-17至2027-09-17")
        self.assertEqual("合成保证公司", amount["entity_name"])
        self.assertEqual("合成保证公司", term["entity_name"])

    def test_guarantee_contract_title_binds_to_following_explicit_guarantor(self) -> None:
        source = self.doc("guarantee-title-owner.docx", "最高额保证合同：CONTRACT-2",
                          "保证人2：合成保证公司")
        row = self.row(absorb(source), "最高额保证合同", "CONTRACT-2")
        self.assertEqual("合成保证公司", row["entity_name"])
        self.assertFalse(row["assumed"])

    def test_unfollowed_guarantee_contract_title_stays_unassigned(self) -> None:
        source = self.doc("unfollowed-guarantee-title.docx", "最高额保证合同：CONTRACT-2")
        row = self.row(absorb(source), "最高额保证合同", "CONTRACT-2")
        self.assertIsNone(row["entity_name"])
        self.assertEqual("指明主体", row["needs"])

    def test_company_and_natural_person_document_values_do_not_cross_owners(self) -> None:
        source = self.doc("company-person.docx", "借款人：合成借款公司",
                          "统一社会信用代码：COMPANY-CREDIT", "保证人：合成个人保证人",
                          "居民身份证号：PERSON-ID")
        rows = absorb(source)
        self.assertEqual("合成借款公司", self.row(rows, "统一社会信用代码", "COMPANY-CREDIT")["entity_name"])
        self.assertEqual("合成个人保证人", self.row(rows, "证件号码", "PERSON-ID")["entity_name"])

    def test_explicit_certificate_label_creates_linked_type_and_number(self) -> None:
        source = self.doc("certificate-pair.docx", "保证人：合成台籍个人",
                          "台湾居民来往大陆通行证：TAIWAN-PASS-001")
        rows = absorb(source)
        number = self.row(rows, "证件号码", "TAIWAN-PASS-001")
        certificate_type = self.row(rows, "证件类型", "台湾居民来往大陆通行证")
        self.assertEqual("合成台籍个人", number["entity_name"])
        self.assertEqual(number["ownership_group"], certificate_type["ownership_group"])
        questions = source_questions(rows)
        self.assertEqual([], questions)

    def test_bank_and_account_in_one_field_are_preserved_with_owner(self) -> None:
        source = self.doc("bank-account.docx", "借款人：合成借款公司",
                          "开户行及账号：中国合成银行 622200000001")
        rows = absorb(source)
        combined = self.row(rows, "开户行及账号", "中国合成银行 622200000001")
        self.assertEqual("合成借款公司", combined["entity_name"])
        # 紧随材料明确的借款人，归属已经由源文直接证明，不再重复提问。
        self.assertFalse(combined["assumed"])
        self.assertEqual([], source_questions(rows))

    def test_one_line_bank_and_account_labels_are_split_or_explicitly_reported(self) -> None:
        source = self.doc("one-line-bank-account.docx", "借款人：合成借款公司",
                          "开户行：中国合成银行 账号：622200000001")
        rows = absorb(source)
        split = (any(row["key"] == "开户行" and row["value"] == "中国合成银行" for row in rows)
                 and any(row["key"] == "收款账号" and row["value"] == "622200000001" for row in rows))
        explicitly_unsplit = any("未拆解" in (question["header"] + question["question"])
                                  for question in source_questions(rows))
        self.assertTrue(split or explicitly_unsplit,
                        "同一行双标签既未拆成开户行/收款账号，也没有公开说明未拆解：%r" % rows)

    def test_same_field_in_two_files_with_different_values_requires_conflict_review(self) -> None:
        first = self.doc("source-a.docx", "借款人：合成借款公司", "联系电话：13800000000")
        second = self.doc("source-b.docx", "借款人：合成借款公司", "联系电话：13900000000")
        target = self.doc("target.docx", "联系电话：")
        started = self.start([first, second], target)
        self.assertEqual("awaiting_source", started["status"], started)
        conflict = next(question for question in started["questions"]
                        if question["id"].startswith("source-conflict-"))
        self.assertIn("source-a.docx", conflict["question"])
        self.assertIn("source-b.docx", conflict["question"])
        self.assertIn("13800000000", conflict["question"])
        self.assertIn("13900000000", conflict["question"])
        # 第 0 项是“暂不采用”；明确选第一份来源，不能依赖其默认位置。
        answered = request({"action": "resume", "work": str(self.work), "task_id": started["task_id"],
                            "answers": [{"id": conflict["id"],
                                         "selected": [conflict["options"][1]["label"]], "custom": ""}]})
        self.assertEqual("completed", answered["status"], answered)
        output = Document(answered["results"][0]["output"])
        self.assertIn("联系电话：13800000000", [p.text for p in output.paragraphs])

    def test_guarantee_contract_boundary_does_not_inherit_borrower(self) -> None:
        source = self.doc("contract-boundary.docx", "借款人：合成借款公司",
                          "联系电话：13800000000", "保证合同", "联系电话：13900000000")
        rows = absorb(source)
        guarantee_phone = self.row(rows, "联系电话", "13900000000")
        self.assertIsNone(guarantee_phone["entity_name"])
        self.assertEqual("指明主体", guarantee_phone["needs"])
        self.assertIn("保证合同", guarantee_phone["context"])

    def test_loan_contract_term_is_not_confused_with_credit_term(self) -> None:
        source = self.doc("loan-term.docx", "借款人：合成借款公司", "流动资金贷款合同", "期限：12个月")
        row = self.row(absorb(source), "贷款期限", "12个月")
        self.assertEqual("合成借款公司", row["entity_name"])
        self.assertIn("贷款合同", row["context"])

    def test_pure_narrative_material_reports_source_failure_not_empty_mapping(self) -> None:
        source = self.doc("narrative-only.docx", "本合同由合成甲方与合成乙方于合成地点签订。")
        target = self.doc("target.docx", "联系电话：")
        result = self.start([source], target)
        self.assertEqual("failed", result["status"], result)
        self.assertEqual("source", result["failed_stage"])
        self.assertEqual(1, result["mapping_pages"]["templates"]["total"])
        self.assertIsNone(result["mapping_pages"]["templates"]["items"][0]["positions"])
        self.assertIn("未识别出键值", result["issues"][0]["reason"])
        self.assertFalse(result["results"])

    def test_unknown_label_is_retained_and_presented_for_review(self) -> None:
        source = self.doc("unknown-label.docx", "借款人：合成借款公司", "风险提示：仅供合成测试")
        rows = absorb(source)
        unknown = self.row(rows, "风险提示", "仅供合成测试")
        self.assertFalse(unknown["known_key"])
        self.assertEqual("合成借款公司", unknown["entity_name"])
        # 未识别标签仍保留，但明确借款人段内的归属无需人工重复确认。
        self.assertFalse(unknown["assumed"])
        self.assertFalse(any("风险提示" in q["question"] for q in source_questions(rows)))

    def test_mixed_business_fact_is_not_completed_without_mapping_or_unparsed_notice(self) -> None:
        source = self.doc("mixed-business-fact.docx", "借款人：合成借款公司",
                          "本次借款金额为人民币300000元。")
        target = self.doc("target.docx", "借款金额：")
        rows = absorb(source)
        started = self.start([source], target)
        amount_candidate = any("300000" in str(row.get("value", "")) for row in rows)
        public_messages = list(started.get("issues", [])) + source_questions(rows)
        explicitly_unparsed = any("300000" in str(message)
                                  and ("未解析" in str(message) or "未识别" in str(message))
                                  for message in public_messages)
        self.assertTrue(amount_candidate or explicitly_unparsed,
                        "300000 未形成候选，也没有包含该数值的公开未解析说明：%r" % started)
        self.assertNotEqual("completed", started["status"], started)


if __name__ == "__main__":
    unittest.main()
