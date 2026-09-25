"""Synthetic multi-owner answers survive process changes and workspace moves."""
from pathlib import Path
from contextlib import closing
import json
import shutil
import sqlite3
import tempfile
import unittest
from docx import Document
from test_source_roles import request, write_docx


class FillReviewRecoveryTests(unittest.TestCase):
    def prepare(self, work, names=('未指定主体的表单.docx',)):
        src = work / '来源.docx'
        write_docx(src, '借款人：合成借款公司', '联系电话：13800000001',
                   '', '保证人1：合成保证公司', '联系电话：13800000002')
        targets = []
        for name in names:
            target = work / name
            write_docx(target, '联系电话：')
            targets.append(str(target))
        payload = dict(action='start', work=str(work), source=[str(src)],
                       targets=targets, batch='20260925-91')
        r = request(payload)
        for _ in range(3):
            if r['status'] != 'awaiting_source':
                break
            r = request(dict(action='resume', work=str(work), task_id=r['task_id'],
                             answers=[dict(id=q['id'], selected=[q['options'][0]['label']], custom='')
                                      for q in r['questions']]))
        self.assertEqual('awaiting_fill', r['status'], r)
        return payload, r

    def test_unknown_owners_are_not_merged_across_forms(self):
        with tempfile.TemporaryDirectory() as tmp:
            _, r = self.prepare(Path(tmp), ('甲表.docx', '乙表.docx'))
            self.assertEqual(2, len(r['questions']), r)
            for q in r['questions']:
                self.assertIn('并非资料缺失', q['question'])
                self.assertFalse('甲表.docx' in q['question'] and '乙表.docx' in q['question'])

    def test_custom_value_fills_word_and_reuses_after_move_without_changing_source_facts(self):
        with tempfile.TemporaryDirectory() as tmp:
            work = Path(tmp) / '随机工作目录'
            work.mkdir()
            _, r = self.prepare(work)
            q = r['questions'][0]
            self.assertIn('多个主体', q['question'])
            r = request(dict(action='resume', work=str(work), task_id=r['task_id'],
                             answers=[dict(id=q['id'], selected=[], custom='13800000999')]))
            self.assertEqual('completed', r['status'], r)
            text = '\n'.join(p.text for p in Document(r['results'][0]['output']).paragraphs)
            self.assertIn('13800000999', text)
            count = r['counters']['fill_processes']
            with closing(sqlite3.connect(work / 'db/workflow.db')) as db:
                values = {x[0] for x in db.execute("SELECT value FROM fact WHERE key='联系电话' AND superseded_by IS NULL")}
                self.assertEqual({'13800000001', '13800000002'}, values)
                self.assertEqual(1, db.execute('SELECT count(*) FROM fill_decision').fetchone()[0])
            moved = Path(tmp) / '另一台电脑的目录'
            shutil.copytree(work, moved)
            r = request(dict(action='status', work=str(moved), task_id=r['task_id']))
            self.assertEqual('completed', r['status'], r)
            self.assertFalse(r['questions'])
            self.assertEqual(count, r['counters']['fill_processes'])
            self.assertTrue(Path(r['results'][0]['output']).is_relative_to(moved))

    def test_unrelated_new_batch_does_not_restart_an_existing_file_set(self):
        with tempfile.TemporaryDirectory() as tmp:
            work = Path(tmp)
            payload, r = self.prepare(work)
            r = request(dict(action='resume', work=str(work), task_id=r['task_id'],
                             answers=[dict(id=q['id'], selected=['留空'], custom='') for q in r['questions']]))
            self.assertEqual('completed', r['status'], r)
            old_id, counts = r['task_id'], r['counters']
            other = work / '临时诊断表.docx'
            write_docx(other, '联系电话：')
            request({**payload, 'targets': [str(other)], 'batch': '20260925-93'})
            payload.pop('batch')
            resumed = request(payload)
            self.assertEqual(old_id, resumed['task_id'])
            self.assertEqual('completed', resumed['status'])
            self.assertFalse(resumed['questions'])
            self.assertEqual(counts, resumed['counters'])

    def test_bad_answer_returns_original_task_for_recovery(self):
        with tempfile.TemporaryDirectory() as tmp:
            work = Path(tmp)
            _, r = self.prepare(work)
            original = r['task_id']
            bad = request(dict(action='resume', work=str(work), task_id=original,
                               answers=[dict(id=r['questions'][0]['id'], selected=['不存在的选项'], custom='')]))
            self.assertEqual('failed', bad['status'], bad)
            self.assertEqual(original, bad['task_id'])
            self.assertEqual('20260925-91', bad['batch'])
            self.assertIn('不要更换批次', bad['recovery'])
            r = request(dict(action='status', work=str(work), task_id=original))
            self.assertEqual('awaiting_fill', r['status'])
            r = request(dict(action='resume', work=str(work), task_id=original,
                             answers=[dict(id=q['id'], selected=['留空'], custom='') for q in r['questions']]))
            self.assertEqual('completed', r['status'], r)

    def test_two_certificates_take_their_own_addresses_and_do_not_borrow_phone(self):
        with tempfile.TemporaryDirectory() as tmp:
            work = Path(tmp)
            src = work / '来源.docx'
            write_docx(src, '借款人：合成借款公司', '地址：借款公司所在路甲号',
                       '法定代表人：合成甲代表', '',
                       '保证人1：合成乙个人', '身份证号：PERSON-SYNTHETIC',
                       '联系电话：13800000011', '',
                       '保证人2：合成保证公司', '地址：保证公司所在路乙号',
                       '法定代表人：合成丙代表', '联系电话：13800000022')
            names = ['法定代表人身份证明书（借款人）.docx', '法定代表人身份证明书(保证人).docx']
            for name in names:
                write_docx(work / name, '办公地点：', '联系电话：')
            payload = dict(action='start', work=str(work), source=[str(src)],
                           targets=[str(work / n) for n in names], batch='20260925-92')
            r = request(payload)
            for _ in range(4):
                if r['status'] not in ('awaiting_source', 'awaiting_fill'):
                    break
                if r['status'] == 'awaiting_fill':
                    for q in r['questions']:
                        self.assertFalse(all(name in q['question'] for name in names), q)
                        self.assertNotIn('13800000011', q['question'], q)
                r = request(dict(action='resume', work=str(work), task_id=r['task_id'],
                                 answers=[dict(id=q['id'], selected=[q['options'][0]['label']], custom='')
                                          for q in r['questions']]))
            self.assertEqual('completed', r['status'], r)
            outputs = {row['template']: '\n'.join(p.text for p in Document(row['output']).paragraphs)
                       for row in r['results']}
            self.assertIn('借款公司所在路甲号', outputs[names[0]])
            self.assertNotIn('保证公司所在路乙号', outputs[names[0]])
            self.assertNotIn('138000000', outputs[names[0]])
            self.assertIn('保证公司所在路乙号', outputs[names[1]])
            self.assertIn('13800000022', outputs[names[1]])
            before = r['counters']
            again = request(payload)
            self.assertEqual('completed', again['status'], again)
            self.assertEqual(before, again['counters'])
            self.assertFalse(again['questions'])


if __name__ == '__main__':
    unittest.main()
