"""定向验证来源材料的多标签拆解与叙述诊断。"""
from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from openpyxl import Workbook

TOOL = Path(__file__).resolve().parents[1]
if str(TOOL) not in sys.path:
    sys.path.insert(0, str(TOOL))

from office_kit.absorb import absorb, absorb_with_diagnostics  # noqa: E402


def write_docx(path: Path, *paragraphs: str) -> None:
    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    document.save(path)


class SourceRepairTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()

    def tearDown(self) -> None:
        self.tmp.cleanup()

    def source(self, *paragraphs: str) -> Path:
        path = Path(self.tmp.name) / "source.docx"
        write_docx(path, *paragraphs)
        return path

    def xlsx_source(self) -> Path:
        path = Path(self.tmp.name) / "horizontal-form.xlsx"
        book = Workbook()
        sheet = book.active
        sheet["A1"] = "借款人"
        sheet["B1"] = "合成借款企业"
        sheet.merge_cells("B1:C1")
        sheet["D1"] = "联系电话"
        sheet["E1"] = "13800000000"
        sheet["A2"] = "开户行"
        sheet.merge_cells("A2:B2")
        sheet["C2"] = "合成银行"
        sheet.merge_cells("C2:D2")
        sheet["E2"] = "收款账号"
        sheet["F2"] = "622200000001"
        sheet.merge_cells("F2:H2")
        book.save(path)
        return path

    def test_one_line_known_labels_split_and_retain_full_quote_and_line(self) -> None:
        text = "开户行：中国合成银行 账号：622200000001"
        rows = absorb(self.source("借款人：合成借款公司", text))
        bank = next(row for row in rows if row["key"] == "开户行")
        account = next(row for row in rows if row["key"] == "收款账号")
        self.assertEqual("中国合成银行", bank["value"])
        self.assertEqual("622200000001", account["value"])
        self.assertEqual(text, bank["quote"])
        self.assertEqual(text, account["quote"])
        self.assertEqual(bank["line"], account["line"])
        self.assertEqual("合成借款公司", account["entity_name"])

    def test_colon_in_bank_value_and_ordinary_explanation_are_not_split_as_fields(self) -> None:
        rows = absorb(self.source("开户行：中国：合成银行 账号：622200000001",
                                  "说明：本材料仅用于合成测试：不得据此付款。"))
        bank = next(row for row in rows if row["key"] == "开户行")
        self.assertEqual("中国：合成银行", bank["value"])
        self.assertFalse(any(row["key"] == "本材料仅用于合成测试" for row in rows))

    def test_complex_business_narrative_is_reported_but_plain_narrative_is_ignored(self) -> None:
        amount = "本次借款金额为人民币300000元，另含保证金10000元。"
        rows, issues = absorb_with_diagnostics(self.source(amount, "本合同说明仅供合成测试。"))
        self.assertEqual([], rows)
        self.assertEqual(1, len(issues))
        self.assertEqual(0, issues[0]["line"])
        self.assertEqual(amount, issues[0]["quote"])
        self.assertIn("未按“标签：值”解析", issues[0]["reason"])

    def test_strict_loan_amount_sentence_becomes_confirmable_candidate(self) -> None:
        text = "本次借款金额为人民币300000元。"
        rows, issues = absorb_with_diagnostics(self.source("借款人：合成借款公司", text))
        amount = next(row for row in rows if row["key"] == "借款金额")
        self.assertEqual("300000元", amount["value"])
        self.assertEqual(text, amount["quote"])
        self.assertTrue(amount["assumed"])
        self.assertEqual([], issues)

    def test_long_explicit_account_is_reported_without_truncation(self) -> None:
        text = "收款账号：" + "6" * 81
        rows, issues = absorb_with_diagnostics(self.source(text))
        self.assertEqual([], rows)
        self.assertEqual(text, issues[0]["quote"])
        self.assertIn("超过 80 个字符", issues[0]["reason"])

    def test_xlsx_horizontal_form_reads_merged_label_value_pairs(self) -> None:
        rows = absorb(self.xlsx_source())
        self.assertEqual("合成借款企业", next(row["value"] for row in rows
                                             if row["key"] == "借款人名称"))
        self.assertEqual("13800000000", next(row["value"] for row in rows
                                               if row["key"] == "联系电话"))
        self.assertEqual("合成银行", next(row["value"] for row in rows
                                          if row["key"] == "开户行"))
        self.assertEqual("622200000001", next(row["value"] for row in rows
                                               if row["key"] == "收款账号"))

    def test_xlsx_inline_fields_are_not_joined_with_the_next_cell(self) -> None:
        path = Path(self.tmp.name) / "inline-fields.xlsx"
        book = Workbook()
        sheet = book.active
        sheet["A1"] = "借款人：合成借款企业"
        sheet["B1"] = "联系电话：13800000000"
        book.save(path)
        rows = absorb(path)
        self.assertEqual("合成借款企业", next(row["value"] for row in rows
                                             if row["key"] == "借款人名称"))
        self.assertEqual("13800000000", next(row["value"] for row in rows
                                               if row["key"] == "联系电话"))

    def test_xlsx_header_cells_are_not_unconditionally_paired_as_facts(self) -> None:
        path = Path(self.tmp.name) / "headers.xlsx"
        book = Workbook()
        sheet = book.active
        sheet["A1"] = "客户名称"
        sheet["B1"] = "贷款金额"
        book.save(path)
        self.assertEqual([], absorb(path))

    def test_xlsx_unmerged_horizontal_pairs_are_retained_without_header_guessing(self) -> None:
        path = Path(self.tmp.name) / "unmerged-horizontal-form.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.append(["客户名称", "合成（上海）科技有限公司", "额度", "500万元", "联系电话", "13800000000"])
        book.save(path)
        rows = absorb(path)
        self.assertEqual({("借款人名称", "合成（上海）科技有限公司"), ("额度", "500万元"),
                          ("联系电话", "13800000000")},
                         {(row["key"], row["value"]) for row in rows})

    def test_xlsx_table_headers_are_not_values_for_neighbouring_headers(self) -> None:
        path = Path(self.tmp.name) / "table-headers.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.append(["协议号", "业务类别", "协议生效日期", "协议到期日期",
                      "业务金额", "当前保证金", "敞口金额", "备注（是否零保证金开票等）"])
        book.save(path)
        self.assertEqual([], absorb(path))

    def test_xlsx_customer_name_and_credit_amount_aliases_are_fields(self) -> None:
        path = Path(self.tmp.name) / "customer-credit.xlsx"
        book = Workbook()
        sheet = book.active
        sheet["A1"] = "客户名称"
        sheet["B1"] = "合成客户企业"
        sheet.merge_cells("B1:C1")
        sheet["D1"] = "授信额度"
        sheet["E1"] = "500万元"
        sheet["F1"] = "已使用额度"
        sheet["G1"] = "100万元"
        book.save(path)
        rows = absorb(path)
        self.assertEqual("合成客户企业", next(row["value"] for row in rows
                                             if row["key"] == "借款人名称"))
        self.assertEqual("500万元", next(row["value"] for row in rows
                                         if row["key"] == "授信额度"))
        self.assertEqual("100万元", next(row["value"] for row in rows
                                         if row["key"] == "已使用额度"))


if __name__ == "__main__":
    unittest.main()
