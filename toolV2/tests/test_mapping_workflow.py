"""Synthetic mapping workflow integration checks.

These tests exercise the public JSON process boundary and use real DOCX files.
They deliberately keep all data synthetic.
"""
from __future__ import annotations

import hashlib
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
import sqlite3

from docx import Document


TOOL = Path(__file__).resolve().parents[1]


def request(payload: dict) -> dict:
    run = subprocess.run([sys.executable, str(TOOL / "office.py")],
                         input=json.dumps(payload), text=True, encoding="utf-8",
                         capture_output=True, timeout=60)
    if run.returncode:
        raise AssertionError(run.stderr or run.stdout)
    return json.loads(run.stdout)


class MappingWorkflowTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.work = Path(self.tmp.name) / "work"
        self.work.mkdir()

    def tearDown(self) -> None:
        self.tmp.cleanup()

    def source(self, *lines: str) -> Path:
        path = self.work / "source.docx"
        doc = Document()
        for line in lines:
            doc.add_paragraph(line)
        doc.save(path)
        return path

    def target(self, *lines: str, name: str = "target.docx") -> Path:
        path = self.work / name
        doc = Document()
        for line in lines:
            doc.add_paragraph(line)
        doc.save(path)
        return path

    def start(self, source: Path, target: Path, batch: str = "20260925-81") -> dict:
        return request({"action": "start", "work": str(self.work),
                        "source": [source.relative_to(self.work).as_posix()],
                        "targets": [target.relative_to(self.work).as_posix()],
                        "batch": batch})

    def answer_sources(self, result: dict) -> dict:
        while result["status"] == "awaiting_source":
            result = request({"action": "resume", "work": str(self.work),
                              "task_id": result["task_id"],
                              "answers": [{"id": q["id"],
                                           "selected": [q["options"][0]["label"]],
                                           "custom": ""} for q in result["questions"]]})
        return result

    def finish_fill(self, result: dict) -> dict:
        while result["status"] == "awaiting_fill":
            result = request({"action": "resume", "work": str(self.work),
                              "task_id": result["task_id"],
                              "answers": [{"id": q["id"],
                                           "selected": [q["options"][0]["label"]],
                                           "custom": ""} for q in result["questions"]]})
        return result

    def test_unknown_sentence_slot_maps_existing_field_and_reuses_execution(self) -> None:
        source = self.source("借款人名称：合成甲公司", "联系电话：13800000000")
        target = self.target("未标注签约信息 ____")
        result = self.answer_sources(self.start(source, target))
        self.assertEqual("needs_mapping", result["status"], result)
        slot = result["mapping_requests"][0]["positions"][0]
        update = request({"action": "update_positions", "work": str(self.work),
                          "task_id": result["task_id"], "updates": [{
                              "template": "target.docx", "slot_id": slot["id"],
                              "field": "联系电话", "target": slot["target"]}]})
        completed = self.finish_fill(update)
        self.assertEqual("completed", completed["status"], completed)
        output = Path(completed["results"][0]["output"])
        self.assertIn("13800000000", Document(output).paragraphs[0].text)
        count = completed["counters"]["fill_processes"]
        again = self.answer_sources(self.start(source, target))
        again = self.finish_fill(again)
        self.assertEqual(count, again["counters"]["fill_processes"])

    def test_invalid_block_update_does_not_change_existing_rule(self) -> None:
        source = self.source("借款人名称：合成甲公司", "联系电话：13800000000")
        target = self.target("联系电话：________")
        result = self.finish_fill(self.answer_sources(self.start(source, target)))
        self.assertEqual("completed", result["status"], result)
        before = json.dumps(result["results"], ensure_ascii=False, sort_keys=True)
        bad = request({"action": "update_positions", "work": str(self.work),
                       "task_id": result["task_id"], "updates": [{
                           "template": "target.docx", "field": "联系电话",
                           "target": {"kind": "block", "block_index": 0}}]})
        self.assertFalse(bad.get("ok", True))
        after = self.answer_sources(self.start(source, target))
        self.assertEqual(before, json.dumps(after["results"], ensure_ascii=False, sort_keys=True))

    def test_explicit_leave_blank_completes_without_writing_value(self) -> None:
        source = self.source("借款人名称：合成甲公司", "联系电话：13800000000")
        target = self.target("未标注字段 ____")
        result = self.answer_sources(self.start(source, target))
        slot = result["mapping_requests"][0]["positions"][0]
        resumed = request({"action": "update_positions", "work": str(self.work),
                           "task_id": result["task_id"], "updates": [{
                               "template": "target.docx", "slot_id": slot["id"],
                               "leave_blank": True, "reason": "来源没有该字段"}]})
        completed = self.finish_fill(resumed)
        self.assertEqual("completed", completed["status"], completed)
        output = Path(completed["results"][0]["output"])
        self.assertNotIn("13800000000", Document(output).paragraphs[0].text)

    def test_duplicate_paragraphs_can_map_different_subject_fields(self) -> None:
        source = self.source("借款人名称：合成甲公司", "借款人法定代表人：甲代表",
                             "保证人名称：合成乙公司", "保证人法定代表人：乙代表")
        target = self.target("法定代表人：____", "法定代表人：____")
        result = self.answer_sources(self.start(source, target, "20260925-82"))
        self.assertEqual("needs_mapping", result["status"], result)
        positions = result["mapping_requests"][0]["positions"]
        self.assertGreaterEqual(len(positions), 2)
        updates = [{"template": "target.docx", "slot_id": positions[0]["id"],
                    "field": "借款人法定代表人", "target": positions[0]["target"]},
                   {"template": "target.docx", "slot_id": positions[1]["id"],
                    "field": "保证人法定代表人", "target": positions[1]["target"]}]
        completed = self.finish_fill(request({"action": "update_positions", "work": str(self.work),
                                              "task_id": result["task_id"], "updates": updates}))
        self.assertEqual("completed", completed["status"], completed)
        doc = Document(Path(completed["results"][0]["output"]))
        self.assertIn("甲代表", doc.paragraphs[0].text)
        self.assertIn("乙代表", doc.paragraphs[1].text)

    def test_invalid_template_update_does_not_discard_other_template_update(self) -> None:
        source = self.source("借款人：合成甲公司", "联系电话：13800000000")
        good = self.target("客户电话：____", name="good.docx")
        bad = self.target("未知信息：____", name="bad.docx")
        result = self.answer_sources(request({"action": "start", "work": str(self.work),
                                               "source": [source.name],
                                               "targets": [good.name, bad.name],
                                               "batch": "20260925-83"}))
        self.assertEqual("needs_mapping", result["status"], result)

        def first_slot(template: str) -> dict:
            page = request({"action": "read_mapping", "work": str(self.work),
                            "task_id": result["task_id"],
                            "mapping_read": {"section": "positions", "template": template}})
            return page["mapping_page"]["items"][0]

        good_slot, bad_slot = first_slot(good.name), first_slot(bad.name)
        updated = request({"action": "update_positions", "work": str(self.work),
                           "task_id": result["task_id"], "updates": [
                               {"template": good.name, "slot_id": good_slot["id"], "field": "联系电话"},
                               {"template": bad.name, "slot_id": bad_slot["id"], "field": "不存在的来源字段"},
                           ]})
        self.assertEqual("needs_mapping", updated["status"], updated)
        self.assertFalse(updated["ok"], updated)
        self.assertTrue(updated["success_with_rejected"], updated)
        self.assertEqual(1, len(updated.get("rejected_updates", [])), updated)
        self.assertEqual(bad.name, updated["rejected_updates"][0]["template"])
        conn = sqlite3.connect(self.work / "db" / "workflow.db")
        try:
            fields = [row[0] for row in conn.execute(
                "SELECT r.field FROM template_rule r JOIN template t ON t.id=r.template_id "
                "WHERE t.path=?", (str(good),))]
        finally:
            conn.close()
        self.assertIn("联系电话", fields)

    def test_invalid_legacy_target_update_keeps_completed_artifact_without_rerun(self) -> None:
        source = self.source("借款人：合成甲公司", "联系电话：13800000000")
        target = self.target("联系电话：____")
        completed = self.finish_fill(self.answer_sources(self.start(source, target, "20260925-84")))
        self.assertEqual("completed", completed["status"], completed)
        rejected = request({"action": "update_positions", "work": str(self.work),
                            "task_id": completed["task_id"], "updates": [{
                                "template": target.name, "field": "不存在的来源字段",
                                "target": {"kind": "anchor", "anchor": "联系电话：", "max_blank": 20},
                            }]})
        self.assertEqual("completed", rejected["status"], rejected)
        self.assertFalse(rejected["ok"], rejected)
        self.assertTrue(rejected["success_with_rejected"], rejected)
        self.assertEqual(completed["counters"], rejected["counters"])
        self.assertIn("来源中没有字段", rejected["rejected_updates"][0]["reason"])


if __name__ == "__main__":
    unittest.main()
