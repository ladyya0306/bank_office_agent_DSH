"""Regression coverage for DOCX source-table structure in ``absorb.read_lines``."""
from __future__ import annotations

import sys
from pathlib import Path

from docx import Document

TOOL = Path(__file__).resolve().parents[1]
if str(TOOL) not in sys.path:
    sys.path.insert(0, str(TOOL))

from office_kit.absorb import absorb, read_lines  # noqa: E402


def _save(document: Document, path: Path) -> Path:
    document.save(path)
    return path


def test_merged_cells_do_not_create_title_to_title_or_value_to_label_pairs(tmp_path):
    document = Document()
    table = document.add_table(rows=3, cols=4)
    table.cell(0, 0).text = "基本信息"
    table.cell(0, 0).merge(table.cell(1, 0))
    table.cell(0, 1).text = "客户名称"
    table.cell(0, 2).text = "合成企业"
    table.cell(0, 2).merge(table.cell(0, 3))
    table.cell(1, 1).text = "联系电话"
    table.cell(1, 2).text = "13800000000"
    table.cell(2, 0).text = "开户行"
    table.cell(2, 0).merge(table.cell(2, 1))
    table.cell(2, 2).text = "合成支行"
    table.cell(2, 2).merge(table.cell(2, 3))
    path = _save(document, tmp_path / "merged.docx")

    lines = [text for _line, text in read_lines(path)]
    pairs = {(row["key"], row["value"]) for row in absorb(path)}

    assert "基本信息" in lines
    assert ("借款人名称", "合成企业") in pairs
    assert ("联系电话", "13800000000") in pairs
    assert ("开户行", "合成支行") in pairs
    assert not any(value in {"客户名称", "联系电话", "开户行"}
                   for _key, value in pairs)


def test_unmerged_multiple_label_value_groups_remain_supported(tmp_path):
    document = Document()
    row = document.add_table(rows=1, cols=6).rows[0].cells
    for cell, text in zip(row, ("客户名称", "合成企业", "授信额度", "500万元",
                                "联系电话", "13800000000")):
        cell.text = text
    path = _save(document, tmp_path / "unmerged.docx")

    assert {(row["key"], row["value"]) for row in absorb(path)} == {
        ("借款人名称", "合成企业"), ("授信额度", "500万元"),
        ("联系电话", "13800000000"),
    }


def test_complete_colon_value_and_space_separated_unknown_text_stay_readable(tmp_path):
    document = Document()
    table = document.add_table(rows=1, cols=3)
    table.cell(0, 0).text = "开户行：中国：合成支行"
    table.cell(0, 1).text = "附加说明 合成内容"
    table.cell(0, 2).text = "备注"
    path = _save(document, tmp_path / "complete-and-unknown.docx")

    lines = [text for _line, text in read_lines(path)]
    rows = absorb(path)

    bank = next(row for row in rows if row["key"] == "开户行")
    assert bank["value"] == "中国：合成支行"
    assert "附加说明 合成内容" in lines
    assert "备注" in lines


def test_plain_unmerged_table_with_unknown_cells_is_not_forced_into_facts(tmp_path):
    document = Document()
    row = document.add_table(rows=1, cols=3).rows[0].cells
    for cell, text in zip(row, ("自定义栏目", "自定义内容", "补充说明")):
        cell.text = text
    path = _save(document, tmp_path / "unknown.docx")

    assert [text for _line, text in read_lines(path)] == ["自定义栏目", "自定义内容", "补充说明"]
    assert absorb(path) == []


def test_unknown_two_column_text_is_preserved_and_explicit_unknown_label_is_available(tmp_path):
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "产品特征"
    table.cell(0, 1).text = "耐腐蚀"
    table.cell(1, 0).text = "新字段："
    table.cell(1, 1).text = "某值"
    path = _save(document, tmp_path / "unknown-form.docx")

    assert "产品特征" in [text for _line, text in read_lines(path)]
    assert "耐腐蚀" in [text for _line, text in read_lines(path)]
    assert {(row["key"], row["value"]) for row in absorb(path)} == {("新字段", "某值")}


def test_same_text_neighbours_are_never_paired(tmp_path):
    document = Document()
    row = document.add_table(rows=1, cols=2).rows[0].cells
    row[0].text = "项目"
    row[1].text = "项目"
    path = _save(document, tmp_path / "same-text.docx")

    assert [text for _line, text in read_lines(path)] == ["项目", "项目"]
    assert absorb(path) == []


def test_empty_colon_label_is_preserved_until_a_neighbour_supplies_its_value(tmp_path):
    document = Document()
    table = document.add_table(rows=1, cols=3)
    table.cell(0, 0).text = "待补字段："
    table.cell(0, 1).text = ""
    table.cell(0, 2).text = "说明"
    path = _save(document, tmp_path / "empty-colon.docx")

    # LINE_RE deliberately requires a non-empty value.  A lone trailing colon
    # must remain source text, and must not consume a non-adjacent cell.
    assert [text for _line, text in read_lines(path)] == ["待补字段：", "说明"]
    assert absorb(path) == []
