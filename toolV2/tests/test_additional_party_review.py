"""Optional signer blocks ask once per physical block, not per source value."""
from pathlib import Path

from docx import Document

from office_kit.harness import build_fill_plan
from office_kit.store_v2 import StoreV2
from office_kit.template_slots import discover_slots


def test_optional_blocks_are_ask_rows_with_stable_physical_block_identity(tmp_path: Path):
    path = tmp_path / 'generic.docx'
    doc = Document()
    doc.add_paragraph('\u501f\u6b3e\u4eba'); doc.add_paragraph('\u4f01\u4e1a\u540d\u79f0\uff1a____')
    for text in ('\u53ef\u6839\u636e\u7b7e\u7f72\u65b9\u4e2a\u6570\u589e\u51cf', '\u4f01\u4e1a\u540d\u79f0\uff1a____',
                 '\u6cd5\u5b9a\u4ee3\u8868\u4eba\uff1a____', '\u4f4f\u6240\uff1a____', '\u8054\u7cfb\u4eba\uff1a____',
                 '\u5f00\u6237\u884c\u53ca\u8d26\u53f7\uff1a____', '\u53ef\u6839\u636e\u7b7e\u7f72\u65b9\u4e2a\u6570\u589e\u51cf', '\u4f01\u4e1a\u540d\u79f0\uff1a____',
                 '\u4fdd\u8bc1\u4eba2', '\u4f01\u4e1a\u540d\u79f0\uff1a____'):
        doc.add_paragraph(text)
    doc.save(path)
    batch = '20260926-95'
    with StoreV2(tmp_path / 'data.db') as store:
        store.batch_scope = batch
        borrower, _ = store.ensure_entity('\u5408\u6210\u501f\u6b3e\u4eba')
        guarantor, _ = store.ensure_entity('\u5408\u6210\u4fdd\u8bc1\u4eba2')
        store.set_role(borrower, '\u501f\u6b3e\u4eba', evidence='\u501f\u6b3e\u4eba', batch_no=batch)
        store.set_role(guarantor, '\u4fdd\u8bc1\u4eba', evidence='\u4fdd\u8bc1\u4eba2', batch_no=batch)
        store.put_fact(batch, '\u501f\u6b3e\u4eba\u540d\u79f0', '\u5408\u6210\u501f\u6b3e\u4eba', entity_id=borrower)
        store.put_fact(batch, '\u4fdd\u8bc1\u4eba\u540d\u79f0', '\u5408\u6210\u4fdd\u8bc1\u4eba2', entity_id=guarantor)
        tid = store.register_template(path, batch)
        slots = discover_slots(path)
        first = next(slot for slot in slots if slot['target']['paragraph_index'] == 1)
        rest = [slot['target'] for slot in slots if slot is not first]
        store.add_rule(tid, '\u501f\u6b3e\u4eba\u540d\u79f0', first['label'], first['target'], confidence=1, decided_by='model', batch_no=batch)
        store.add_rule(tid, '\u4fdd\u8bc1\u4eba2\u540d\u79f0', '\u4f01\u4e1a\u540d\u79f0',
                       {'kind': 'multi', 'targets': rest}, confidence=1, decided_by='model', batch_no=batch)
        plan = build_fill_plan(store, [path], batch_no=batch, run_id='synthetic')
    rows = [row for row in plan['rows'] if row['kind'] == 'slot']
    optional = [row for row in rows if row['target_subject'].get('additional_party')]
    assert len(optional) == 6 and all(row['decision'] == 'ask' for row in optional)
    reviews = [row['relationship_review'] for row in optional]
    first_block = [review for review in reviews if review['block_id'] == reviews[0]['block_id']]
    assert len(first_block) == 5
    assert {review['owner_eid'] for review in reviews} == {guarantor}
    assert {review['owner_name'] for review in reviews} == {'\u5408\u6210\u4fdd\u8bc1\u4eba2'}
    explicit = [row for row in rows if row['target_subject'].get('number') == 2]
    assert explicit and all(row['decision'] == 'auto' and row['relationship_review'] is None for row in explicit)
    second = next(row for row in rows if row['target']['paragraph_index'] == 9)
    assert second['relationship_review']['block_id'] != reviews[0]['block_id']
