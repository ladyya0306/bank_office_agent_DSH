"""Workbook label references are bounded context evidence, never mappings."""
from __future__ import annotations

from openpyxl import Workbook

from workflow.mapping_view import page


def context_task(tmp_path, *, with_references: bool):
    book = Workbook()
    first = book.active
    first.title = 'Sheet1'
    second = book.create_sheet('Sheet2')
    other = book.create_sheet('Other')
    second['C2'] = '\u7ecf\u529e\u673a\u6784\u540d\u79f0\uff1a____'
    if with_references:
        first['D2'] = '\u7ecf\u529e\u673a\u6784\u540d\u79f0\uff1a\u5408\u6210\u652f\u884c\u7532'
        other['A1'] = '\u7ecf\u529e\u673a\u6784\u540d\u79f0\uff1a\u5408\u6210\u652f\u884c\u4e59'
        other['A2'] = '\u7ecf\u529e\u673a\u6784\u540d\u79f0\uff1a____'
        other['A3'] = '\u5176\u4ed6\u7ecf\u529e\u673a\u6784\u540d\u79f0\uff1a\u65e0\u5173\u652f\u884c'
        other['A4'] = '\u65e0\u5173\u4fe1\u606f\uff1a\u4e0d\u5e94\u6cc4\u9732'
    path = tmp_path / 'reference.xlsx'
    book.save(path)
    task = {
        'last_work': str(tmp_path),
        'source_signature': 'synthetic',
        'available_fields': [],
        'documents': {'reference.xlsx': {
            'status': 'needs_mapping', 'slot_hash': 'synthetic', 'slot_version': 1,
            'current_positions': [], 'blank_slots': {},
            'slots': [{'id': 'sheet2-c2', 'label': '\u7ecf\u529e\u673a\u6784\u540d\u79f0', 'context': {},
                       'target': {'kind': 'xlsx_cell', 'sheet': 'Sheet2', 'cell': 'C2',
                                  'expected_text': '\u7ecf\u529e\u673a\u6784\u540d\u79f0\uff1a____',
                                  'span_start': 7, 'span_end': 11}}],
            'unmapped_slots': [{'id': 'sheet2-c2', 'label': '\u7ecf\u529e\u673a\u6784\u540d\u79f0', 'context': {},
                                'target': {'kind': 'xlsx_cell', 'sheet': 'Sheet2', 'cell': 'C2',
                                           'expected_text': '\u7ecf\u529e\u673a\u6784\u540d\u79f0\uff1a____',
                                           'span_start': 7, 'span_end': 11}}],
        }},
    }
    return task


def test_xlsx_context_shows_only_exact_same_label_references(tmp_path):
    task = context_task(tmp_path, with_references=True)
    result = page(task,
                  {'section': 'context', 'template': 'reference.xlsx', 'slot_id': 'sheet2-c2'})
    evidence = next(item['same_label_references'] for item in result['items']
                    if 'same_label_references' in item)
    assert evidence == [
        {'sheet': 'Sheet1', 'cell': 'D2', 'text': '\u7ecf\u529e\u673a\u6784\u540d\u79f0\uff1a\u5408\u6210\u652f\u884c\u7532'},
        {'sheet': 'Other', 'cell': 'A1', 'text': '\u7ecf\u529e\u673a\u6784\u540d\u79f0\uff1a\u5408\u6210\u652f\u884c\u4e59'},
    ]
    assert '\u65e0\u5173\u652f\u884c' not in str(result)
    assert '____' not in str(evidence)
    positions = page(task, {'section': 'positions', 'template': 'reference.xlsx'})
    assert positions['items'][0]['same_label_references'] == evidence


def test_xlsx_context_omits_reference_key_when_none_exist(tmp_path):
    result = page(context_task(tmp_path, with_references=False),
                  {'section': 'context', 'template': 'reference.xlsx', 'slot_id': 'sheet2-c2'})
    assert not any('same_label_references' in item for item in result['items'])
