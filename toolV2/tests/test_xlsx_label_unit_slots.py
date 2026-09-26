"""Synthetic Excel label/value-cell pairing and amount-unit regressions."""
from pathlib import Path
from tempfile import TemporaryDirectory

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Border, Side

from office_kit.common import OfficeKitError
from office_kit.doc_fill import XlsxEngine
from office_kit.target_validation import value_target_issue
from office_kit.template_slots import discover_slots
from office_kit.value_fit import fit_value, unit_wrapper_span, value_fit_issue


def _template(path: Path) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = 'Sheet1'
    sheet['A3'] = '客户名称：'
    sheet['B3'].border = Border(bottom=Side(style='thin'))
    sheet['D3'] = '额度：'
    sheet['E3'] = '人民币  万元'
    sheet['A4'] = '本次业务金额：'
    sheet['B4'] = '人民币   万元'
    sheet['D4'] = '已使用额度：'
    sheet['E4'] = '人民币  万元'
    book.save(path)
    book.close()


def test_right_unit_value_cells_inherit_label_and_describe_replaced_targets():
    with TemporaryDirectory() as tmp:
        template = Path(tmp) / '合成台账.xlsx'
        _template(template)

        slots = discover_slots(template)
        by_cell = {slot['target']['cell']: slot for slot in slots}
        assert {'B3', 'E3', 'B4', 'E4'} <= set(by_cell)
        assert not {'A4', 'D3', 'D4'} & set(by_cell)
        assert by_cell['E3']['label'] == '额度'
        assert by_cell['B4']['label'] == '本次业务金额'
        assert by_cell['E4']['label'] == '已使用额度'
        assert by_cell['B4']['target']['label_cell'] == 'A4'
        assert by_cell['B4']['replaces_targets'] == [
            {'kind': 'xlsx_cell', 'sheet': 'Sheet1', 'cell': 'A4', 'anchor': '本次业务金额：'},
            {'kind': 'xlsx_cell', 'sheet': 'Sheet1', 'cell': 'A4',
             'expected_text': '本次业务金额：', 'span_start': 7, 'span_end': 7},
        ]


def test_unit_wrappers_fill_once_and_keep_zero_without_repeating_units():
    with TemporaryDirectory() as tmp:
        template = Path(tmp) / '合成台账.xlsx'
        _template(template)
        slots = {slot['target']['cell']: slot for slot in discover_slots(template)}
        engine = XlsxEngine(template)
        engine.fill_xlsx_cell(slots['E3']['target'], '800万元')
        engine.fill_xlsx_cell(slots['B4']['target'], '800万元')
        engine.fill_xlsx_cell(slots['E4']['target'], '0万元')
        output = Path(tmp) / 'result.xlsx'
        engine.save(output)

        check = load_workbook(output, data_only=False)
        sheet = check['Sheet1']
        assert sheet['D3'].value == '额度：'
        assert sheet['E3'].value == '人民币800万元'
        assert sheet['A4'].value == '本次业务金额：'
        assert sheet['B4'].value == '人民币800万元'
        assert sheet['D4'].value == '已使用额度：'
        assert sheet['E4'].value == '人民币0万元'
        check.close()


def test_legacy_label_cell_write_is_rejected_when_unit_value_cell_exists():
    with TemporaryDirectory() as tmp:
        template = Path(tmp) / '合成台账.xlsx'
        _template(template)
        engine = XlsxEngine(template)
        with pytest.raises(OfficeKitError, match='右侧 E4 已有单位和空位'):
            engine.fill_xlsx_cell(
                {'kind': 'xlsx_cell', 'sheet': 'Sheet1', 'cell': 'D4',
                 'anchor': '已使用额度：'}, '0万元')
        with pytest.raises(OfficeKitError, match='右侧 E3 已有单位和空位'):
            engine.fill_xlsx_cell(
                {'kind': 'xlsx_cell', 'sheet': 'Sheet1', 'cell': 'D3',
                 'anchor': '额度：'}, '800万元')
        with pytest.raises(OfficeKitError, match='右侧 B4 已有单位和空位'):
            engine.fill_xlsx_cell(
                {'kind': 'xlsx_cell', 'sheet': 'Sheet1', 'cell': 'A4',
                 'expected_text': '本次业务金额：', 'span_start': 7, 'span_end': 7}, '800万元')


def test_amount_unit_fit_is_strict_and_zero_is_not_dropped():
    text = '人民币   万元'
    start, end = unit_wrapper_span(text)
    assert (start, end) == (3, 6)
    assert fit_value(text, start, end, '0万元') == '0'
    assert fit_value(text, start, end, '800万元') == '800'
    assert fit_value(text, start, end, '人民币800万元') == '800'
    assert fit_value(text, start, end, '800元') == '0.08'
    assert fit_value('金额____元人民币', 2, 6, '800万元') == '8000000'
    assert fit_value('金额____亿元', 2, 6, '125000000元') == '1.25'
    assert fit_value('金额____万元', 2, 6, '0.25亿元') == '2500'
    assert value_fit_issue('金额____元人民币', 2, 6, '800万元') is None
    with pytest.raises(OfficeKitError, match='不是可识别的人民币数字金额'):
        fit_value(text, start, end, '800美元')
    assert '跨币种' in value_fit_issue('金额____元人民币', 2, 6, '800美元')
    underscore_text = '人民币____万元'
    underscore_span = unit_wrapper_span(underscore_text)
    assert underscore_span is not None
    assert fit_value(underscore_text, *underscore_span, '800万元') == '800'
    assert unit_wrapper_span('人民币____万元（大写）____') is None


def test_preview_value_check_accepts_exact_conversion_and_reports_cross_currency():
    with TemporaryDirectory() as tmp:
        template = Path(tmp) / '金额模板.xlsx'
        book = Workbook()
        book.active['A1'] = '金额____元人民币'
        book.save(template)
        book.close()
        target = {'kind': 'xlsx_cell', 'sheet': 'Sheet', 'cell': 'A1',
                  'expected_text': '金额____元人民币', 'span_start': 2, 'span_end': 6}
        assert value_target_issue(template, target, '800万元') is None
        issue = value_target_issue(template, target, '800美元')
        assert issue and '跨币种' in issue
