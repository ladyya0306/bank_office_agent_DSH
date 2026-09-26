"""Confirm one uncertain relationship without asking for known facts again."""
import unittest
import tempfile
from pathlib import Path
from unittest.mock import patch
from workflow import review
from office_kit.fill_decisions import fingerprint


class RelationshipCardsTests(unittest.TestCase):
    def test_real_company_and_representative_facts_share_one_saved_relationship_answer(self):
        from docx import Document
        from office_kit.store_v2 import StoreV2
        from office_kit.harness import build_fill_plan
        from office_kit.template_slots import discover_slots
        from office_kit.fill_decisions import saved_choices
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / '合同.docx'; db = Path(tmp) / 'workflow.db'
            doc = Document(); doc.add_paragraph('可根据签署方个数增减')
            labels = ('企业名称', '法定代表人', '住所', '联系电话', '开户行及账号')
            for label in labels: doc.add_paragraph(label + '：____')
            doc.save(path)
            batch = '20260926-96'
            with StoreV2(db) as store:
                store.batch_scope = batch
                company, _ = store.ensure_entity('合成保证企业')
                person, _ = store.ensure_entity('合成代表', personal=True)
                store.set_role(company, '保证人', evidence='保证人2', batch_no=batch)
                values = [('保证人名称', '合成保证企业'), ('保证人法定代表人', '合成代表'),
                          ('地址', '合成注册地址'), ('联系电话', '13800000000'), ('开户行及账号', '合成银行 SYNTH-123')]
                for key, value in values: store.put_fact(batch, key, value, entity_id=company)
                fields = ('保证人2名称', '保证人2法定代表人', '保证人2地址', '保证人2联系电话', '保证人2开户行及账号')
                tid = store.register_template(path, batch)
                for slot, field in zip(discover_slots(path), fields):
                    store.add_rule(tid, field, slot['label'], slot['target'], confidence=1, decided_by='model', batch_no=batch)
                plan = build_fill_plan(store, [path], batch_no=batch, run_id='synthetic')
            plan.update(db_path=str(db), template_files=[str(path)])
            task = {'batch': batch, 'documents': {'合同.docx': {'status': 'pending', 'plan': plan}}}
            cards = review.questions(task)
            self.assertEqual(1, len(cards))
            self.assertEqual(5, len(cards[0]['locations']))
            self.assertEqual(5, len(set(r['value'] for r in plan['rows'])))
            self.assertEqual(person, next(r['entity_id'] for r in plan['rows'] if r['field'].endswith('法定代表人')))
            answers = [{'id': cards[0]['id'], 'selected': [review.CONFIRM_PARTY]}]
            review.validate(cards, answers); review.save(task, answers)
            self.assertEqual([], review.questions(task))
            self.assertEqual('1,2,3,4,5', saved_choices(plan)[0]['select'])

    def task(self):
        rows = []
        for n, (field, value) in enumerate((('名称', '合成企业'), ('法人', '合成人员'),
                                           ('地址', '合成地址'), ('电话', '13800000000'), ('银行', '合成银行')), 1):
            rows.append({'n': n, 'field': field, 'value': value, 'kind': 'slot',
                         'decision': 'ask', 'template_sha256': 'synthetic', 'entity_id': 7,
                         'relationship_review': {'block_id': 'paragraph-8', 'owner_eid': 7, 'owner_name': '合成企业'},
                         'target': {'kind': 'anchor', 'anchor': field}})
        return {'batch': 'synthetic', 'documents': {'合同.docx': {'status': 'pending', 'plan': {'rows': rows}}}}

    def questions(self, task):
        with patch.object(review, '_fact_catalog', return_value={}), \
             patch.object(review, 'saved_choices', side_effect=lambda p: ({}, [r['n'] for r in p['rows']])):
            return review.questions(task)

    def test_one_relationship_card_covers_distinct_source_values(self):
        task = self.task()
        cards = self.questions(task)
        self.assertEqual(1, len(cards))
        self.assertEqual(5, len(cards[0]['locations']))
        self.assertIn('合同.docx', cards[0]['header'])
        self.assertIn('合成企业', cards[0]['header'])
        self.assertIn('合成人员', cards[0]['detail'])

    def test_accept_selects_each_value_blank_covers_only_this_block(self):
        for option, key, expected in ((review.CONFIRM_PARTY, 'select', '1,2,3,4,5'),
                                      (review.BLANK, 'blank', ['1', '2', '3', '4', '5'])):
            task = self.task(); cards = self.questions(task)
            answers = [{'id': cards[0]['id'], 'selected': [option]}]
            self.assertTrue(review.validate(cards, answers))
            choices = {'apply_all': True, 'select': '', 'blank': [], 'new': [], 'use': []}
            with patch.object(review, 'saved_choices', return_value=(choices, [])), \
                 patch.object(review, 'save_choices') as save:
                review.save(task, answers)
            self.assertEqual(expected, save.call_args.args[1][key])
            self.assertEqual([], save.call_args.args[1]['new'])

    def test_separate_blocks_and_owners_remain_separate(self):
        task = self.task(); rows = task['documents']['合同.docx']['plan']['rows']
        rows[-1]['relationship_review'] = {**rows[-1]['relationship_review'], 'block_id': 'paragraph-20'}
        rows[-2]['relationship_review'] = {**rows[-2]['relationship_review'], 'owner_eid': 9}
        self.assertEqual(3, len(self.questions(task)))

    def test_custom_text_cannot_be_copied_into_all_different_fields(self):
        cards = self.questions(self.task())
        with self.assertRaisesRegex(ValueError, '主体关系'):
            review.validate(cards, [{'id': cards[0]['id'], 'custom': '另一种地址', 'selected': []}])
        self.assertTrue(review.validate(cards, [{'id': cards[0]['id'], 'custom': '合成企业', 'selected': []}]))

    def test_relationship_change_invalidates_saved_answer_identity(self):
        task = self.task(); plan = task['documents']['合同.docx']['plan']; row = plan['rows'][0]
        before = fingerprint(plan, row)
        row['relationship_review']['owner_eid'] = 99
        self.assertNotEqual(before, fingerprint(plan, row))
