"""Committed evidence snapshots remain readable while a writer holds its lock."""
from docx import Document
import pytest
from workflow import storage
from workflow.runner import dispatch


def test_evidence_read_does_not_take_writer_lock_or_mutate_task(tmp_path):
    source = Document(); source.add_paragraph('借款人：合成企业'); source.save(tmp_path / 'source.docx')
    target = Document(); target.add_paragraph('补充说明：____'); target.save(tmp_path / 'target.docx')
    result = dispatch({'action': 'start', 'work': str(tmp_path), 'source': ['source.docx'], 'targets': ['target.docx']})
    counters = result['counters']
    query = {'work': str(tmp_path), 'task_id': result['task_id']}
    with storage.work_lock(tmp_path):
        for section in ('fields', 'positions', 'source', 'document'):
            page = dispatch({**query, 'action': 'read_mapping', 'mapping_read': {'section': section, 'template': 'target.docx'}})
            assert page['ok'] and page['counters'] == counters
        with pytest.raises(RuntimeError, match='已有填报'):
            dispatch({**query, 'action': 'status'})
    assert dispatch({**query, 'action': 'status'})['counters'] == counters
