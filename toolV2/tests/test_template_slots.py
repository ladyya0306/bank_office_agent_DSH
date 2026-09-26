from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Border, Side

from office_kit.template_slots import discover_slots


def test_word_discovers_inline_bracket_colon_and_protected_slots():
    with TemporaryDirectory() as temp:
        path = Path(temp) / "template.docx"
        doc = Document()
        doc.add_paragraph("借款人：  合成公司（    ）")
        doc.add_paragraph("法定代表人签字：")
        doc.sections[0].header.paragraphs[0].text = "联系电话："
        doc.save(path)
        slots = discover_slots(path)
        assert any(x["label"] == "借款人" and x["target"]["span_start"] < x["target"]["span_end"] for x in slots)
        assert any(x["target"]["part"].startswith("word/header") for x in slots)
        protected = next(x for x in slots if "签字" in x["label"])
        assert protected["protected"] is True
        assert protected["target"]["span_start"] == protected["target"]["span_end"]


def test_word_discovers_blank_inside_chinese_corner_brackets():
    with TemporaryDirectory() as temp:
        path = Path(temp) / "template.docx"
        doc = Document()
        doc.add_paragraph("编号为【        】的《综合授信合同》")
        doc.save(path)

        slots = discover_slots(path)
        assert len(slots) == 1
        target = slots[0]["target"]
        assert target["span_start"] == len("编号为【")
        assert target["span_end"] == len("编号为【        ")


def test_word_table_merged_cell_is_once_and_is_addressed_as_cell():
    with TemporaryDirectory() as temp:
        path = Path(temp) / "template.docx"
        doc = Document()
        table = doc.add_table(rows=1, cols=3)
        table.cell(0, 0).merge(table.cell(0, 1))
        table.cell(0, 0).text = "借款人名称"
        doc.save(path)
        slots = [x for x in discover_slots(path) if x["target"]["kind"] == "cell"]
        assert len(slots) == 1
        assert slots[0]["target"] == {"kind": "cell", "table": 0, "row": 0, "col": 2}


def test_xlsx_discovers_text_blank_and_formatted_adjacent_label_cell():
    with TemporaryDirectory() as temp:
        path = Path(temp) / "template.xlsx"
        book = Workbook(); sheet = book.active
        sheet["A1"] = "借款人："
        sheet["B1"].border = Border(bottom=Side(style="thin"))
        sheet["A2"] = "金额（    ）万元"
        book.save(path); book.close()
        slots = discover_slots(path)
        assert any(x["target"].get("cell") == "B1" and x["target"]["kind"] == "xlsx_cell" for x in slots)
        assert any(x["target"].get("cell") == "A2" and x["target"]["kind"] == "xlsx_cell"
                   and "expected_text" in x["target"] for x in slots)


def test_ids_ignore_file_path_but_keep_exact_target_content():
    with TemporaryDirectory() as temp:
        one, two = Path(temp) / "one.docx", Path(temp) / "two.docx"
        for path in (one, two):
            doc = Document(); doc.add_paragraph("借款人：  "); doc.save(path)
        assert [x["id"] for x in discover_slots(one)] == [x["id"] for x in discover_slots(two)]


def test_word_column_spacing_without_placeholder_syntax_is_not_a_slot():
    with TemporaryDirectory() as temp:
        path = Path(temp) / "template.docx"
        doc = Document()
        doc.add_paragraph("说明文字第一栏                    第二栏标题")
        doc.add_paragraph("已填写公司名称                    已填写合同编号")
        doc.save(path)
        assert discover_slots(path) == []


def test_mixed_company_and_signature_only_protects_signature_slot():
    with TemporaryDirectory() as temp:
        path = Path(temp) / "template.docx"
        doc = Document()
        doc.add_paragraph("经核实，    公司法定代表人签字真实有效，合同号：")
        doc.save(path)
        slots = discover_slots(path)
        company = next(x for x in slots if x["target"]["span_start"] < 8)
        contract = next(x for x in slots if x["target"]["span_start"] == x["target"]["span_end"])
        assert company["protected"] is False
        assert contract["protected"] is False


def test_nested_word_table_does_not_shift_top_level_xml_engine_table_index():
    with TemporaryDirectory() as temp:
        path = Path(temp) / "template.docx"
        doc = Document()
        outer = doc.add_table(rows=1, cols=1)
        nested = outer.cell(0, 0).add_table(rows=1, cols=1)
        nested.cell(0, 0).text = "内层名称：____"
        top_level = doc.add_table(rows=1, cols=2)
        top_level.cell(0, 0).text = "外层名称"
        doc.save(path)
        slots = discover_slots(path)
        cells = [x["target"] for x in slots if x["target"]["kind"] == "cell"]
        assert cells == [{"kind": "cell", "table": 1, "row": 0, "col": 1}]
        assert any(x["target"]["kind"] == "anchor" and "内层名称" in x["target"]["expected_text"] for x in slots)


def test_underlined_continuous_blank_and_colon_insertion_become_one_slot():
    with TemporaryDirectory() as temp:
        path = Path(temp) / "template.docx"
        doc = Document(); paragraph = doc.add_paragraph("名称：")
        first = paragraph.add_run(" "); first.font.underline = True
        second = paragraph.add_run(" "); second.font.underline = True
        doc.save(path)
        slots = discover_slots(path)
        assert len(slots) == 1
        target = slots[0]["target"]
        assert target["span_start"] == len("名称：") and target["span_end"] == len("名称：  ")


def test_underlined_padding_around_printed_text_is_not_a_fill_slot():
    with TemporaryDirectory() as temp:
        path = Path(temp) / "template.docx"
        doc = Document()
        paragraph = doc.add_paragraph()
        for blank in ("       ", "        "):
            run = paragraph.add_run(blank)
            run.font.underline = True
        printed = paragraph.add_run(" 财务主管  ")
        printed.font.underline = True
        interior = doc.add_paragraph()
        semantic_gap = interior.add_run("岗位  名称")
        semantic_gap.font.underline = True
        doc.save(path)

        slots = discover_slots(path)
        assert len(slots) == 2
        by_text = {slot["target"]["expected_text"]: slot["target"] for slot in slots}
        assert by_text["                财务主管  "]["span_start"] == 0
        assert by_text["                财务主管  "]["span_end"] == 15
        assert by_text["岗位  名称"]["span_start"] == 2
        assert by_text["岗位  名称"]["span_end"] == 4


def test_merged_table_and_two_paragraph_short_label_use_adjacent_cell():
    with TemporaryDirectory() as temp:
        path = Path(temp) / "template.docx"
        doc = Document(); table = doc.add_table(rows=1, cols=4)
        label = table.cell(0, 0).merge(table.cell(0, 1))
        label.text = "借款人"
        label.add_paragraph("名称：")
        table.cell(0, 2).merge(table.cell(0, 3))
        doc.save(path)
        slots = [x for x in discover_slots(path) if x["target"]["kind"] == "cell"]
        assert slots[0]["target"] == {"kind": "cell", "table": 0, "row": 0, "col": 2}
        assert slots[0]["logical_location"] == "table[0].r[0].c[1]"


def test_merged_table_labels_use_only_the_adjacent_value_cell():
    with TemporaryDirectory() as temp:
        path = Path(temp) / "template.docx"
        doc = Document()
        for label in ("授信申请人：", "担保人1：", "担保人2："):
            table = doc.add_table(rows=1, cols=5)
            table.cell(0, 0).merge(table.cell(0, 3)).text = label
            # Column 4 is the explicit blank value cell.
        sentence = doc.add_paragraph("合同编号：")
        blank = sentence.add_run("    ")
        blank.font.underline = True
        sentence.add_run("请核对")
        doc.save(path)

        slots = discover_slots(path)
        cells = [slot for slot in slots if slot["target"]["kind"] == "cell"]
        anchors = [slot for slot in slots if slot["target"]["kind"] == "anchor"]
        assert [(slot["label"], slot["target"]["table"], slot["target"]["col"])
                for slot in cells] == [
                    ("授信申请人", 0, 4), ("担保人1", 1, 4), ("担保人2", 2, 4)]
        assert not any(slot["target"]["expected_text"].endswith("：")
                       and slot["label"] in {"授信申请人", "担保人1", "担保人2"}
                       for slot in anchors)
        assert any(slot["target"]["expected_text"] == "合同编号：    请核对"
                   and slot["target"]["span_start"] < slot["target"]["span_end"]
                   for slot in anchors)


def test_company_name_with_unit_seal_is_fillable_but_signature_is_protected():
    with TemporaryDirectory() as temp:
        path = Path(temp) / "template.docx"
        doc = Document()
        doc.add_paragraph("公司名称（单位公章）：")
        doc.add_paragraph("法定代表人签字：")
        doc.save(path)
        slots = discover_slots(path)
        company = next(x for x in slots if x["label"] == "公司名称（单位公章）")
        signature = next(x for x in slots if "签字" in x["label"])
        assert company["protected"] is False
        assert signature["protected"] is True
