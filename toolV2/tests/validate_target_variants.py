"""Independent synthetic target checks; records failures without changing expectations."""
from pathlib import Path
import hashlib
import json
import subprocess
import sys
import tempfile
import time
from docx import Document
from docx.shared import Pt
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Border, Side

ROOT = Path(__file__).resolve().parents[1]
VALUE = '13800001234'


def call(payload):
    cp = subprocess.run([sys.executable, str(ROOT / 'office.py')],
                        input=json.dumps(payload), encoding='utf-8',
                        capture_output=True, timeout=90)
    return json.loads(cp.stdout)


def run_case(kind):
    started = time.monotonic()
    with tempfile.TemporaryDirectory(prefix='toolv2_target_validation_') as tmp:
        work = Path(tmp)
        source = work / '合成来源.docx'
        doc = Document()
        doc.add_paragraph('借款人：合成验证公司')
        doc.add_paragraph('联系电话：' + VALUE)
        doc.save(source)
        target = work / ('目标.xlsx' if kind.startswith('excel') else '目标.docx')
        if kind.startswith('excel'):
            wb = Workbook()
            ws = wb.active
            ws['A1'] = '联系电话：' if kind == 'excel_inline' else '联系电话'
            if kind != 'excel_inline':
                ws['B1'].border = Border(bottom=Side(style='thin'))
            ws['F5'] = '=1+2'
            ws.column_dimensions['B'].width = 27
            if kind == 'excel_merged':
                ws.merge_cells('B1:D1')
            if kind == 'excel_second_sheet':
                ws['A1'] = '保留说明'
                ws = wb.create_sheet('填写页')
                ws['A1'] = '联系电话'
                ws['B1'].border = Border(bottom=Side(style='thin'))
            wb.save(target)
            wb.close()
        else:
            doc = Document()
            if kind == 'word_table':
                row = doc.add_table(rows=1, cols=2).rows[0]
                row.cells[0].text = '联系电话'
            elif kind == 'word_header':
                doc.sections[0].header.paragraphs[0].text = '联系电话：'
                doc.add_paragraph('正文保留')
            else:
                p = doc.add_paragraph()
                for text in (['联系', '电话', '：', '________'] if kind == 'word_split_runs'
                             else ['联系电话：________']):
                    r = p.add_run(text)
                    r.bold = True
                    r.font.size = Pt(11)
            doc.save(target)
        original_hashes = [hashlib.sha256(p.read_bytes()).hexdigest() for p in (source, target)]
        payload = dict(action='start', work=str(work), source=[str(source)],
                       targets=[str(target)], batch='20260925-01')
        result = call(payload)
        seen = set()
        rounds = 0
        for _ in range(5):
            if result.get('status') not in ('awaiting_source', 'awaiting_fill'):
                break
            signature = json.dumps(result['questions'], sort_keys=True)
            assert signature not in seen, '重复提出已回答的同一组问题'
            seen.add(signature)
            rounds += 1
            # Synthetic, predeclared truth only. Never run against user materials.
            result = call(dict(action='resume', work=str(work), task_id=result['task_id'],
                               answers=[dict(id=q['id'], selected=[q['options'][0]['label']], custom='')
                                        for q in result['questions']]))
        assert original_hashes == [hashlib.sha256(p.read_bytes()).hexdigest() for p in (source, target)], '原件发生变化'
        assert result.get('status') == 'completed', '未填完: ' + json.dumps(result.get('issues', result), ensure_ascii=False)[:1600]
        output = Path(result['results'][0]['output'])
        if kind.startswith('excel'):
            wb = load_workbook(output, data_only=False)
            ws = wb['填写页'] if kind == 'excel_second_sheet' else wb.active
            if kind == 'excel_inline':
                assert ws['A1'].value == '联系电话：' + VALUE, '标签后面的值不正确'
            else:
                assert str(ws['B1'].value) == VALUE, f'B1落点错误: {ws["B1"].value!r}'
                assert ws['A1'].value == '联系电话', '标签格被修改'
            assert wb.worksheets[0]['F5'].value == '=1+2', '原公式变动'
            assert wb.worksheets[0].column_dimensions['B'].width == 27, '列宽变动'
            if kind == 'excel_merged':
                assert 'B1:D1' in [str(r) for r in ws.merged_cells.ranges], '合并区域变动'
            wb.close()
        else:
            doc = Document(output)
            if kind == 'word_table':
                assert doc.tables[0].cell(0, 1).text == VALUE, 'Word表格右侧格未填入正确值'
                assert doc.tables[0].cell(0, 0).text == '联系电话', '标签被修改'
            elif kind == 'word_header':
                assert VALUE in doc.sections[0].header.paragraphs[0].text, '页眉填写位置遗漏'
                assert doc.paragraphs[0].text == '正文保留', '正文被修改'
            else:
                assert VALUE in doc.paragraphs[0].text, '段落漏填'
                assert all(r.bold and r.font.size == Pt(11) for r in doc.paragraphs[0].runs if r.text), '字体或加粗改变'
        again = call(payload)
        assert not again.get('questions'), '相同任务重新提问'
        assert again['counters']['fill_processes'] == result['counters']['fill_processes'], '相同任务重复填报'
        return dict(case=kind, passed=True, question_rounds=rounds,
                    seconds=round(time.monotonic() - started, 2))


def main():
    results = []
    for kind in ('word_inline', 'word_split_runs', 'word_table', 'word_header',
                 'excel_inline', 'excel_adjacent', 'excel_merged', 'excel_second_sheet'):
        try:
            results.append(run_case(kind))
        except Exception as exc:
            results.append(dict(case=kind, passed=False, error=str(exc)))
    report = dict(scope='合成目标变体；未验证真实浏览器、真实银行全部模板或其他电脑',
                  cases=results, passed=sum(r['passed'] for r in results), total=len(results))
    report_path = ROOT / 'validation_target_results.json'
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if report['passed'] == report['total'] else 1


if __name__ == '__main__':
    sys.exit(main())
