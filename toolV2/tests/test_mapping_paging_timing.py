"""Synthetic paging through the public process boundary and honest timer accounting."""
import json
from pathlib import Path
import shutil
import sys
import tempfile
import unittest
from unittest.mock import patch

from docx import Document

TOOL = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(TOOL))
from workflow import mapping_view, timing
from test_mapping_workflow import request


class PagingTests(unittest.TestCase):
    def test_fifteen_templates_sixty_positions_read_once_then_move(self):
        with tempfile.TemporaryDirectory() as tmp:
            work = Path(tmp) / 'random-project'
            work.mkdir()
            source = Document()
            source.add_paragraph('借款人名称：合成分页公司')
            source.add_paragraph('联系电话：13800000000')
            source.save(work / 'source.docx')
            targets = []
            for i in range(15):
                doc = Document()
                for j in range(4):
                    doc.add_paragraph(f'未知含义条目{j} ____')
                name = f'target-{i:02d}.docx'
                doc.save(work / name)
                targets.append(name)
            result = request({'action': 'start', 'work': str(work), 'source': ['source.docx'],
                              'targets': targets, 'batch': '20260925-88'})
            while result['status'] == 'awaiting_source':
                result = request({'action': 'resume', 'work': str(work), 'task_id': result['task_id'],
                                  'answers': [{'id': q['id'], 'selected': [q['options'][0]['label']], 'custom': ''}
                                              for q in result['questions']]})
            self.assertEqual('needs_mapping', result['status'], result)
            self.assertLess(len(json.dumps(result, ensure_ascii=False)), 18000)
            default_chars = len(json.dumps(result, ensure_ascii=False))
            max_page_chars = 0
            self.assertEqual(1, result['counters']['source_reads'])
            phones = [f for f in result['available_fields'] if f['value'] == '13800000000']
            self.assertEqual(1, len(phones))
            self.assertIn('借款人联系电话', [phones[0]['field'], *phones[0]['aliases']])
            before = result['counters']
            updates = []
            for name in targets:
                query = {'section': 'positions', 'template': name}
                while query:
                    r = request({'action': 'read_mapping', 'work': str(work), 'task_id': result['task_id'], 'mapping_read': query})
                    self.assertEqual(before, r['counters'])
                    self.assertLess(len(json.dumps(r, ensure_ascii=False)), 7500)
                    max_page_chars = max(max_page_chars, len(json.dumps(r, ensure_ascii=False)))
                    for slot in r['mapping_page']['items']:
                        self.assertIn('【填写处】', str(slot['context']))
                        updates.append({'template': name, 'slot_id': slot['id'], 'field': '联系电话'})
                    query = r['mapping_page']['next']
            self.assertEqual(60, len(updates))
            # A new process and relocated random workspace use the same task and pages.
            moved = Path(tmp) / 'relocated'
            shutil.copytree(work, moved)
            r = request({'action': 'read_mapping', 'work': str(moved), 'task_id': result['task_id'],
                         'mapping_read': {'section': 'fields'}})
            self.assertEqual(before, r['counters'])
            self.assertTrue(any(f['value'] == '13800000000' for f in r['mapping_page']['items']))
            changed = request({'action': 'update_positions', 'work': str(moved), 'task_id': result['task_id'], 'updates': updates})
            self.assertEqual('awaiting_fill', changed['status'], changed)
            stale = request({'action': 'read_mapping', 'work': str(moved), 'task_id': result['task_id'],
                             'mapping_read': {'section': 'positions', 'template': targets[0], 'revision': r['mapping_page']['revision']}})
            self.assertFalse(stale['ok'])
            self.assertIn('已经变化', stale['error'])
            print(json.dumps({'mapping_templates': 15, 'mapping_positions': len(updates),
                              'default_response_characters': default_chars,
                              'largest_position_page_characters': max_page_chars,
                              'source_reads': before['source_reads']}, ensure_ascii=False))

    def test_large_fact_remains_completely_readable_without_serialized_payload(self):
        value = '合成长文本' * 2000
        task = {'documents': {}, 'available_fields': [
            {'field': '说明', 'value': value, 'entity': '合成公司', 'source': '来源.docx', 'fact_id': 1},
            {'field': '借款人说明', 'value': value, 'entity': '合成公司', 'source': '来源.docx', 'fact_id': 1}]}
        first = mapping_view.page(task, {'section': 'fields'})
        self.assertEqual(1, first['total'])
        query = first['items'][0]['details_read']
        recovered = ''
        while query:
            page = mapping_view.page(task, query)
            self.assertLess(mapping_view.size(page), 6500)
            recovered += ''.join(i['text'] for i in page['items'] if i.get('property') == '.value')
            query = page['next']
        self.assertEqual(value, recovered)


class TimingTests(unittest.TestCase):
    def test_cancel_replay_and_mapping_gaps_are_not_double_counted(self):
        task = {'status': 'awaiting_fill'}
        with patch.object(timing, 'now', return_value=100):
            start = timing.begin(task)
        with patch.object(timing, 'now', return_value=101):
            timing.end(task, start)
        event = {'id': 'native-wait-1', 'milliseconds': 3000, 'outcome': 'cancelled'}
        with patch.object(timing, 'now', return_value=106):
            start = timing.begin(task, event)
            timing.end(task, start)
        with patch.object(timing, 'now', return_value=108):
            start = timing.begin(task, event)
            task['status'] = 'needs_mapping'
            timing.end(task, start)
        with patch.object(timing, 'now', return_value=118):
            timing.begin(task)
        stages = timing.summary(task)['stages']
        self.assertEqual(3, stages['user_confirmation_wait']['seconds'])
        self.assertEqual(1, stages['user_confirmation_wait']['count'])
        self.assertEqual(4, stages['other_interval_unclassified']['seconds'])
        self.assertEqual(10, stages['mapping_interval_unclassified']['seconds'])
        self.assertEqual(1, timing.summary(task)['cancelled_waits'])


if __name__ == '__main__':
    unittest.main()
