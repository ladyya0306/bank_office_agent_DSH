"""Optional repeated signer blocks do not inherit an earlier borrower."""
from __future__ import annotations

from pathlib import Path

from docx import Document

from office_kit.fact_catalog import qualified_facts, target_subject_context
from office_kit.store_v2 import StoreV2
from office_kit.template_slots import discover_slots
from workflow import mapping
from workflow.runner import dispatch


BORROWER = '\u5408\u6210\u501f\u6b3e\u4eba\u4f01\u4e1a'
GUARANTOR_TWO = '\u5408\u6210\u4fdd\u8bc1\u4eba\u4e8c\u4f01\u4e1a'
BATCH = '20260926-97'


def template(path: Path) -> None:
    doc = Document()
    doc.add_paragraph('\u53ef\u6839\u636e\u7b7e\u7f72\u65b9\u4e2a\u6570\u589e\u51cf')
    doc.add_paragraph('\u4f01\u4e1a\u540d\u79f0\uff1a____')
    doc.add_paragraph('\u6cd5\u5b9a\u4ee3\u8868\u4eba\uff1a____')
    doc.add_paragraph('\u4f4f\u6240\uff1a____')
    doc.add_paragraph('\u8054\u7cfb\u4eba\uff1a____')
    doc.add_paragraph('\u5f00\u6237\u884c\u53ca\u8d26\u53f7\uff1a____')
    doc.add_paragraph('')
    doc.add_paragraph('\u7ecf\u8fc7\u5bf9____\u7b7e\u7ea6\u6838\u5b9e')
    doc.add_paragraph('\u4fdd\u8bc1\u4eba2')
    doc.add_paragraph('\u4f01\u4e1a\u540d\u79f0\uff1a____')
    doc.save(path)


def store_with_two_subjects(path: Path) -> tuple[StoreV2, Path]:
    store = StoreV2(path / 'workflow.db')
    store.batch_scope = BATCH
    borrower, _ = store.ensure_entity(BORROWER)
    guarantor, _ = store.ensure_entity(GUARANTOR_TWO)
    store.set_role(borrower, '\u501f\u6b3e\u4eba', evidence='\u501f\u6b3e\u4eba', batch_no=BATCH)
    store.set_role(guarantor, '\u4fdd\u8bc1\u4eba', evidence='\u4fdd\u8bc1\u4eba2', batch_no=BATCH)
    store.put_fact(BATCH, '\u501f\u6b3e\u4eba\u540d\u79f0', BORROWER, entity_id=borrower)
    store.put_fact(BATCH, '\u4fdd\u8bc1\u4eba\u540d\u79f0', GUARANTOR_TWO, entity_id=guarantor)
    doc = path / 'generic.docx'
    template(doc)
    return store, doc


def test_optional_block_is_unbound_but_explicit_second_subject_and_conclusion_map(tmp_path: Path):
    store, doc = store_with_two_subjects(tmp_path)
    try:
        slots = discover_slots(doc)
        assert len(slots) == 7
        facts = qualified_facts(store)
        contexts = [(slot, target_subject_context(store, doc, slot['target'])) for slot in slots]
        optional_slots = [slot for slot, context in contexts if context['additional_party']]
        conclusion_slot, conclusion = next((slot, context) for slot, context in contexts
                                           if '\u7ecf\u8fc7\u5bf9' in slot['target']['expected_text'])
        guarantor_slot, guarantor = next((slot, context) for slot, context in contexts
                                         if context['role'] == '\u4fdd\u8bc1\u4eba' and context['number'] == 2)
        assert conclusion['role'] == '\u501f\u6b3e\u4eba'
        assert conclusion['additional_party'] is False
        assert mapping.scoped_candidates(store, doc, conclusion_slot, facts)[0]['field'] == '\u501f\u6b3e\u4eba\u540d\u79f0'
        assert len(optional_slots) == 5
        for optional in optional_slots:
            scope = target_subject_context(store, doc, optional['target'])
            assert scope['role'] is None and scope['number'] is None
            assert '\u672a\u9ed8\u8ba4\u590d\u5236\u501f\u6b3e\u4eba' in scope['reason']
            assert mapping.scoped_candidates(store, doc, optional, facts) == []
        assert (guarantor['role'], guarantor['number']) == ('\u4fdd\u8bc1\u4eba', 2)
        assert mapping.scoped_candidates(store, doc, guarantor_slot, facts)[0]['field'] == '\u4fdd\u8bc1\u4eba2\u540d\u79f0'
    finally:
        store.close()


def test_optional_instruction_inside_company_label_applies_to_whole_block(tmp_path: Path):
    store, path = store_with_two_subjects(tmp_path)
    try:
        doc = Document(path)
        doc.paragraphs[0].text = ''
        doc.paragraphs[1].text = '公司名称（按签署方数量调整）：____'
        doc.save(path)
        slots = discover_slots(path)
        facts = qualified_facts(store)
        for slot in slots[:5]:
            scope = target_subject_context(store, path, slot['target'])
            assert scope['additional_party'] and scope['role'] is None, scope
            assert mapping.scoped_candidates(store, path, slot, facts) == []
        assert target_subject_context(store, path, slots[5]['target'])['role'] == '借款人'
        assert target_subject_context(store, path, slots[6]['target'])['number'] == 2
    finally:
        store.close()


def test_auto_history_is_removed_but_model_mapping_in_optional_block_is_kept(tmp_path: Path):
    store, doc = store_with_two_subjects(tmp_path)
    try:
        record: dict = {}
        slots = mapping.inspect(store, doc, record, qualified_facts(store))
        optional = next(slot for slot in slots
                        if slot['target_subject']['additional_party'])
        tid = store.register_template(doc, BATCH)
        store.add_rule(tid, '\u501f\u6b3e\u4eba\u540d\u79f0', '\u4f01\u4e1a\u540d\u79f0', optional['target'],
                       confidence=1, decided_by='auto', batch_no=BATCH)
        mapping.auto_map(store, doc, record, qualified_facts(store), BATCH)
        optional_position = mapping.physical(doc, optional['target'])
        assert not any(rule['decided_by'] == 'auto' and mapping.intersects(
            optional_position, mapping.physical(doc, rule['target']))
            for rule in store.rules_for(doc))

        target = {**optional['target'], 'slot_id': optional['id']}
        replacement = mapping.combine(store, doc, {'field': '\u4fdd\u8bc1\u4eba2\u540d\u79f0',
                                                   'slot_id': optional['id'], 'target': target})
        store.add_rule(tid, '\u4fdd\u8bc1\u4eba2\u540d\u79f0', '\u4f01\u4e1a\u540d\u79f0', replacement,
                       confidence=1, decided_by='model', batch_no=BATCH)
        mapping.auto_map(store, doc, record, qualified_facts(store), BATCH)
        assert any(rule['field'] == '\u4fdd\u8bc1\u4eba2\u540d\u79f0' and rule['decided_by'] == 'model'
                   for rule in store.rules_for(doc))
    finally:
        store.close()


def answer_all(result: dict) -> dict:
    while result['status'] in ('awaiting_source', 'awaiting_fill'):
        result = dispatch({'action': 'resume', 'work': result['work'], 'task_id': result['task_id'],
                           'answers': [{'id': item['id'], 'selected': [item['options'][0]['label']], 'custom': ''}
                                       for item in result['questions']]})
    return result


def test_dispatch_keeps_explicit_guarantor_two_mapping_in_optional_block(tmp_path: Path):
    source = Document()
    source.add_paragraph('\u501f\u6b3e\u4eba\uff1a' + BORROWER)
    source.add_paragraph('\u4fdd\u8bc1\u4eba2\uff1a' + GUARANTOR_TWO)
    source.save(tmp_path / 'source.docx')
    target = tmp_path / 'generic.docx'
    template(target)
    result = answer_all(dispatch({'action': 'start', 'work': str(tmp_path), 'source': ['source.docx'],
                                  'targets': ['generic.docx'], 'batch': BATCH}))
    assert result['status'] == 'needs_mapping', result
    optional = next(item for item in result['mapping_requests'][0]['positions']
                    if item['target_subject']['additional_party'])
    optional_slots = [item for item in result['mapping_requests'][0]['positions']
                      if item['target_subject']['additional_party']]
    assert len(optional_slots) == 5
    assert all(item['target_subject']['role'] is None and item['target_subject']['number'] is None
               and not item['candidates'] for item in optional_slots)
    assert optional['target_subject']['role'] is None
    result = dispatch({'action': 'update_positions', 'work': str(tmp_path), 'task_id': result['task_id'],
                       'updates': ([{'template': 'generic.docx', 'slot_id': optional['id'],
                                     'field': '\u4fdd\u8bc1\u4eba2\u540d\u79f0'}] +
                                   [{'template': 'generic.docx', 'slot_id': item['id'], 'leave_blank': True,
                                     'reason': '\u6765\u6e90\u672a\u63d0\u4f9b\u8be5\u53ef\u9009\u7b7e\u7ea6\u65b9\u4fe1\u606f'}
                                    for item in optional_slots if item['id'] != optional['id']])})
    assert result['status'] == 'awaiting_fill', result
    assert any(question.get('relationship_review') for question in result['questions'])
    result = answer_all(result)
    assert result['status'] == 'completed', result
    lines = [paragraph.text for paragraph in Document(result['results'][0]['output']).paragraphs]
    assert lines[1] == '\u4f01\u4e1a\u540d\u79f0\uff1a' + GUARANTOR_TWO
    assert lines[7] == '\u7ecf\u8fc7\u5bf9' + BORROWER + '\u7b7e\u7ea6\u6838\u5b9e'
    assert lines[9] == '\u4f01\u4e1a\u540d\u79f0\uff1a' + GUARANTOR_TWO
