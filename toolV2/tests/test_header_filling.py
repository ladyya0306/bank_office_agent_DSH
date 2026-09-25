"""Existing Word header parts must be scoped and preserve all unrelated bytes."""
from pathlib import Path
import sys
import tempfile
import unittest
import zipfile

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from office_kit.doc_fill import XmlEngine
from office_kit.common import OfficeKitError
from office_kit.xml_fill import prove_fidelity


class HeaderFillingTests(unittest.TestCase):
    def test_header_and_body_same_label_do_not_cross_write(self):
        with tempfile.TemporaryDirectory() as tmp:
            src, dst = Path(tmp) / 'template.docx', Path(tmp) / 'filled.docx'
            doc = Document()
            doc.add_paragraph('联系电话：________')
            doc.sections[0].header.paragraphs[0].text = '联系电话：________'
            doc.save(src)
            engine = XmlEngine(src)
            hits = engine.find_anchor_hits({'anchor': '联系电话：', 'part': 'word/header1.xml'})
            self.assertEqual(1, len(hits))
            self.assertTrue(engine.fill_span(*hits[0], '13800001234'))
            engine.save(dst)
            filled = Document(dst)
            self.assertEqual('联系电话：________', filled.paragraphs[0].text)
            self.assertIn('13800001234', filled.sections[0].header.paragraphs[0].text)
            with zipfile.ZipFile(src) as a, zipfile.ZipFile(dst) as b:
                self.assertEqual(a.namelist(), b.namelist())
                self.assertEqual(['word/header1.xml'], [n for n in a.namelist() if a.read(n) != b.read(n)])
            self.assertTrue(prove_fidelity(src, dst)['ok'])

    def test_default_anchor_only_searches_body(self):
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / 'template.docx'
            doc = Document()
            doc.add_paragraph('联系电话：________')
            doc.sections[0].header.paragraphs[0].text = '联系电话：________'
            doc.save(src)
            engine = XmlEngine(src)
            self.assertEqual(1, len(engine.find_anchor_hits({'anchor': '联系电话：'})))
            with self.assertRaises(OfficeKitError):
                engine.find_anchor_hits({'anchor': '联系电话：', 'part': 'word/header999.xml'})

    def test_body_only_template_does_not_gain_header_parts(self):
        with tempfile.TemporaryDirectory() as tmp:
            src, dst = Path(tmp) / 'template.docx', Path(tmp) / 'filled.docx'
            doc = Document()
            doc.add_paragraph('联系电话：________')
            doc.save(src)
            engine = XmlEngine(src)
            hits = engine.find_anchor_hits({'anchor': '联系电话：'})
            self.assertTrue(engine.fill_span(*hits[0], '13800001234'))
            engine.save(dst)
            with zipfile.ZipFile(src) as a, zipfile.ZipFile(dst) as b:
                self.assertEqual(a.namelist(), b.namelist())
                self.assertEqual(['word/document.xml'], [n for n in a.namelist() if a.read(n) != b.read(n)])


if __name__ == '__main__':
    unittest.main()
