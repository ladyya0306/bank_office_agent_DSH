"""Synthetic Word regressions for source ownership and conflict review."""
from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import time
import unittest
from pathlib import Path

from docx import Document

TOOL = Path(__file__).resolve().parents[1]
if str(TOOL) not in sys.path:
    sys.path.insert(0, str(TOOL))
from office_kit.absorb import absorb  # noqa: E402


def write_docx(path: Path, *paragraphs: str) -> None:
    doc = Document()
    for paragraph in paragraphs:
        doc.add_paragraph(paragraph)
    doc.save(path)


def request(payload: dict) -> dict:
    completed = subprocess.run(
        [sys.executable, str(TOOL / "office.py")],
        input=json.dumps(payload, ensure_ascii=False), text=True, encoding="utf-8",
        capture_output=True, timeout=60,
    )
    if completed.returncode:
        raise AssertionError(f"office.py failed: {completed.stderr}\n{completed.stdout}")
    return json.loads(completed.stdout)


class SourceRoleTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.work = Path(self.tmp.name) / "synthetic-work"
        self.work.mkdir()
        self.source = self.work / "来源角色样例.docx"
        self.target = self.work / "目标.docx"

    def tearDown(self) -> None:
        for _ in range(20):
            try:
                self.tmp.cleanup()
                return
            except PermissionError:
                time.sleep(0.25)
        self.tmp.cleanup()

    def payload(self, batch: str = "20260925-01") -> dict:
        return {"action": "start", "work": str(self.work),
                "source": [str(self.source)], "targets": [str(self.target)],
                "batch": batch}

    def drive_to_conflict(self, result: dict) -> tuple[dict, dict]:
        """Answer ownership cards first, then return the native conflict card."""
        for _ in range(3):
            conflict = next((q for q in result["questions"]
                             if q["id"].startswith("source-conflict-")
                             and "13800000000" in q["question"]
                             and ("13900000000" in q["question"] or
                                  "15800000000" in q["question"])), None)
            if conflict is not None:
                return result, conflict
            self.assertEqual("awaiting_source", result["status"], result)
            result = request({"action": "resume", "work": str(self.work),
                              "task_id": result["task_id"],
                              "answers": [{"id": q["id"],
                                           "selected": [q["options"][0]["label"]],
                                           "custom": ""}
                                          for q in result["questions"]]})
        self.fail(f"三轮内未出现冲突问题：{result}")

    def test_guarantor_company_and_person_keep_legal_identity_documents_with_owner(self) -> None:
        write_docx(
            self.source,
            "借款人：合成借款公司",
            "统一社会信用代码：BORROWER-CREDIT",
            "法定代表人：借款法人",
            "身份证号码：BORROWER-ID",
            "保证人1：合成保证公司",
            "法定代表人：保证法人",
            "统一社会信用代码：GUARANTOR-CREDIT",
            "身份证号码：GUARANTOR-ID",
            "保证人2：合成自然人",
            "身份证号码：PERSON-ID",
        )
        rows = absorb(self.source)
        def row_for(key: str, value: str) -> dict:
            return next(row for row in rows if row["key"] == key and row["value"] == value)

        self.assertEqual("合成保证公司", row_for("统一社会信用代码", "GUARANTOR-CREDIT")["entity_name"])
        self.assertEqual("保证法人", row_for("证件号码", "GUARANTOR-ID")["entity_name"])
        self.assertEqual("合成自然人", row_for("证件号码", "PERSON-ID")["entity_name"])
        self.assertNotEqual("合成借款公司", row_for("统一社会信用代码", "GUARANTOR-CREDIT")["entity_name"])
        self.assertNotEqual("合成借款公司", row_for("证件号码", "GUARANTOR-ID")["entity_name"])
        self.assertNotEqual("合成借款公司", row_for("证件号码", "PERSON-ID")["entity_name"])

    def test_taiwan_pass_and_company_passport_pairs_keep_type_and_number_owner(self) -> None:
        write_docx(
            self.source,
            "保证人1：合成台籍自然人",
            "台湾居民来往大陆通行证：TAIWAN-PASS-001",
            "证件号码：GENERIC-001",
            "保证人2：合成保证公司",
            "法定代表人：合成公司法人",
            "护照：PASSPORT-001",
        )
        rows = absorb(self.source)

        taiwan_number = next(row for row in rows if row["value"] == "TAIWAN-PASS-001"
                             and row["key"] == "证件号码")
        taiwan_type = next(row for row in rows if row["value"] == "台湾居民来往大陆通行证"
                           and row["key"] == "证件类型")
        passport_number = next(row for row in rows if row["value"] == "PASSPORT-001"
                               and row["key"] == "证件号码")
        passport_type = next(row for row in rows if row["value"] == "护照"
                             and row["key"] == "证件类型")
        self.assertEqual("合成台籍自然人", taiwan_number["entity_name"])
        self.assertEqual("合成台籍自然人", taiwan_type["entity_name"])
        self.assertEqual("合成公司法人", passport_number["entity_name"])
        self.assertEqual("合成公司法人", passport_type["entity_name"])
        self.assertEqual(taiwan_number["ownership_group"], taiwan_type["ownership_group"])
        self.assertEqual(passport_number["ownership_group"], passport_type["ownership_group"])
        generic = next(row for row in rows if row["value"] == "GENERIC-001")
        self.assertEqual("证件号码", generic["key"])
        self.assertIsNone(generic.get("certificate_type"))

        write_docx(self.target, "证件信息：")
        started = request(self.payload())
        self.assertEqual("awaiting_source", started["status"], started)
        pair_questions = [q for q in started["questions"]
                          if "TAIWAN-PASS-001" in q["question"] or
                          "PASSPORT-001" in q["question"]]
        self.assertEqual(2, len(pair_questions), started["questions"])
        self.assertEqual(1, sum("TAIWAN-PASS-001" in q["question"] for q in pair_questions))
        self.assertEqual(1, sum("PASSPORT-001" in q["question"] for q in pair_questions))
        self.assertTrue(all(self.source.name in q["question"] for q in pair_questions))
        generic_question = next(q for q in started["questions"] if "GENERIC-001" in q["question"])
        self.assertNotIn("居民身份证", generic_question["question"])

    def test_source_question_shows_file_original_quote_and_context_basis(self) -> None:
        write_docx(self.source, "联系电话：13800000000")
        write_docx(self.target, "联系电话：")
        started = request(self.payload())
        self.assertEqual("awaiting_source", started["status"], started)
        self.assertEqual(1, len(started["questions"]))
        question = started["questions"][0]["question"]
        self.assertIn(self.source.name, question)
        self.assertIn("联系电话：13800000000", question)
        self.assertTrue(any(marker in question for marker in ("上下文", "主体", "归属")), question)

    def test_same_subject_different_values_becomes_native_conflict_review_and_recovers(self) -> None:
        write_docx(self.source, "借款人：合成借款公司",
                   "联系电话：13800000000", "联系电话：13900000000")
        write_docx(self.target, "联系电话：")
        started, conflict = self.drive_to_conflict(request(self.payload()))
        self.assertIn(self.source.name, conflict["question"])
        self.assertTrue(conflict["options"], conflict)

        resumed = request({"action": "resume", "work": str(self.work),
                           "task_id": started["task_id"],
                           "answers": [{"id": conflict["id"],
                                        "selected": ["采用第 1 项"],
                                        "custom": ""}]})
        self.assertNotEqual("failed", resumed["status"], resumed)
        if resumed["status"] == "awaiting_fill":
            resumed = request({"action": "resume", "work": str(self.work),
                               "task_id": resumed["task_id"],
                               "answers": [{"id": q["id"],
                                            "selected": [q["options"][0]["label"]],
                                            "custom": ""} for q in resumed["questions"]]})
        self.assertEqual("completed", resumed["status"], resumed)
        output = Path(resumed["results"][0]["output"])
        text = "\n".join(p.text for p in Document(output).paragraphs)
        self.assertIn("13800000000", text)
        self.assertNotIn("13900000000", text)
        reread = request({"action": "status", "work": str(self.work),
                          "task_id": started["task_id"]})
        self.assertEqual("completed", reread["status"], reread)
        self.assertFalse(reread["questions"])

    def test_conflict_can_be_left_blank_and_remains_resolved_after_new_process(self) -> None:
        write_docx(self.source, "借款人：合成借款公司",
                   "联系电话：13800000000", "联系电话：13900000000")
        write_docx(self.target, "联系电话：")
        started, conflict = self.drive_to_conflict(request(self.payload()))
        resumed = request({"action": "resume", "work": str(self.work),
                           "task_id": started["task_id"],
                           "answers": [{"id": conflict["id"],
                                        "selected": ["这些值暂不采用，留空"],
                                        "custom": ""}]})
        self.assertNotEqual("failed", resumed["status"], resumed)
        reread = request({"action": "status", "work": str(self.work),
                          "task_id": started["task_id"]})
        self.assertNotEqual("failed", reread["status"], reread)
        self.assertFalse(any(q["id"].startswith("source-conflict-")
                             for q in reread["questions"]), reread["questions"])

    def test_changed_conflict_value_reopens_review_without_reusing_old_answer(self) -> None:
        write_docx(self.source, "借款人：合成借款公司",
                   "联系电话：13800000000", "联系电话：13900000000")
        write_docx(self.target, "联系电话：")
        started, conflict = self.drive_to_conflict(request(self.payload()))
        resumed = request({"action": "resume", "work": str(self.work),
                           "task_id": started["task_id"],
                           "answers": [{"id": conflict["id"],
                                        "selected": ["采用第 1 项"],
                                        "custom": ""}]})
        self.assertNotEqual("failed", resumed["status"], resumed)

        write_docx(self.source, "借款人：合成借款公司",
                   "联系电话：13800000000", "联系电话：15800000000")
        changed, changed_conflict = self.drive_to_conflict(request(self.payload()))
        self.assertTrue(any("13800000000" in q["question"] and "15800000000" in q["question"]
                            for q in changed["questions"]), changed["questions"])


if __name__ == "__main__":
    unittest.main()
