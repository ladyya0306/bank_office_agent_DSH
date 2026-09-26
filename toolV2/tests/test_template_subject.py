from pathlib import Path

from docx import Document

from office_kit.fact_catalog import target_subject_context
from office_kit.harness import _guard_fill, _template_subject_scope
from office_kit.store_v2 import StoreV2


def _store(tmp_path):
    store = StoreV2(tmp_path / "workflow.db")
    # Deliberately use the legacy/default entity type for all four entities.
    # Company filtering must rely on source facts, not that default.
    borrower, _ = store.ensure_entity("合成借款主体")
    personal, _ = store.ensure_entity("合成个人主体")
    company_one, _ = store.ensure_entity("合成保证主体一")
    company_two, _ = store.ensure_entity("合成保证主体二")
    store.set_role(borrower, "借款人", batch_no="20260925-01")
    for eid in (personal, company_one, company_two):
        store.set_role(eid, "保证人", batch_no="20260925-01")
    for eid, value in ((borrower, "借款地址"), (personal, "个人地址"),
                       (company_one, "保证公司一地址"), (company_two, "保证公司二地址")):
        store.put_fact("20260925-01", "地址", value, entity_id=eid)
    store.put_fact("20260925-01", "借款人法定代表人", "借款公司代表", entity_id=borrower)
    store.put_fact("20260925-01", "保证人法定代表人", "保证公司一代表", entity_id=company_one)
    store.put_fact("20260925-01", "保证人法定代表人", "保证公司二代表", entity_id=company_two)
    store.put_fact("20260925-01", "联系电话", "个人保证人电话", entity_id=personal)
    store.put_fact("20260925-01", "联系电话", "保证公司一电话", entity_id=company_one)
    store.put_fact("20260925-01", "联系电话", "保证公司二电话", entity_id=company_two)
    store.batch_scope = "20260925-01"
    return store, borrower, personal, company_one, company_two


def test_company_template_filters_personal_guarantor_and_never_uses_first(tmp_path):
    store, _borrower, personal, company_one, company_two = _store(tmp_path)
    try:
        eid, scope, label, note = _template_subject_scope(
            store, Path("保证人法定代表人身份证明.docx"), [])
        assert eid is None
        assert label == "保证人"
        assert scope == [company_one, company_two]
        assert personal not in scope
        assert "企业" in note

        facts = store.facts_for_subject(eid, scope_entity_ids=scope)
        assert facts["地址"]["_ambiguous"] is True
        assert {c["entity_id"] for c in facts["地址"]["_candidates"]} == {company_one, company_two}
        assert {c["entity_id"] for c in facts["联系电话"]["_candidates"]} == {company_one, company_two}
    finally:
        store.close()


def test_borrower_scope_does_not_borrow_a_single_guarantor_phone(tmp_path):
    store, borrower, _personal, _company_one, _company_two = _store(tmp_path)
    try:
        eid, scope, _label, _note = _template_subject_scope(
            store, Path("借款人法定代表人身份证明.docx"), [])
        assert eid == borrower
        facts = store.facts_for_subject(eid, scope_entity_ids=scope)
        assert facts["地址"]["entity_id"] == borrower
        assert "联系电话" not in facts
    finally:
        store.close()


def test_company_template_without_evidence_keeps_candidate_for_explicit_choice(tmp_path):
    store = StoreV2(tmp_path / "empty-company-evidence.db")
    try:
        guarantor, _ = store.ensure_entity("默认类型保证主体")
        store.set_role(guarantor, "保证人", batch_no="20260925-01")
        store.put_fact("20260925-01", "联系电话", "唯一但无关的电话", entity_id=guarantor)
        store.batch_scope = "20260925-01"
        eid, scope, label, note = _template_subject_scope(
            store, Path("保证人法定代表人身份证明.docx"), [])
        assert (eid, scope, label) == (None, [guarantor], "保证人")
        assert "未明确企业身份" in note
        phone = store.facts_for_subject(eid, scope_entity_ids=scope)["联系电话"]
        assert phone["_ambiguous"] is True
        assert [c["entity_id"] for c in phone["_candidates"]] == [guarantor]
    finally:
        store.close()


def test_single_rule_role_is_fallback_but_multiple_declarations_do_not_choose_borrower(tmp_path):
    store, borrower, _personal, _company_one, _company_two = _store(tmp_path)
    try:
        eid, scope, label, _note = _template_subject_scope(
            store, Path("未命名材料.docx"), [{"field": "借款人名称"}])
        assert (eid, scope, label) == (borrower, [borrower], "借款人")

        eid, scope, label, note = _template_subject_scope(
            store, Path("借款人及保证人材料.docx"), [{"field": "借款人名称"}])
        assert (eid, scope, label) == (None, None, "认不出")
        assert "同时出现" in note
    finally:
        store.close()


def test_ambiguous_custom_value_is_a_local_override_not_a_source_entity():
    row = {"n": 1, "kind": "slot", "field": "联系电话", "decision": "ask",
           "ambiguous": True, "candidates": [{"entity_id": 7, "entity_name": "候选甲", "value": "旧值"}],
           "entity_id": None, "subject_eid": 42, "subject_label": "借款人"}
    accepted, problems = _guard_fill({"rows": [row]}, {
        "chosen": [1], "apply_all": False, "actions": {1: {"action": "new", "value": "新值"}}})
    assert problems == []
    assert accepted[0]["value"] == "新值"
    assert accepted[0]["entity_id"] == 42
    assert accepted[0]["_target_override"] is True

    row["subject_eid"] = None
    accepted, problems = _guard_fill({"rows": [row]}, {
        "chosen": [1], "apply_all": False, "actions": {1: {"action": "new", "value": "新值"}}})
    assert problems == []
    assert accepted[0]["entity_id"] is None


def test_same_role_scope_does_not_borrow_a_field_from_outside_the_scope(tmp_path):
    store, _borrower, personal, company_one, company_two = _store(tmp_path)
    try:
        store.put_fact("20260925-01", "保证人联系电话", "仅个人保证人的电话", entity_id=personal)
        facts = store.facts_for_subject(None, scope_entity_ids=[company_one, company_two])
        assert "保证人联系电话" not in facts
    finally:
        store.close()


def test_mixed_role_declaration_leaves_candidates_visible_for_a_question(tmp_path):
    store, _borrower, _personal, _company_one, _company_two = _store(tmp_path)
    try:
        eid, scope, label, _note = _template_subject_scope(
            store, Path("借款人及保证人材料.docx"), [])
        assert (eid, scope, label) == (None, None, "认不出")
        phone = store.facts_for_subject(eid, scope_entity_ids=scope)["联系电话"]
        assert phone["_ambiguous"] is True
        assert len(phone["_candidates"]) == 3
    finally:
        store.close()


def test_table_position_uses_numbered_guarantor_evidence_not_role_order(tmp_path):
    store, borrower, _personal, company_one, company_two = _store(tmp_path)
    try:
        store.set_role(company_one, "保证人", evidence="保证人1：合成保证主体一", batch_no="20260925-01")
        store.set_role(company_two, "保证人", evidence="保证人2：合成保证主体二", batch_no="20260925-01")
        template = tmp_path / "核保书.docx"
        doc = Document(); table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "担保人2"
        table.cell(1, 0).text = "地址"
        doc.save(template)
        context = target_subject_context(store, template,
                                         {"kind": "cell", "table": 0, "row": 1, "col": 1},
                                         label="地址")
        assert context["role"] == "保证人"
        assert context["number"] == 2
        assert context["entity_ids"] == [company_two]
        assert borrower not in context["entity_ids"]
    finally:
        store.close()


def test_anonymous_repeated_position_does_not_default_to_unique_borrower(tmp_path):
    store, _borrower, _personal, _company_one, _company_two = _store(tmp_path)
    try:
        template = tmp_path / "甲表.docx"
        doc = Document(); doc.add_paragraph("法定代表人：____"); doc.add_paragraph("法定代表人：____")
        doc.save(template)
        context = target_subject_context(store, template,
                                         {"kind": "anchor", "part": "word/document.xml",
                                          "paragraph_index": 1,
                                          "expected_text": "法定代表人：____", "span_start": 6, "span_end": 10})
        assert context["role"] is None
        assert context["entity_ids"] == []
    finally:
        store.close()


def test_generic_guarantor_representative_inherits_nearest_number_but_own_hint(tmp_path):
    store, _borrower, _personal, _company_one, company_two = _store(tmp_path)
    try:
        store.set_role(company_two, "保证人", evidence="保证人2：合成保证主体二", batch_no="20260925-01")
        template = tmp_path / "核保书.docx"
        doc = Document(); doc.add_paragraph("担保人2："); doc.add_paragraph("担保人法定代表人或授权代理人身份证号：")
        doc.save(template)
        context = target_subject_context(store, template,
                                         {"kind": "anchor", "part": "word/document.xml",
                                          "paragraph_index": 1,
                                          "expected_text": "担保人法定代表人或授权代理人身份证号：",
                                          "span_start": len("担保人法定代表人或授权代理人身份证号："),
                                          "span_end": len("担保人法定代表人或授权代理人身份证号：")})
        assert (context["role"], context["number"], context["field_hint"]) == ("保证人", 2, "法定代表人证件号码")
        assert context["entity_ids"] == [company_two]
    finally:
        store.close()
