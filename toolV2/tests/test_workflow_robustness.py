"""Public-workflow robustness probes using only synthetic documents."""
from __future__ import annotations

import hashlib
import tempfile
import unittest
from unittest.mock import patch
from pathlib import Path

from docx import Document
from openpyxl import Workbook, load_workbook

from workflow.runner import dispatch
from workflow import storage as state
from office_kit import harness
from office_kit.store_v2 import StoreV2
from workflow import runner


def word(path: Path, *lines: str) -> None:
    doc = Document()
    for line in lines:
        doc.add_paragraph(line)
    doc.save(path)


class WorkflowRobustnessTests(unittest.TestCase):
    def test_execution_summary_recovers_existing_run_as_reused(self):
        before = {'documents': {'target.xlsx': {
            'status': 'running', 'run_id': 'prior-successful-run',
        }}}
        task = {'documents': {'target.xlsx': {
            'status': 'completed', 'run_id': 'prior-successful-run',
            'output': 'out/recovered.xlsx', 'output_hash': 'saved-hash',
        }}}
        self.assertEqual({'generated_files': 0, 'reused_files': 1, 'generated_templates': []},
                         runner.execution_summary(before, task))

    def test_disbursement_amount_uses_current_loan_not_credit_limit(self):
        with tempfile.TemporaryDirectory() as tmp:
            work = Path(tmp)
            word(work/'source.docx', '借款人：合成公司', '授信额度：900万元', '本次借款金额：325万元')
            book = Workbook(); book.active['A1'] = '放款金额（万元）：____'
            book.save(work/'target.xlsx')
            result = self.start(work, 'source.docx', ['target.xlsx'], '20260926-99')
            self.assertEqual('completed', result['status'], result)
            self.assertEqual({'generated_files': 1, 'reused_files': 0},
                             {key: result['execution_summary'][key] for key in ('generated_files', 'reused_files')})
            out = load_workbook(result['results'][0]['output'])
            self.assertEqual('放款金额（万元）：325', out.active['A1'].value)

    def start(self, work: Path, source: str, targets: list[str], batch: str) -> dict:
        return dispatch({'action': 'start', 'work': str(work), 'source': [source],
                         'targets': targets, 'batch': batch})

    def test_numbered_guarantors_follow_explicit_numbers_not_entity_order(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            work = Path(tmp)
            word(work/'source.docx', '保证人2：合成乙公司', '联系电话：222', '保证金额：200万元',
                 '保证人1：合成甲公司', '联系电话：111', '保证金额：100万元')
            word(work/'target.docx', '保证人1名称：____', '保证人2名称：____')
            result = self.start(work, 'source.docx', ['target.docx'], '20260926-93')
            self.assertEqual('completed', result['status'], result)
            text = '\n'.join(p.text for p in Document(result['results'][0]['output']).paragraphs)
            self.assertIn('保证人1名称：合成甲公司', text)
            self.assertIn('保证人2名称：合成乙公司', text)

    def test_labeled_xlsx_values_keep_formula_and_sheet_name(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            work = Path(tmp)
            word(work/'source.docx', '借款人：合成公司', '授信额度：900万元',
                 '本次借款金额：325万元', '已使用额度：40万元')
            book = Workbook(); sheet = book.active; sheet.title = '自定义页'
            for row, label in enumerate(('授信额度', '本次借款金额', '已使用额度'), 1):
                sheet.cell(row, 1, label); sheet.cell(row, 2, '____')
            sheet['D1'] = '=SUM(1,1)'; book.save(work/'target.xlsx'); book.close()
            result = self.start(work, 'source.docx', ['target.xlsx'], '20260926-94')
            self.assertEqual('completed', result['status'], result)
            out = load_workbook(result['results'][0]['output'], data_only=False)
            try:
                sheet = out['自定义页']
                self.assertEqual('900万元', sheet['B1'].value)
                self.assertEqual('325万元', sheet['B2'].value)
                self.assertEqual('40万元', sheet['B3'].value)
                self.assertEqual('=SUM(1,1)', sheet['D1'].value)
            finally:
                out.close()

    def test_formula_workbook_repeated_status_reuses_output(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            work = Path(tmp)
            word(work/'source.docx', '借款人：合成公司', '联系电话：13800000000')
            book = Workbook(); sheet = book.active; sheet.title = '另一页'
            sheet['A1'], sheet['B1'], sheet['D1'] = '联系电话', '____', '=1+1'
            book.save(work/'target.xlsx'); book.close()
            first = self.start(work, 'source.docx', ['target.xlsx'], '20260926-95')
            self.assertEqual('completed', first['status'], first)
            again = dispatch({'action': 'status', 'work': str(work), 'task_id': first['task_id']})
            self.assertEqual(first['counters'], again['counters'])
            self.assertEqual(first['results'][0]['output'], again['results'][0]['output'])
            self.assertEqual({'generated_files': 0, 'reused_files': 1},
                             {key: again['execution_summary'][key] for key in ('generated_files', 'reused_files')})

    def test_one_changed_target_preserves_other_output_hash(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            work = Path(tmp)
            word(work/'source.docx', '借款人：合成公司', '联系电话：13800000000')
            word(work/'one.docx', '联系电话：____')
            book = Workbook(); book.active['A1'] = '联系电话：____'; book.save(work/'two.xlsx')
            targets = ['one.docx', 'two.xlsx']
            first = self.start(work, 'source.docx', targets, '20260926-96')
            self.assertEqual('completed', first['status'], first)
            self.assertEqual({'generated_files': 2, 'reused_files': 0},
                             {key: first['execution_summary'][key] for key in ('generated_files', 'reused_files')})
            before = {r['template']: r['output'] for r in first['results']}
            original = hashlib.sha256(Path(before['one.docx']).read_bytes()).hexdigest()
            book = load_workbook(work/'two.xlsx'); book.active['A2'] = '附注：第二版'; book.save(work/'two.xlsx')
            second = self.start(work, 'source.docx', targets, '20260926-96')
            self.assertEqual('completed', second['status'], second)
            self.assertEqual({'generated_files': 1, 'reused_files': 1},
                             {key: second['execution_summary'][key] for key in ('generated_files', 'reused_files')})
            after = {r['template']: r['output'] for r in second['results']}
            self.assertEqual(before['one.docx'], after['one.docx'])
            self.assertEqual(original, hashlib.sha256(Path(after['one.docx']).read_bytes()).hexdigest())
            self.assertEqual(1, second['counters']['fill_processes'] - first['counters']['fill_processes'])

    def test_missing_data_blank_reopens_after_subject_plan_upgrade_but_user_blank_does_not(self) -> None:
        for reason, expected in [('来源未提供放款金额', 'completed'), ('来源只有借款金额，未给出放款金额', 'completed'), ('用户要求不填', 'completed')]:
            with self.subTest(reason=reason), tempfile.TemporaryDirectory() as tmp:
                work = Path(tmp)
                word(work/'source.docx', '借款人：合成公司', '授信额度：900万元', '本次借款金额：325万元')
                book = Workbook(); book.active['A1'] = '放款金额（万元）：____'; book.save(work/'target.xlsx')
                original = list(harness.SYNONYMS['借款金额'])
                with patch.dict(harness.SYNONYMS, {'借款金额': [x for x in original if x != '放款金额']}):
                    first = self.start(work, 'source.docx', ['target.xlsx'], '20260926-92')
                slot = first['mapping_requests'][0]['positions'][0]
                blank = dispatch({'action': 'update_positions', 'work': str(work), 'task_id': first['task_id'],
                                  'updates': [{'template': 'target.xlsx', 'slot_id': slot['id'],
                                               'leave_blank': True, 'reason': reason}]})
                self.assertEqual('completed', blank['status'], blank)
                with patch.object(runner, 'SUBJECT_PLAN_VERSION', runner.SUBJECT_PLAN_VERSION + 1):
                    after = dispatch({'action': 'status', 'work': str(work), 'task_id': first['task_id']})
                self.assertEqual(expected, after['status'], after)
                self.assertEqual(1, after['counters']['source_reads'])
                if reason.startswith('来源'):
                    text = load_workbook(after['results'][0]['output']).active['A1'].value
                    self.assertEqual('放款金额（万元）：325', text)
                    self.assertEqual([], after['questions'])
                else:
                    text = load_workbook(after['results'][0]['output']).active['A1'].value
                    self.assertNotIn('325', str(text))
                    with StoreV2(work/'db/workflow.db') as store:
                        task = state.get(store.conn, 'office_v2_task', first['task_id'])
                    self.assertIn(slot['id'], task['documents']['target.xlsx']['blank_slots'])


if __name__ == '__main__':
    unittest.main()
