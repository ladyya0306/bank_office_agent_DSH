"""A confirmed value can fill several explicit locations without duplicate questions."""
from pathlib import Path
import tempfile
import unittest
from docx import Document
from openpyxl import Workbook, load_workbook
from test_source_diagnostics_workflow import request


class MultipleLocationTests(unittest.TestCase):
    def test_two_worksheets_and_word_cells_share_one_fill_question(self):
        with tempfile.TemporaryDirectory() as tmp:
            work = Path(tmp)
            source = work / '来源.docx'
            doc = Document()
            doc.add_paragraph('借款人：合成甲公司')
            doc.add_paragraph('联系电话：13800001234')
            doc.save(source)
            word = work / '两处.docx'
            doc = Document()
            table = doc.add_table(rows=2, cols=2)
            table.cell(0, 0).text = table.cell(1, 0).text = '联系电话'
            doc.save(word)
            excel = work / '两张表.xlsx'
            book = Workbook()
            book.active['A1'] = '联系电话：'
            book.create_sheet('第二张')['A1'] = '联系电话：'
            book.save(excel)
            book.close()
            payload = dict(action='start', work=str(work), source=[str(source)], targets=[str(word), str(excel)], batch='20260925-75')
            r = request(payload)
            fill_questions = 0
            for _ in range(4):
                if r.get('status') not in ('awaiting_source', 'awaiting_fill'):
                    break
                if r['status'] == 'awaiting_fill':
                    fill_questions += len(r['questions'])
                    rendered = r['questions'][0]['question'] + '\n' + r['questions'][0].get('detail', '')
                    self.assertIn('第二张', rendered)
                    self.assertIn('第 2 行', rendered)
                r = request(dict(action='resume', work=str(work), task_id=r['task_id'],
                                 answers=[dict(id=q['id'], selected=[q['options'][0]['label']], custom='')
                                          for q in r['questions']]))
            self.assertEqual('completed', r['status'], r)
            self.assertEqual(1, fill_questions)
            for row in r['results']:
                if row['template'].endswith('.docx'):
                    doc = Document(row['output'])
                    self.assertEqual(['13800001234', '13800001234'],
                                     [doc.tables[0].cell(i, 1).text for i in range(2)])
                else:
                    book = load_workbook(row['output'])
                    self.assertEqual(['联系电话：13800001234'] * 2, [ws['A1'].value for ws in book])
                    book.close()
            repeated = request(payload)
            self.assertEqual('completed', repeated['status'], repeated)
            self.assertFalse(repeated['questions'])
            self.assertEqual(r['counters']['fill_processes'], repeated['counters']['fill_processes'])


if __name__ == '__main__':
    unittest.main()
