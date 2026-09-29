"""Public-protocol regressions for user-confirmed learning methods.

Every fixture here is synthetic.  The tests intentionally use ``office.py`` in
separate processes so confirmation, SQLite persistence, source hashes, and
reusable target rules follow the same boundary as DSH calls.
"""
from __future__ import annotations

import json
import sqlite3
import subprocess
import sys
import tempfile
import time
import unittest
from hashlib import sha256
from pathlib import Path

from docx import Document
from openpyxl import Workbook, load_workbook


TOOL = Path(__file__).resolve().parents[1]


def request(payload: dict) -> dict:
    completed = subprocess.run([sys.executable, str(TOOL / "office.py")],
                               input=json.dumps(payload, ensure_ascii=False), text=True,
                               encoding="utf-8", capture_output=True, timeout=45)
    if completed.returncode:
        raise AssertionError("office.py failed: %s\n%s" % (completed.stderr, completed.stdout))
    return json.loads(completed.stdout)


def sha256_file(path: Path) -> str:
    return sha256(path.read_bytes()).hexdigest()


class LearningWorkflowTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.work = Path(self.tmp.name) / "synthetic-learning-work"
        self.work.mkdir()
        self.source = self.work / "narrative.docx"
        document = Document()
        document.add_paragraph("本次资金用于流动资金。")
        document.save(self.source)
        self.target = self.work / "target.xlsx"
        book = Workbook()
        book.active["A1"] = "用途："
        book.save(self.target)

    def tearDown(self) -> None:
        for attempt in range(10):
            try:
                self.tmp.cleanup()
                return
            except PermissionError:
                if attempt == 9:
                    raise
                time.sleep(0.2)

    def start_failed(self, batch: str = "20260929-01") -> dict:
        result = request({"action": "start", "work": str(self.work),
                          "source": [self.source.name], "targets": [self.target.name], "batch": batch})
        self.assertFalse(result["ok"])
        self.assertEqual("failed", result["status"])
        self.assertEqual("source", result["failed_stage"])
        return result

    def start_awaiting_source(self, batch: str = "20260929-03") -> dict:
        document = Document()
        document.add_paragraph("用途：流动资金")
        document.save(self.source)
        result = request({"action": "start", "work": str(self.work),
                          "source": [self.source.name], "targets": [self.target.name], "batch": batch})
        self.assertTrue(result["ok"], result)
        self.assertEqual("awaiting_source", result["status"])
        self.assertTrue(result["questions"])
        return result

    def updates(self, value: str = "流动资金") -> list[dict]:
        return [{"source": self.source.name, "source_sha256": sha256_file(self.source), "rows": [{
            "key": "用途", "value": value, "evidence": {"paragraph_index": 0},
        }]}]

    def propose(self, task_id: str, *, name: str = "合成用途识别", script_path: str | None = None,
                save_target_rules: bool = False, source_updates: list[dict] | None = None) -> dict:
        learning = {"action": "propose", "name": name, "source_updates": source_updates or self.updates(),
                    "save_target_rules": save_target_rules}
        if script_path:
            learning["script_path"] = script_path
        return request({"action": "learning", "work": str(self.work), "task_id": task_id,
                        "learning": learning})

    def confirm(self, proposed: dict, *, selected: str = "确认保存") -> dict:
        self.assertEqual("awaiting_method", proposed["status"])
        self.assertTrue(proposed["questions"])
        return request({"action": "resume", "work": str(self.work), "task_id": proposed["task_id"],
                        "answers": [{"id": question["id"], "selected": [selected], "custom": ""}
                                    for question in proposed["questions"]]})

    def method_list(self, task_id: str) -> list[dict]:
        result = request({"action": "learning", "work": str(self.work), "task_id": task_id,
                          "learning": {"action": "list"}})
        self.assertTrue(result["ok"], result)
        return result["learning"]["methods"]

    def fact_count(self) -> int:
        connection = sqlite3.connect(self.work / "db" / "workflow.db")
        try:
            return connection.execute("SELECT COUNT(*) FROM fact").fetchone()[0]
        finally:
            connection.close()

    def write_legacy_not_saved_task(self, task_id: str) -> None:
        connection = sqlite3.connect(self.work / "db" / "workflow.db")
        try:
            encoded = connection.execute("SELECT payload FROM office_v2_task WHERE id=?", (task_id,)).fetchone()[0]
            task = json.loads(encoded)
            task.update(status="cancelled", questions=[], learning={"status": "not_saved"})
            connection.execute("UPDATE office_v2_task SET payload=? WHERE id=?",
                               (json.dumps(task, ensure_ascii=False), task_id))
            connection.commit()
        finally:
            connection.close()

    @staticmethod
    def run_recipe(script: Path, source: Path) -> str:
        completed = subprocess.run([sys.executable, str(script), str(source)], text=True,
                                   encoding="utf-8", capture_output=True, timeout=20)
        if completed.returncode:
            raise AssertionError("synthetic recipe failed: %s" % completed.stderr)
        return completed.stdout.strip()

    def test_failed_source_still_exposes_source_and_target_document_pages(self) -> None:
        failed = self.start_failed()
        source_page = request({"action": "read_mapping", "work": str(self.work), "task_id": failed["task_id"],
                               "mapping_read": {"section": "source_document", "source": self.source.name}})
        self.assertTrue(source_page["ok"], source_page)
        self.assertEqual("本次资金用于流动资金。", source_page["mapping_page"]["items"][0]["text"])
        target_page = request({"action": "read_mapping", "work": str(self.work), "task_id": failed["task_id"],
                               "mapping_read": {"section": "document", "template": self.target.name}})
        self.assertTrue(target_page["ok"], target_page)
        self.assertEqual("用途：", target_page["mapping_page"]["items"][0]["text"])

    def test_confirmation_persists_evidenced_values_fills_xlsx_and_reuses_script_method(self) -> None:
        failed = self.start_failed()
        script = self.work / "recognise.py"
        script.write_text(
            "from docx import Document\nimport re\nimport sys\n"
            "text = '\\n'.join(p.text for p in Document(sys.argv[1]).paragraphs)\n"
            "print(re.search(r'用于(.+?)。', text).group(1))\n", encoding="utf-8")
        self.assertEqual("流动资金", self.run_recipe(script, self.source))
        proposed = self.propose(failed["task_id"], script_path=script.name)
        self.assertEqual(0, self.fact_count())

        confirmed = self.confirm(proposed)
        self.assertTrue(confirmed["ok"], confirmed)
        self.assertEqual("completed", confirmed["status"], confirmed)
        output = Path(confirmed["results"][0]["output"])
        reopened = load_workbook(output, data_only=True)
        try:
            self.assertIn("流动资金", "\n".join(str(cell.value or "")
                                                 for row in reopened.active.iter_rows() for cell in row))
        finally:
            reopened.close()
        self.assertEqual(1, self.fact_count())

        methods = self.method_list(failed["task_id"])
        self.assertEqual(1, len(methods))
        method_id = methods[0]["id"]
        read = request({"action": "learning", "work": str(self.work), "task_id": failed["task_id"],
                        "learning": {"action": "read", "id": method_id}})
        self.assertTrue(Path(read["learning"]["script_path"]).is_file())

        repeated = self.propose(failed["task_id"], script_path=script.name)
        self.assertNotEqual("awaiting_method", repeated["status"], repeated)
        self.assertEqual(1, len(self.method_list(failed["task_id"])))
        self.assertEqual(confirmed["counters"]["source_imports"], repeated["counters"]["source_imports"])
        applied = request({"action": "learning", "work": str(self.work), "task_id": failed["task_id"],
                           "learning": {"action": "apply", "id": method_id}})
        self.assertTrue(applied["ok"], applied)
        self.assertEqual(repeated["counters"]["source_imports"], applied["counters"]["source_imports"])

    def test_legacy_not_saved_snapshot_without_next_action_can_resume(self) -> None:
        started = self.start_awaiting_source(batch="20260929-04")
        self.write_legacy_not_saved_task(started["task_id"])

        restored = request({"action": "status", "work": str(self.work), "task_id": started["task_id"]})
        self.assertTrue(restored["ok"], restored)
        self.assertEqual("awaiting_source", restored["status"])
        self.assertTrue(restored["questions"])
        self.assertIn("同一 task_id", restored["next_action"])

    def test_script_change_during_confirmation_rejects_the_pending_method(self) -> None:
        failed = self.start_failed()
        script = self.work / "recognise.py"
        script.write_text("# first version\n", encoding="utf-8")
        proposed = self.propose(failed["task_id"], script_path=script.name)
        script.write_text("# changed version\n", encoding="utf-8")
        rejected = self.confirm(proposed)
        self.assertFalse(rejected["ok"])
        self.assertIn("确认期间", rejected["error"])
        self.assertEqual([], self.method_list(failed["task_id"]))

    def test_not_saved_method_does_not_write_facts_or_method_record(self) -> None:
        failed = self.start_failed()
        cancelled = self.confirm(self.propose(failed["task_id"]), selected="暂不保存")
        self.assertEqual("cancelled", cancelled["status"])
        self.assertEqual("not_saved", cancelled["learning"]["status"])
        self.assertIn("原任务", cancelled["next_action"])
        self.assertEqual(0, self.fact_count())
        self.assertEqual([], self.method_list(failed["task_id"]))

    def test_not_saved_method_cancels_this_call_without_reopening_existing_source_questions(self) -> None:
        started = self.start_awaiting_source()
        cancelled = self.confirm(self.propose(started["task_id"]), selected="暂不保存")
        self.assertTrue(cancelled["ok"], cancelled)
        self.assertEqual("cancelled", cancelled["status"])
        self.assertEqual("not_saved", cancelled["learning"]["status"])
        self.assertEqual([], cancelled["questions"])
        self.assertIn("原任务", cancelled["next_action"])
        self.assertEqual(0, cancelled["execution_summary"]["generated_files"])
        self.assertEqual(0, self.fact_count())
        self.assertEqual([], self.method_list(started["task_id"]))

        # A normal status call restores the preserved source workflow instead
        # of treating the transient method cancellation as a permanent block.
        restored = request({"action": "status", "work": str(self.work), "task_id": started["task_id"]})
        self.assertEqual("awaiting_source", restored["status"])
        self.assertTrue(restored["questions"])
        self.assertEqual(0, restored["execution_summary"]["generated_files"])

        # A corrected method can still be proposed and confirmed on this task.
        revised = self.propose(started["task_id"], name="修订后的合成用途识别")
        self.assertEqual("awaiting_method", revised["status"])
        confirmed = self.confirm(revised)
        self.assertTrue(confirmed["ok"], confirmed)
        self.assertEqual("completed", confirmed["status"])
        self.assertEqual(1, self.fact_count())

    def test_apply_rejects_old_evidence_after_source_changes(self) -> None:
        failed = self.start_failed()
        confirmed = self.confirm(self.propose(failed["task_id"]))
        self.assertEqual("completed", confirmed["status"], confirmed)
        method_id = self.method_list(failed["task_id"])[0]["id"]
        changed = Document()
        changed.add_paragraph("本次资金用于其他合成用途。")
        changed.save(self.source)
        rejected = request({"action": "learning", "work": str(self.work), "task_id": failed["task_id"],
                            "learning": {"action": "apply", "id": method_id}})
        self.assertTrue(rejected["ok"], rejected)
        self.assertEqual("needs_source_update", rejected["status"])
        self.assertEqual("needs_refresh", rejected["learning"]["status"])
        self.assertIn(self.source.name, {item["source"] for item in rejected["learning"]["changed_sources"]})
        self.assertEqual(0, rejected["execution_summary"]["generated_files"])
        self.assertEqual(1, self.fact_count())

    def test_saved_target_rules_apply_to_a_new_batch(self) -> None:
        failed = self.start_failed()
        first = self.confirm(self.propose(failed["task_id"]))
        self.assertEqual("completed", first["status"], first)
        rules_method = self.confirm(self.propose(failed["task_id"], name="合成用途和位置", save_target_rules=True))
        self.assertEqual("completed", rules_method["status"], rules_method)
        method = next(item for item in self.method_list(failed["task_id"])
                      if item["name"] == "合成用途和位置")

        fresh = request({"action": "start", "work": str(self.work), "source": [self.source.name],
                         "targets": [self.target.name], "batch": "20260929-02"})
        self.assertTrue(fresh["ok"], fresh)
        self.assertEqual("completed", fresh["status"], fresh)
        self.assertEqual([], fresh["questions"])
        self.assertEqual(1, fresh["counters"]["source_imports"])
        reopened = load_workbook(fresh["results"][0]["output"], data_only=True)
        try:
            self.assertIn("流动资金", str(reopened.active["A1"].value))
        finally:
            reopened.close()
        again = request({"action": "learning", "work": str(self.work), "task_id": fresh["task_id"],
                         "learning": {"action": "apply", "id": method["id"]}})
        self.assertTrue(again["ok"], again)
        self.assertEqual("completed", again["status"], again)
        self.assertEqual(fresh["counters"]["source_imports"], again["counters"]["source_imports"])

    def test_archived_recipe_extracts_changed_source_for_a_new_confirmed_fill(self) -> None:
        failed = self.start_failed()
        script = self.work / "recognise.py"
        script.write_text(
            "from docx import Document\nimport re\nimport sys\n"
            "text = '\\n'.join(p.text for p in Document(sys.argv[1]).paragraphs)\n"
            "print(re.search(r'用于(.+?)。', text).group(1))\n", encoding="utf-8")
        first_value = self.run_recipe(script, self.source)
        first = self.confirm(self.propose(failed["task_id"], script_path=script.name,
                                          source_updates=self.updates(first_value)))
        self.assertEqual("completed", first["status"], first)
        method = self.method_list(failed["task_id"])[0]
        archived = Path(request({"action": "learning", "work": str(self.work), "task_id": failed["task_id"],
                                 "learning": {"action": "read", "id": method["id"]}})["learning"]["script_path"])

        changed = Document()
        changed.add_paragraph("本次资金用于设备采购。")
        changed.save(self.source)
        second_value = self.run_recipe(archived, self.source)
        self.assertEqual("设备采购", second_value)
        second = self.confirm(self.propose(failed["task_id"], name="合成用途识别（新来源）",
                                           script_path=str(archived),
                                           source_updates=self.updates(second_value)))
        self.assertEqual("completed", second["status"], second)
        output = load_workbook(second["results"][0]["output"], data_only=True)
        try:
            text = str(output.active["A1"].value)
            self.assertIn("设备采购", text)
            self.assertNotIn("流动资金", text)
        finally:
            output.close()


if __name__ == "__main__":
    unittest.main()
