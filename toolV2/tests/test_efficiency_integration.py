"""Public synthetic fill, local unit mismatch, versions and legacy-rule recovery."""
import hashlib
import json
from pathlib import Path
import sqlite3
import sys
import tempfile
import unittest
from docx import Document
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Border, Side

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from office_kit.store_v2 import StoreV2
from test_mapping_workflow import request


def answer(result):
    while result['status'] in ('awaiting_source', 'awaiting_fill'):
        result = request({'action': 'resume', 'work': result['work'], 'task_id': result['task_id'],
                          'answers': [{'id': q['id'], 'selected': [q['options'][0]['label']], 'custom': ''}
                                      for q in result['questions']]})
    return result


class EfficiencyIntegrationTests(unittest.TestCase):
    def test_aliases_share_one_native_question_and_reuse_answer(self):
        with tempfile.TemporaryDirectory() as tmp:
            work = Path(tmp)
            doc = Document(); doc.add_paragraph('借款人名称：合成别名公司')
            doc.add_paragraph('联系电话：13800000000'); doc.save(work / 'source.docx')
            for name in ('first.docx', 'second.docx'):
                target = Document(); target.add_paragraph('待关联内容 ____'); target.save(work / name)
            payload = {'action': 'start', 'work': str(work), 'source': ['source.docx'],
                       'targets': ['first.docx', 'second.docx'], 'batch': '20260925-85'}
            result = request(payload)
            while result['status'] == 'awaiting_source':
                result = request({'action': 'resume', 'work': str(work), 'task_id': result['task_id'],
                                  'answers': [{'id': q['id'], 'selected': [q['options'][0]['label']], 'custom': ''}
                                              for q in result['questions']]})
            self.assertEqual('needs_mapping', result['status'])
            updates = []
            for name, field in [('first.docx', '联系电话'), ('second.docx', '借款人联系电话')]:
                page = request({'action': 'read_mapping', 'work': str(work), 'task_id': result['task_id'],
                                'mapping_read': {'section': 'positions', 'template': name}})
                updates.append({'template': name, 'field': field, 'slot_id': page['mapping_page']['items'][0]['id']})
            ready = request({'action': 'update_positions', 'work': str(work), 'task_id': result['task_id'], 'updates': updates})
            self.assertEqual('awaiting_fill', ready['status'], ready)
            self.assertEqual(1, len(ready['questions']), ready['questions'])
            self.assertIn('first.docx', ready['questions'][0]['detail'])
            self.assertIn('second.docx', ready['questions'][0]['detail'])
            completed = answer(ready)
            self.assertEqual('completed', completed['status'], completed)
            repeated = request(payload)
            self.assertEqual('completed', repeated['status'])
            self.assertEqual([], repeated['questions'])
            self.assertEqual(completed['counters']['fill_processes'], repeated['counters']['fill_processes'])

    def test_date_range_blanks_only_affected_slot_and_preserves_old_versions(self):
        with tempfile.TemporaryDirectory() as tmp:
            work = Path(tmp)
            doc = Document()
            doc.add_paragraph('借款人名称：合成公司')
            doc.add_paragraph('借款期限：2026年1月1日至2027年1月1日')
            doc.add_paragraph('联系电话：13800000000')
            doc.save(work / 'source.docx')
            duration = Document(); duration.add_paragraph('借款期限：____年')
            duration.add_paragraph('联系电话：____'); duration.save(work / 'term.docx')
            phone = Document(); phone.add_paragraph('联系电话：____'); phone.save(work / 'phone.docx')
            hashes = {n: hashlib.sha256((work / n).read_bytes()).hexdigest() for n in ('term.docx', 'phone.docx')}
            data = {'action': 'start', 'work': str(work), 'source': ['source.docx'],
                    'targets': ['term.docx', 'phone.docx'], 'batch': '20260925-87'}
            first = answer(request(data))
            self.assertEqual('completed', first['status'], first)
            paths = {r['template']: r['output'] for r in first['results']}
            text = '\n'.join(p.text for p in Document(paths['term.docx']).paragraphs)
            self.assertIn('借款期限：____年', text)
            self.assertIn('13800000000', text)
            self.assertNotIn('2026年', text)
            report = Path(first['report']).read_text(encoding='utf-8')
            self.assertIn('日期区间', report)
            self.assertIn('阶段用时记录', report)
            old_hash = hashlib.sha256(Path(paths['term.docx']).read_bytes()).hexdigest()
            # Change only one template. Its current version advances; the other is reused.
            duration.add_paragraph('附注：这是合成模板的第二版。'); duration.save(work / 'term.docx')
            second = answer(request(data))
            self.assertEqual('completed', second['status'], second)
            changed = {r['template']: r['output'] for r in second['results']}
            self.assertEqual(paths['phone.docx'], changed['phone.docx'])
            self.assertNotEqual(paths['term.docx'], changed['term.docx'])
            self.assertEqual(old_hash, hashlib.sha256(Path(paths['term.docx']).read_bytes()).hexdigest())
            self.assertEqual(1, second['counters']['fill_processes'] - first['counters']['fill_processes'])
            versions = second['delivery']['versions']
            self.assertEqual(2, sum(v['status'] == 'current' for v in versions))
            self.assertEqual(1, sum(v['status'] == 'superseded' for v in versions))
            self.assertTrue(all(Path(p).is_absolute() for p in second['delivery']['attachments']))
            self.assertTrue(all(len(g) <= 8 for g in second['delivery']['attachment_groups']))
            self.assertEqual(hashes['phone.docx'], hashlib.sha256((work / 'phone.docx').read_bytes()).hexdigest())

    def test_old_double_xlsx_rule_is_migrated_to_one_value_cell(self):
        with tempfile.TemporaryDirectory() as tmp:
            work = Path(tmp)
            doc = Document(); doc.add_paragraph('借款人名称：合成甲公司'); doc.save(work / 'source.docx')
            book = Workbook(); book.active['A1'] = '借款人名称：'
            book.active['B1'].border = Border(bottom=Side(style='thin'))
            book.save(work / 'target.xlsx'); book.close()
            data = {'action': 'start', 'work': str(work), 'source': ['source.docx'],
                    'targets': ['target.xlsx'], 'batch': '20260925-86'}
            first = answer(request(data))
            self.assertEqual('completed', first['status'], first)
            with StoreV2(work / 'db/workflow.db', actor='synthetic-test') as store:
                tid = store.register_template(work / 'target.xlsx', first['batch'])
                store.add_rule(tid, '借款人名称', '借款人名称', {'kind': 'multi', 'targets': [
                    {'kind': 'xlsx_cell', 'sheet': 'Sheet', 'cell': 'A1', 'anchor': '借款人名称：'},
                    {'kind': 'xlsx_cell', 'sheet': 'Sheet', 'cell': 'B1'}]},
                    confidence=1, decided_by='model', batch_no=first['batch'])
                task = json.loads(store.conn.execute('SELECT payload FROM office_v2_task WHERE id=?', (first['task_id'],)).fetchone()[0])
                task['documents']['target.xlsx']['proposed_version'] = 4
                store.conn.execute('UPDATE office_v2_task SET payload=? WHERE id=?', (json.dumps(task), first['task_id']))
                store.conn.commit()
            resumed = answer(request({'action': 'status', 'work': str(work), 'task_id': first['task_id']}))
            self.assertEqual('completed', resumed['status'], resumed)
            self.assertEqual(1, resumed['results'][0]['filled'])
            output = load_workbook(resumed['results'][0]['output'])
            self.assertEqual('借款人名称：', output.active['A1'].value)
            self.assertEqual('合成甲公司', output.active['B1'].value)
            output.close()


if __name__ == '__main__':
    unittest.main()
