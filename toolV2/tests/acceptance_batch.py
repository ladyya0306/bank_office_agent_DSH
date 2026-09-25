"""Synthetic fifteen-template acceptance; never opens customer workspaces."""
from __future__ import annotations

import hashlib
import json
from pathlib import Path
import shutil
import sqlite3
import subprocess
import sys
import tempfile
import time
from contextlib import closing

from docx import Document
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]


def request(payload):
    p = subprocess.run([sys.executable, str(ROOT / 'office.py')],
                       input=json.dumps(payload, ensure_ascii=False), encoding='utf-8',
                       capture_output=True, timeout=240, cwd=ROOT)
    try:
        result = json.loads(p.stdout)
    except ValueError as exc:
        raise AssertionError(f'Invalid response: {p.stdout[-500:]} {p.stderr[-1000:]}') from exc
    assert result.get('ok'), result
    return result


def finish(result):
    questions = 0
    versions = set()
    while result['status'] in ('awaiting_source', 'awaiting_fill'):
        signature = json.dumps(result['questions'], sort_keys=True, ensure_ascii=False)
        assert signature not in versions, 'Repeated unchanged questions'
        versions.add(signature)
        questions += len(result['questions'])
        result = request({'action': 'resume', 'work': result['work'],
                          'task_id': result['task_id'],
                          'answers': [{'id': q['id'], 'selected': [q['options'][0]['label']],
                                       'custom': ''} for q in result['questions']]})
    assert result['status'] == 'completed', result
    return result, questions


def counts(work):
    with closing(sqlite3.connect(work / 'db/workflow.db')) as conn:
        return dict(conn.execute('select event_type,count(*) from event group by event_type'))


def main():
    started = time.monotonic()
    with tempfile.TemporaryDirectory(prefix='toolv2_acceptance_') as tmp:
        work = Path(tmp) / '项目甲'
        work.mkdir()
        source = work / '来源.docx'
        doc = Document()
        doc.add_paragraph('借款人名称：演示甲企业')
        doc.add_paragraph('收款账号：123456789')
        doc.save(source)
        targets = []
        for i in range(12):
            p = work / f'目标{i + 1}.docx'
            doc = Document()
            doc.add_paragraph('收款账号：')
            doc.save(p)
            targets.append(p)
        for i in range(3):
            p = work / f'台账{i + 1}.xlsx'
            wb = Workbook()
            wb.active['A1'] = '收款账号：'
            wb.save(p)
            wb.close()
            targets.append(p)
        originals = {str(p): hashlib.sha256(p.read_bytes()).hexdigest() for p in [source, *targets]}
        payload = {'action': 'start', 'work': str(work), 'source': [str(source)],
                   'targets': [str(p) for p in targets], 'batch': '20260925-01'}
        first, asked = finish(request(payload))
        assert len(first['results']) == 15, first
        produced = list((work / 'out').rglob('*_已填写.docx')) + list((work / 'out').rglob('*_已填写.xlsx'))
        assert len(produced) == 15, produced
        for p in produced:
            if p.suffix == '.docx':
                assert '123456789' in '\n'.join(x.text for x in Document(p).paragraphs), p
            else:
                wb = load_workbook(p, data_only=True)
                assert any('123456789' in str(c.value) for row in wb.active for c in row), p
                wb.close()
        before = counts(work)
        second, repeated_questions = finish(request(payload))
        assert repeated_questions == 0, second
        assert counts(work).get('fill_executed', 0) == before.get('fill_executed', 0), 'Unchanged files were filled twice'
        assert originals == {str(p): hashlib.sha256(p.read_bytes()).hexdigest() for p in [source, *targets]}
        # Simulate the same project being moved to another computer's folder.
        moved = Path(tmp) / '另一个位置' / '项目甲'
        shutil.copytree(work, moved)
        migrated = {**payload, 'work': str(moved), 'source': [str(moved / source.name)],
                    'targets': [str(moved / p.name) for p in targets]}
        third, moved_questions = finish(request(migrated))
        assert moved_questions == 0, third
        assert counts(moved).get('fill_executed', 0) == before.get('fill_executed', 0), 'Move reran unchanged outputs'
        print(json.dumps({'passed': True, 'templates': 15, 'first_questions': asked,
                          'repeat_questions': repeated_questions, 'moved_questions': moved_questions,
                          'seconds': round(time.monotonic() - started, 2)}, ensure_ascii=False))


if __name__ == '__main__':
    main()
