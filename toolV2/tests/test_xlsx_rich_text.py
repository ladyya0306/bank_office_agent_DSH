from openpyxl import Workbook, load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont

from office_kit.doc_fill import XlsxEngine


def rich(*parts):
    return CellRichText(*[TextBlock(font, text) for text, font in parts])


def test_rich_text_roundtrip_and_span_replacements(tmp_path):
    path = tmp_path / 'input.xlsx'; out = tmp_path / 'out.xlsx'
    wb = Workbook(); ws = wb.active
    normal, bold = InlineFont(sz=11, b=False), InlineFont(sz=14, b=True)
    ws['A1'] = rich(('说明', bold), ('保持', normal))
    ws['A2'] = rich(('标签：', bold), ('____', normal), ('后缀', bold))
    ws['A3'] = rich(('甲____乙____', normal))
    ws['A4'] = '=1+1'; wb.save(path)
    engine = XlsxEngine(path)
    engine.fill_xlsx_cell({'sheet': 'Sheet', 'cell': 'A2', 'anchor': '标签：', 'before': '后缀'}, '值')
    engine.fill_xlsx_cell({'sheet': 'Sheet', 'cell': 'A3', 'expected_text': '甲____乙____', 'span_start': 6, 'span_end': 10}, '二')
    engine.fill_xlsx_cell({'sheet': 'Sheet', 'cell': 'A3', 'expected_text': '甲____乙____', 'span_start': 1, 'span_end': 5}, '一')
    engine.save(out)
    check = load_workbook(out, rich_text=True); ws = check.active
    assert str(ws['A1'].value) == '说明保持'
    assert ws['A1'].value[0].font.b is True and ws['A1'].value[1].font.b is False
    assert ws['A1'].value[0].font.sz == 14 and ws['A1'].value[1].font.sz == 11
    assert str(ws['A2'].value) == '标签：值后缀'
    assert ws['A2'].value[-1].font.b is True
    assert ws['A2'].value[1].text == '值' and ws['A2'].value[1].font.b is False
    assert str(ws['A3'].value) == '甲一乙二'
    assert ws['A4'].value == '=1+1'
