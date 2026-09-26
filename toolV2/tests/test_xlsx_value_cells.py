"""End-to-end migration of legacy Excel rules into unit-bearing value cells."""
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import Workbook, load_workbook

from office_kit.store_v2 import StoreV2, sha256_file
from office_kit.workroot import init_workroot
from workflow import runner, storage


BATCH = '20260926-01'


def _template(path: Path) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = 'Sheet1'
    sheet['D3'] = '额度：'
    sheet['E3'] = '人民币  万元'
    sheet['A4'] = '本次业务金额：'
    sheet['B4'] = '人民币   万元'
    sheet['D4'] = '已使用额度：'
    sheet['E4'] = '人民币  万元'
    book.save(path)
    book.close()


def _old_multi(*targets):
    return {'kind': 'multi', 'targets': list(targets)}


def test_prepare_migrates_legacy_rules_and_writes_each_distinct_amount_once():
    with TemporaryDirectory() as tmp:
        work = Path(tmp)
        init_workroot(work)
        template = work / 'targets' / '额度台账.xlsx'
        template.parent.mkdir(parents=True, exist_ok=True)
        _template(template)
        name = template.relative_to(work).as_posix()

        with StoreV2(work / 'db' / 'workflow.db', actor='test') as store:
            storage.initialise(store.conn)
            store.batch_scope = BATCH
            tid = store.register_template(template, BATCH)
            # Deliberately seed the old duplicate and mislabeled locations.
            store.add_rule(tid, '授信额度', '额度', _old_multi(
                {'kind': 'xlsx_cell', 'sheet': 'Sheet1', 'cell': 'D3', 'anchor': '额度：'},
                {'kind': 'xlsx_cell', 'sheet': 'Sheet1', 'cell': 'E3',
                 'expected_text': '人民币  万元', 'span_start': 5, 'span_end': 5}),
                confidence=.99, decided_by='human', batch_no=BATCH)
            store.add_rule(tid, '借款金额', '本次业务金额', _old_multi(
                {'kind': 'xlsx_cell', 'sheet': 'Sheet1', 'cell': 'A4',
                 'expected_text': '本次业务金额：', 'span_start': 7, 'span_end': 7},
                {'kind': 'xlsx_cell', 'sheet': 'Sheet1', 'cell': 'E4',
                 'expected_text': '人民币  万元', 'span_start': 5, 'span_end': 5}),
                confidence=.99, decided_by='human', batch_no=BATCH)
            store.add_rule(tid, '保证金额', '保证金额',
                {'kind': 'xlsx_cell', 'sheet': 'Sheet1', 'cell': 'B4',
                 'expected_text': '人民币   万元', 'span_start': 6, 'span_end': 6},
                confidence=.99, decided_by='human', batch_no=BATCH)
            store.add_rule(tid, '已用额度', '已使用额度',
                {'kind': 'xlsx_cell', 'sheet': 'Sheet1', 'cell': 'D4', 'anchor': '已使用额度：'},
                confidence=.99, decided_by='human', batch_no=BATCH)

            for field, value in [('授信额度', '800万元'), ('借款金额', '300万元'),
                                 ('已使用额度', '0万元')]:
                store.put_fact(BATCH, field, value, source_kind='source',
                               provenance='合成测试源文件')

            record = {'proposed_version': runner.POSITION_PARSE_VERSION,
                      'proposed_hash': sha256_file(template)}
            task = {'id': 'synthetic-value-cell-migration', 'batch': BATCH,
                    'source': [], 'targets': [name], 'documents': {name: record},
                    'counts': {'source_reads': 0, 'source_imports': 0,
                               'position_proposals': 0, 'previews': 0,
                               'fill_processes': 0}, 'last_work': str(work),
                    'source_signature': None, 'source_content_signature': None,
                    'equivalent_source_signatures': []}

            runner.prepare_documents(store, work, task)
            record = task['documents'][name]
            assert record['status'] == 'pending', record.get('error')
            assert record['value_cell_version']

            positions = {rule['field']: rule['target'] for rule in store.rules_for(template)}
            assert set(positions) == {'授信额度', '借款金额', '已使用额度'}
            assert positions['授信额度']['cell'] == 'E3'
            assert positions['借款金额']['cell'] == 'B4'
            assert positions['已使用额度']['cell'] == 'E4'
            assert all(target.get('kind') != 'multi' for target in positions.values())

            # Confirm the old B4 rule is disabled in history, then accept the
            # suggested source values and run the normal fill pipeline.
            disabled = store.conn.execute(
                'SELECT field FROM template_rule_disabled WHERE template_id=?', (tid,)
            ).fetchall()
            assert {row[0] for row in disabled} >= {'保证金额', '已用额度'}
            assert [(row['field'], row['value'], row['decision'])
                    for row in record['plan']['rows']] == [
                        ('授信额度', '800万元', 'auto'),
                        ('借款金额', '300万元', 'auto'),
                        ('已使用额度', '0万元', 'auto')]
            runner.execute_documents(store, work, task)
            assert record['status'] == 'completed', record.get('error')
            output = load_workbook(work / record['output'], data_only=False)
            sheet = output['Sheet1']
            assert sheet['D3'].value == '额度：'
            assert sheet['E3'].value == '人民币800万元'
            assert sheet['A4'].value == '本次业务金额：'
            assert sheet['B4'].value == '人民币300万元'
            assert sheet['D4'].value == '已使用额度：'
            assert sheet['E4'].value == '人民币0万元'
            output.close()
