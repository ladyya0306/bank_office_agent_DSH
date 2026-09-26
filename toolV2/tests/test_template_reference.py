"""Explicit same-template references stay traceable and never auto-copy."""
from docx import Document
from openpyxl import Workbook, load_workbook
import pytest

from office_kit.store_v2 import StoreV2
from office_kit.template_reference import resolve
from workflow.runner import dispatch


def answer(result):
    while result['status'] in ('awaiting_source', 'awaiting_fill'):
        result = dispatch({'action': 'resume', 'work': result['work'], 'task_id': result['task_id'],
                           'answers': [{'id': q['id'], 'selected': [q['options'][0]['label']], 'custom': ''}
                                       for q in result['questions']]})
    return result


def workbook(path):
    book = Workbook(); one = book.active; one.title = 'Sheet1'
    one['D2'] = '\u7ecf\u529e\u673a\u6784\u540d\u79f0\uff1a\u5408\u6210\u652f\u884c'
    two = book.create_sheet('Sheet2'); two['C2'] = '\u7ecf\u529e\u673a\u6784\u540d\u79f0\uff1a____'
    book.save(path)


def test_explicit_template_reference_fills_once_and_refuses_formula_refresh(tmp_path):
    source = Document(); source.add_paragraph('\u501f\u6b3e\u4eba\uff1a\u5408\u6210\u4f01\u4e1a'); source.save(tmp_path / 'source.docx')
    target = tmp_path / 'target.xlsx'; workbook(target)
    started = answer(dispatch({'action': 'start', 'work': str(tmp_path), 'source': ['source.docx'],
                               'targets': ['target.xlsx'], 'batch': '20260926-94'}))
    slot = started['mapping_requests'][0]['positions'][0]
    done = dispatch({'action': 'update_positions', 'work': str(tmp_path), 'task_id': started['task_id'],
                     'updates': [{'template': 'target.xlsx', 'slot_id': slot['id'],
                                  'template_reference': {'sheet': 'Sheet1', 'cell': 'D2'}}]})
    done = answer(done)
    assert done['status'] == 'completed', done
    output = load_workbook(done['results'][0]['output']); assert output['Sheet2']['C2'].value.endswith('\u5408\u6210\u652f\u884c')
    counters = done['counters'].copy()
    assert dispatch({'action': 'status', 'work': str(tmp_path), 'task_id': done['task_id']})['counters'] == counters
    book = load_workbook(target); book['Sheet1']['D2'] = '=1+1'; book.save(target)
    changed = dispatch({'action': 'start', 'work': str(tmp_path), 'source': ['source.docx'],
                        'targets': ['target.xlsx'], 'batch': '20260926-94'})
    assert changed['status'] != 'completed' or changed.get('error')


@pytest.mark.parametrize('value', ['', '____', '□', '=1+1', '\u5176\u4ed6\u6807\u7b7e\uff1a\u5408\u6210\u652f\u884c'])
def test_reference_rejects_blank_placeholder_formula_or_different_label(tmp_path, value):
    path = tmp_path / 'bad.xlsx'; workbook(path)
    book = load_workbook(path); book['Sheet1']['D2'] = value; book.save(path)
    with StoreV2(tmp_path / 'data.db') as store:
        with pytest.raises(ValueError):
            resolve(store, path, {'kind': 'xlsx_cell', 'sheet': 'Sheet2', 'cell': 'C2'},
                    {'sheet': 'Sheet1', 'cell': 'D2'}, '\u7ecf\u529e\u673a\u6784\u540d\u79f0')


def test_reference_rejects_explicit_borrower_guarantor_context_mismatch(tmp_path):
    path = tmp_path / 'roles.xlsx'; workbook(path)
    book = load_workbook(path)
    book['Sheet1']['A2'] = '\u501f\u6b3e\u4eba'
    book['Sheet2']['A2'] = '\u4fdd\u8bc1\u4eba2'
    book.save(path)
    with StoreV2(tmp_path / 'data.db') as store:
        with pytest.raises(ValueError, match='\u660e\u786e\u4e3b\u4f53'):
            resolve(store, path, {'kind': 'xlsx_cell', 'sheet': 'Sheet2', 'cell': 'C2'},
                    {'sheet': 'Sheet1', 'cell': 'D2'}, '\u7ecf\u529e\u673a\u6784\u540d\u79f0')


@pytest.mark.parametrize('extra', [{'field': '\u501f\u6b3e\u4eba\u540d\u79f0'}, {'leave_blank': True}])
def test_reference_is_mutually_exclusive_with_field_or_blank(tmp_path, extra):
    source = Document(); source.add_paragraph('\u501f\u6b3e\u4eba\uff1a\u5408\u6210\u4f01\u4e1a'); source.save(tmp_path / 'source.docx')
    target = tmp_path / 'target.xlsx'; workbook(target)
    started = answer(dispatch({'action': 'start', 'work': str(tmp_path), 'source': ['source.docx'],
                               'targets': ['target.xlsx'], 'batch': '20260926-93'}))
    slot = started['mapping_requests'][0]['positions'][0]
    result = dispatch({'action': 'update_positions', 'work': str(tmp_path), 'task_id': started['task_id'],
                       'updates': [{**extra, 'template': 'target.xlsx', 'slot_id': slot['id'],
                                    'template_reference': {'sheet': 'Sheet1', 'cell': 'D2'}}]})
    assert 'template_reference' in result['rejected_updates'][0]['reason']
