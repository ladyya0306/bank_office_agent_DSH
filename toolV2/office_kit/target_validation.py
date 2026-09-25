"""Read-only validation of fill targets before a rule is persisted or used."""
from __future__ import annotations

from pathlib import Path
from typing import Any

from .common import OfficeKitError

_BLANK_MARKS = "_＿-—.·…□☐"


def _blank_slot(text: str) -> bool:
    return all(ch.isspace() or ch in _BLANK_MARKS for ch in str(text))


def _word_target(path: Path, target: dict[str, Any]) -> dict[str, Any]:
    from .doc_fill import XmlEngine, paragraph_text

    kind = target.get("kind")
    if kind == "multi":
        children = target.get("targets")
        if not isinstance(children, list) or not children:
            raise OfficeKitError("多位置规则必须提供非空 targets 数组")
        if any(isinstance(child, dict) and child.get("kind") == "multi" for child in children):
            raise OfficeKitError("multi 不允许嵌套 multi")
        return {"kind": "multi", "targets":
                [_word_target(path, child) for child in children]}
    if kind == "cell":
        required = ("table", "row", "col")
        if any(k not in target for k in required):
            raise OfficeKitError("cell 目标必须同时提供 table、row、col")
        engine = XmlEngine(path)
        try:
            table_i, row_i, col_i = (int(target[k]) for k in ("table", "row", "col"))
        except (TypeError, ValueError) as exc:
            raise OfficeKitError("cell 的 table、row、col 必须是整数") from exc
        if min(table_i, row_i, col_i) < 0:
            raise OfficeKitError("cell 的 table、row、col 不能为负数")
        try:
            cell = engine.document.tables[table_i].cell(row_i, col_i)
        except (IndexError, TypeError, ValueError) as exc:
            raise OfficeKitError("Word 表格位置不存在") from exc
        if not cell.paragraphs:
            raise OfficeKitError("目标单元格没有段落")
        for paragraph in cell.paragraphs:
            if engine.xml_for(paragraph) is None:
                raise OfficeKitError("目标单元格无法对应到原文件")
            if not _blank_slot(paragraph_text(paragraph)):
                raise OfficeKitError("目标单元格已有文字，拒绝覆盖")
        return {"kind": "cell"}
    if kind != "anchor":
        raise OfficeKitError("不支持的 Word 目标 kind：%r" % kind)
    precise = ("paragraph_index", "expected_text", "span_start", "span_end")
    if not target.get("anchor") and not target.get("pattern"):
        if not all(k in target for k in precise):
            raise OfficeKitError("anchor 目标必须提供 anchor/pattern，或完整精确段落区间")
        try:
            if int(target["paragraph_index"]) < 0:
                raise OfficeKitError("paragraph_index 不能为负数")
            int(target["span_start"])
            int(target["span_end"])
        except (TypeError, ValueError) as exc:
            raise OfficeKitError("paragraph_index/span_start/span_end 必须是整数") from exc
    engine = XmlEngine(path)
    hits = engine.find_anchor_hits(target)
    if not hits:
        raise OfficeKitError("找不到指定填写位置")
    max_matches = int(target.get("max_matches", 1))
    if max_matches < 1:
        raise OfficeKitError("max_matches 必须为正数")
    if len(hits) > max_matches and not target.get("only_first") and target.get("occurrence") is None:
        raise OfficeKitError("找到 %d 个位置，未指定 occurrence/only_first，拒绝写入" % len(hits))
    for par, start, end in hits:
        raw = engine.xml_for(par)
        text = raw.text if raw is not None else paragraph_text(par)
        if not _blank_slot(text[start:end]):
            raise OfficeKitError("目标范围包含已有文字，拒绝覆盖")
    return {"kind": "anchor", "matches": len(hits)}


def _xlsx_target(path: Path, target: dict[str, Any]) -> dict[str, Any]:
    import openpyxl

    required = ("cell",)
    if "cell" not in target:
        raise OfficeKitError("xlsx_cell 目标必须提供 cell")
    if "pattern" in target:
        raise OfficeKitError("xlsx_cell 不支持 pattern，只能使用 anchor/before")
    wb = openpyxl.load_workbook(str(path), read_only=False, data_only=False)
    try:
        sheet = target.get("sheet")
        try:
            ws = wb.worksheets[sheet] if isinstance(sheet, int) else (
                wb.worksheets[0] if sheet is None else wb[str(sheet)])
        except (IndexError, KeyError) as exc:
            raise OfficeKitError("Excel 工作表不存在") from exc
        cell = ws[str(target["cell"])]
        current = "" if cell.value is None else str(cell.value)
        if cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("=")):
            raise OfficeKitError("Excel 目标是公式单元格，禁止填充")
        if any(k in target for k in ("expected_text", "span_start", "span_end")):
            if not all(k in target for k in ("expected_text", "span_start", "span_end")):
                raise OfficeKitError("xlsx 精确位置必须提供 expected_text、span_start、span_end")
            try:
                start, end = int(target["span_start"]), int(target["span_end"])
            except (TypeError, ValueError) as exc:
                raise OfficeKitError("xlsx span_start/span_end 必须是整数") from exc
            if str(target["expected_text"]) != current:
                raise OfficeKitError("Excel 目标原文不匹配，拒绝修改")
            if start < 0 or end < start or end > len(current):
                raise OfficeKitError("Excel 精确区间越界")
            if not _blank_slot(current[start:end]):
                raise OfficeKitError("Excel 精确区间含已有文字，拒绝覆盖")
            return {"kind": "xlsx_cell", "sheet": ws.title, "cell": cell.coordinate}
        if not current.strip():
            return {"kind": "xlsx_cell", "sheet": ws.title, "cell": cell.coordinate}
        anchor = target.get("anchor")
        if not anchor or current.count(str(anchor)) != 1:
            raise OfficeKitError("Excel 目标已有内容且 anchor 未唯一命中，拒绝覆盖")
        start = current.index(str(anchor)) + len(str(anchor))
        before = target.get("before")
        end = current.index(str(before), start) if before else len(current)
        if not _blank_slot(current[start:end]):
            raise OfficeKitError("Excel anchor 后已有文字，拒绝覆盖")
        return {"kind": "xlsx_cell", "sheet": ws.title, "cell": cell.coordinate}
    finally:
        wb.close()


def validate_target(path: str | Path, target: dict[str, Any]) -> dict[str, Any]:
    """Validate target shape and current template state without writing anything."""
    if not isinstance(target, dict):
        raise OfficeKitError("target 必须是对象")
    path = Path(path)
    if path.suffix.lower() == ".docx":
        return _word_target(path, target)
    if path.suffix.lower() == ".xlsx":
        if target.get("kind") == "multi":
            children = target.get("targets")
            if not isinstance(children, list) or not children:
                raise OfficeKitError("多位置规则必须提供非空 targets 数组")
            if any(isinstance(child, dict) and child.get("kind") == "multi" for child in children):
                raise OfficeKitError("multi 不允许嵌套 multi")
            return {"kind": "multi", "targets":
                    [_xlsx_target(path, child) for child in children]}
        if target.get("kind") != "xlsx_cell":
            raise OfficeKitError("Excel 目标 kind 必须为 xlsx_cell")
        return _xlsx_target(path, target)
    raise OfficeKitError("不支持的模板类型：%s" % path.suffix)


__all__ = ["validate_target"]
