"""Check that an edited workbook kept the template's layout and other cells.

Unlike DOCX, openpyxl rewrites XLSX XML on save.  Byte comparisons would reject
ordinary cell edits, so this compares the workbook features users actually see.
"""
from __future__ import annotations

from pathlib import Path
from typing import Any, Iterable


def _cell_state(cell: Any, include_value: bool) -> tuple:
    state = (
        tuple(cell._style) if cell.has_style else None,
        cell.number_format,
        cell.hyperlink.target if cell.hyperlink else None,
        cell.comment.text if cell.comment else None,
    )
    if include_value:
        return (cell.value, cell.data_type, *state)
    return state


def _dimensions(ws: Any) -> tuple:
    # openpyxl discards explicit row tags that carry no height/style/visibility
    # when saving. Those are visually identical to absent row tags.
    rows = tuple(sorted((i, d.height, d.hidden, d.outlineLevel, d.collapsed,
                         tuple(d._style) if d.has_style else None)
                        for i, d in ws.row_dimensions.items()
                        if d.height is not None or d.hidden or d.outlineLevel
                        or d.collapsed or d.has_style))
    cols = tuple(sorted((i, d.width, d.hidden, d.outlineLevel, d.collapsed,
                         d.min, d.max) for i, d in ws.column_dimensions.items()))
    return rows, cols


def prove_xlsx_fidelity(template: Path, written: Path,
                        edited_cells: Iterable[tuple[str, str, str]]) -> dict[str, Any]:
    """Allow values only at addressed cells; compare styles and visible layout.

    This is a semantic check, not a claim that OOXML bytes or rendered pixels are
    identical.  Unsupported package features must be assessed separately.
    """
    import openpyxl

    changed = {(sheet, address.upper()) for sheet, address, _ in edited_cells}
    a = openpyxl.load_workbook(template)
    b = openpyxl.load_workbook(written)
    issues: list[str] = []
    try:
        if a.sheetnames != b.sheetnames:
            issues.append("工作表名称或顺序改变")
        for name in a.sheetnames:
            if name not in b:
                continue
            left, right = a[name], b[name]
            if str(left.merged_cells) != str(right.merged_cells):
                issues.append(f"{name} 合并单元格改变")
            if _dimensions(left) != _dimensions(right):
                issues.append(f"{name} 行高或列宽改变")
            for attr, label in (("freeze_panes", "冻结窗格"),
                                ("sheet_state", "隐藏状态"),
                                ("print_area", "打印区域")):
                if str(getattr(left, attr)) != str(getattr(right, attr)):
                    issues.append(f"{name} {label}改变")
            for attr, label in (("page_setup", "页面设置"),
                                ("page_margins", "页边距"),
                                ("print_options", "打印选项")):
                if str(getattr(left, attr)) != str(getattr(right, attr)):
                    issues.append(f"{name} {label}改变")
            # Iterate stored cells only.  Walking the rectangular sheet bounds can
            # create millions of empty Cell objects in a sparse bank ledger.
            addresses = {cell.coordinate for cell in left._cells.values()}
            addresses.update(cell.coordinate for cell in right._cells.values())
            for address in sorted(addresses):
                editable = (name, address) in changed
                before = _cell_state(left[address], not editable)
                after = _cell_state(right[address], not editable)
                if before != after:
                    issues.append(f"{name}!{address} 的"
                                  + ("样式或附属信息" if editable else "内容或样式")
                                  + "改变")
            if len(left._charts) != len(right._charts) or len(left._images) != len(right._images):
                issues.append(f"{name} 图表或图片数量改变")
        for sheet, address in changed:
            if sheet not in b or b[sheet][address].value is None:
                issues.append(f"{sheet}!{address} 没有写入值")
    finally:
        a.close()
        b.close()
    return {"ok": not issues, "reasons": issues,
            "note": ("单元格格式、其他单元格、合并区域、行列尺寸和页面设置未变；"
                     "XLSX 文件内部字节未作相同承诺" if not issues else
                     "工作簿内容或布局改变：" + "；".join(issues[:8])),
            "parts_structurally_changed": issues,
            "check_kind": "xlsx_semantic"}
