"""Expose concrete template gaps so the model maps facts, not hand-written code."""
import re
import hashlib
from office_kit import doc_fill
from office_kit.common import OfficeKitError
from office_kit.store_v2 import sha256_file
from office_kit.harness import candidates_for
from office_kit.template_slots import discover_slots
from office_kit.target_validation import validate_target, read_word_template, read_excel_template

SLOT_DISCOVERY_VERSION = 8


class MappingUpdateError(ValueError):
    """Identify the actual failed proposal inside an atomic template update."""
    def __init__(self, message, item, slot=None):
        super().__init__(message)
        self.slot_id = item.get('slot_id')
        self.label = (slot or {}).get('label') or item.get('label')


def contains(outer, inner):
    return outer[:2] == inner[:2] and outer[2] <= inner[2] and outer[3] >= inner[3]


def intersects(left, right):
    return any(a[:2] == b[:2] and (contains(a, b) or contains(b, a)
               or max(a[2], b[2]) < min(a[3], b[3])) for a in left for b in right)


def fact_catalog(store):
    from office_kit.fact_catalog import qualified_facts
    return {**store.facts_by_key(), **qualified_facts(store)}


def physical(path, target, engine=None):
    if target.get('kind') == 'multi':
        return set().union(*(physical(path, t, engine) for t in target['targets']))
    if path.suffix.lower() == '.xlsx':
        import openpyxl
        wb = read_excel_template(path)
        try:
            sheet = target.get('sheet', 0)
            ws = wb.worksheets[sheet] if isinstance(sheet, int) else wb[sheet]
            text = str(ws[target['cell']].value or '')
            start = target.get('span_start')
            end = target.get('span_end')
            if start is None:
                start = text.find(target['anchor']) + len(target['anchor']) if target.get('anchor') else 0
                end = text.index(target['before'], start) if target.get('before') else len(text)
            return {(ws.title, target['cell'], start, end)}
        finally:
            wb.close()
    engine = engine or read_word_template(path)
    if target.get('kind') == 'cell':
        cell = engine.document.tables[int(target['table'])].cell(int(target['row']), int(target['col']))
        raw = engine.xml_for(cell.paragraphs[0])
        return {('word/document.xml', raw.start, 0, len(raw.text))}
    result = set()
    for par, start, end in engine.find_anchor_hits(target):
        raw = engine.xml_for(par)
        result.add((target.get('part', 'word/document.xml'), raw.start, start, end))
    return result


def inspect(store, path, record, facts):
    """Every discovered writable position is filled, expressly blank, or exposed."""
    from office_kit.fact_catalog import target_subject_context

    if (record.get('slot_hash') != sha256_file(path)
            or record.get('slot_version') != SLOT_DISCOVERY_VERSION):
        record['slots'] = discover_slots(path)
        record['slot_hash'] = sha256_file(path)
        record['slot_version'] = SLOT_DISCOVERY_VERSION
    engine = read_word_template(path) if path.suffix.lower() == '.docx' else None
    covered = set()
    for rule in store.rules_for(path):
        try:
            validate_target(path, rule['target'])
            covered.update(physical(path, rule['target'], engine))
        except (OfficeKitError, ValueError, KeyError, IndexError, TypeError):
            continue
    skipped = record.get('blank_slots', {})
    missing = []
    for slot in record['slots']:
        if slot.get('protected') or slot['id'] in skipped:
            continue
        positions = physical(path, slot['target'], engine)
        if positions and all(any(contains(c, p) for c in covered) for p in positions):
            continue
        subject = target_subject_context(store, path, slot['target'],
                                         label=slot.get('label', ''),
                                         context=slot.get('context'))
        target_subject = {
            'role': subject.get('role'),
            'number': subject.get('number'),
            'additional_party': bool(subject.get('additional_party')),
            'reason': subject.get('reason'),
            'field_hint': subject.get('field_hint'),
            'entity_names': [store.entity_label(eid) for eid in subject.get('entity_ids', [])],
        }
        missing.append({**slot, 'target_subject': target_subject,
                        'candidates': scoped_candidates(store, path, slot, facts)[:4]})
    record['unmapped_slots'] = missing
    record['coverage'] = {'detected': len(record['slots']), 'unmapped': len(missing),
                          'protected': sum(bool(s.get('protected')) for s in record['slots']),
                          'left_blank': len(skipped)}
    return missing


def visible_facts(facts):
    return [{'field': key, 'value': meta.get('value'), 'entity': meta.get('entity_name'),
             'source': meta.get('provenance'), 'fact_id': meta.get('fact_id', meta.get('id')),
             'entity_id': meta.get('entity_id'),
             'candidates': meta.get('_candidates', [])}
            for key, meta in facts.items()]


def visible_slot(slot):
    """Keep one useful excerpt, not repeated full paragraphs and fact payloads."""
    target = slot['target']
    context = dict(slot.get('context', {}))
    subject = slot.get('target_subject') or {}
    if 'expected_text' in target:
        raw = target['expected_text']
        start, end = target['span_start'], target['span_end']
        context['paragraph'] = raw[max(0, start-160):start] + '【填写处】' + raw[end:end+160]
    return {'id': slot['id'], 'label': slot.get('label'),
            'context': {key: value[:360] for key, value in context.items()},
            'target_subject': {key: subject.get(key) for key in
                               ('role', 'number', 'additional_party', 'reason', 'field_hint', 'entity_names')},
            'target': {key: value for key, value in target.items() if key != 'expected_text'},
            'candidates': [{'field': c['field'], 'score': c['score']}
                           for c in slot.get('candidates', [])[:3]],
            'full_context': {'section': 'context', 'slot_id': slot['id']}}


def semantic_compatible(field, slot):
    """Reject an obviously different *position meaning*, not just a low score.

    Labels on old shared/multi rules can describe another blank.  A source
    ``授信期限`` is not a board-meeting date, and a contract number is not a
    product/count field.  These are generic document semantics, evaluated
    from the visible position text rather than template names or coordinates.
    """
    target = slot.get('target', {})
    if target.get('template_reference'):
        return True  # The shared resolver checks the exact source/target label.
    raw = str(target.get('expected_text', ''))
    before = raw[:int(target.get('span_start', 0))]
    if re.search(r'(?:合同编号|编号为)[^，。；;（）()]{0,40}$', before) and not re.search(r'合同|协议', field):
        return False
    # Semantic guarding is position-local.  Paragraph/table context can
    # describe a neighbouring blank, so it must not reject a valid contract
    # number merely because another field in that context says “品种”.
    start, end = int(target.get('span_start', 0)), int(target.get('span_end', 0))
    # A paragraph can have both “合同编号：____；品种：____”.  Limit the
    # guard to this blank's clause, rather than letting the second blank's
    # label reject the first one.
    before_clause = re.split(r'[；;。\n]', str(target.get('expected_text', ''))[:start])[-1]
    after_clause = re.split(r'[；;。\n]', str(target.get('expected_text', ''))[end:])[0]
    local_text = ' '.join(str(value) for value in
                          (slot.get('label', ''), before_clause[-80:], after_clause[:80]) if value)
    if re.search(r'会议(?:时间|日期)', local_text) and re.search(r'期限', field):
        return False
    if re.search(r'(?:品种|份数)', local_text) and re.search(r'合同|协议|编号', field):
        return False
    return True


def scoped_candidates(store, path, slot, facts):
    """Rank source identities inside this position's actual subject context."""
    from office_kit.fact_catalog import target_subject_context
    from office_kit.harness import score_candidate
    scope = target_subject_context(store, path, slot['target'], label=slot.get('label', ''),
                                   context=slot.get('context'))
    if scope.get('additional_party'):
        return []
    identities = {}
    target = slot['target']
    raw = str(target.get('expected_text') or '')
    after = raw[int(target.get('span_end', 0)):]
    contract = re.match(r'[】\]）)\s]*的?[《【]\s*[【]?([^》】]+)[】]?[》】]', after)
    contract_title = contract.group(1).strip() if contract and '合同' in contract.group(1) else None
    for field, meta in facts.items():
        if not semantic_compatible(field, slot):
            continue
        owner = meta.get('_relation_owner_eid') or meta.get('entity_id')
        if scope['entity_ids'] and owner not in scope['entity_ids'] and owner is not None:
            continue
        if meta.get('_ambiguous'):
            continue
        score = score_candidate(slot.get('label', ''), field)
        hint = scope.get('field_hint')
        prefix = (scope.get('role') or '') + (str(scope['number']) if scope.get('number') is not None else '')
        # The same position-local meaning used by the fill planner must also
        # inform first-time mapping; otherwise clear name/ID slots wait for a
        # model to recreate the relationship the parser has already resolved.
        if hint and scope['entity_ids'] and (field == prefix + hint or meta.get('_base_field') == hint):
            score = 1.0
        if contract_title and re.sub(r'^(借款人|保证人|担保人)\d*', '', field) == contract_title:
            score = 1.0
        if meta.get('_base_field'):
            score = max(score, score_candidate(slot.get('label', ''), meta['_base_field']))
        if not scope['entity_ids']:
            # Generic “法定代表人” cannot favor the borrower's alias simply
            # because the old synonym table gives that role a higher score.
            neutral = re.sub(r'^(借款人|保证人|担保人)\d*', '', field)
            score = max(score, score_candidate(slot.get('label', ''), neutral))
        identity = (meta.get('fact_id', meta.get('id')), owner, meta.get('_base_field') or field)
        # Qualified aliases of a known fact must not create artificial ties.
        if identity[0] is not None:
            identity = identity[:2]
        old = identities.get(identity)
        if old is None or (score, bool(meta.get('_qualified'))) > (old['score'], old['qualified']):
            identities[identity] = {'field': field, 'score': score,
                                    'qualified': bool(meta.get('_qualified'))}
    return sorted(identities.values(), key=lambda x: -x['score'])


def _retire_semantic_mismatch(store, path, record, batch):
    """Retire only legacy rule targets whose current visible slot disagrees."""
    engine = read_word_template(path) if path.suffix.lower() == '.docx' else None
    slots = record.get('slots', [])
    for rule in list(store.rules_for(path)):
        children = rule['target'].get('targets', [rule['target']])
        kept, removed = [], False
        for target in children:
            try:
                position = physical(path, target, engine)
                matches = [slot for slot in slots
                           if intersects(position, physical(path, slot['target'], engine))]
            except (OfficeKitError, KeyError, IndexError, ValueError, TypeError):
                matches = []
            if matches and any(slot.get('protected') or (not target.get('template_reference') and
                               not semantic_compatible(rule['field'], slot)) for slot in matches):
                removed = True
            else:
                kept.append(target)
        if not removed:
            continue
        store.disable_rule(path, rule['field'], rule['target'],
                           '当前位置语义与来源字段不兼容，停用旧映射并保留历史')
        if kept:
            replacement = kept[0] if len(kept) == 1 else {'kind': 'multi', 'targets': kept}
            store.add_rule(rule['template_id'], rule['field'], rule['label'], replacement,
                           confidence=rule['confidence'], decided_by=rule['decided_by'], batch_no=batch)


def _retire_unbound_additional_party_auto_rules(store, path, record, batch):
    """Remove only prior automatic guesses in optional signer blocks.

    User/model rules remain editable and usable: this migration corrects the
    former automatic borrower-copy behavior without treating an optional block
    as a protected or permanently unfillable position.
    """
    from office_kit.fact_catalog import target_subject_context

    engine = read_word_template(path) if path.suffix.lower() == '.docx' else None
    slots = record.get('slots', [])
    for rule in list(store.rules_for(path)):
        if rule.get('decided_by') != 'auto':
            continue
        children = rule['target'].get('targets', [rule['target']])
        kept, removed = [], False
        for target in children:
            try:
                position = physical(path, target, engine)
                matches = [slot for slot in slots
                           if intersects(position, physical(path, slot['target'], engine))]
                additional = any(target_subject_context(store, path, slot['target'],
                                                        label=slot.get('label', ''),
                                                        context=slot.get('context')).get('additional_party')
                                 for slot in matches)
            except (OfficeKitError, KeyError, IndexError, ValueError, TypeError):
                additional = False
            if additional:
                removed = True
            else:
                kept.append(target)
        if not removed:
            continue
        store.disable_rule(path, rule['field'], rule['target'],
                           '可选签署方区块未声明主体；停用历史自动映射并保留历史')
        if kept:
            replacement = kept[0] if len(kept) == 1 else {'kind': 'multi', 'targets': kept}
            store.add_rule(rule['template_id'], rule['field'], rule['label'], replacement,
                           confidence=rule['confidence'], decided_by=rule['decided_by'], batch_no=batch)


def _best(candidates):
    if not candidates or candidates[0]['score'] < .85:
        return None
    if len(candidates) > 1 and candidates[1]['score'] == candidates[0]['score']:
        return None
    return candidates[0]


def _refresh_value_cells(store, path, record, facts, batch):
    """Replace old label/value duplicate mappings using the visible left label."""
    if path.suffix.lower() != '.xlsx' or record.get('value_cell_version') == SLOT_DISCOVERY_VERSION:
        return
    for slot in record.get('slots', []):
        target = slot['target']
        if not target.get('label_cell') or not slot.get('replaces_targets') or slot.get('protected'):
            continue
        picked = _best(scoped_candidates(store, path, slot, facts))
        for rule in list(store.rules_for(path)):
            children = rule['target'].get('targets', [rule['target']])
            kept = [c for c in children if not (
                c.get('sheet') == target.get('sheet') and
                c.get('cell') in (target['cell'], target['label_cell']))]
            if len(kept) == len(children):
                continue
            store.disable_rule(path, rule['field'], rule['target'],
                               '重新按左侧标签与右侧值格关联，原位置保留历史')
            if kept:
                store.add_rule(rule['template_id'], rule['field'], rule['label'],
                               kept[0] if len(kept) == 1 else {'kind':'multi','targets':kept},
                               confidence=rule['confidence'], decided_by=rule['decided_by'], batch_no=batch)
        if picked:
            item = {'field':picked['field'], 'slot_id':slot['id'],
                    'target':{**target,'slot_id':slot['id']}}
            replacement = combine(store,path,item)
            store.add_rule(store.register_template(path,batch),picked['field'],slot['label'],replacement,
                           confidence=picked['score'],decided_by='auto',batch_no=batch)
    record['value_cell_version'] = SLOT_DISCOVERY_VERSION


def _retire_obsolete_slots(store, path, record, batch):
    """Retire old detected blanks that the current parser identifies as layout."""
    if record.get('slot_migration_version') == SLOT_DISCOVERY_VERSION:
        return
    current = {s['id'] for s in record.get('slots', [])}
    engine = read_word_template(path) if path.suffix.lower() == '.docx' else None
    positions = [physical(path, s['target'], engine) for s in record.get('slots', [])]
    for rule in list(store.rules_for(path)):
        children = rule['target'].get('targets', [rule['target']])
        kept = []
        for target in children:
            if not target.get('slot_id') or target['slot_id'] in current:
                kept.append(target)
                continue
            try:
                previous = physical(path, target, engine)
                if any(intersects(previous, p) for p in positions):
                    kept.append(target)
            except (OfficeKitError, KeyError, ValueError, IndexError, TypeError):
                pass
        if len(kept) != len(children):
            store.disable_rule(path, rule['field'], rule['target'],
                               '旧检测空位已不再是可填写位置；保留原位置历史')
            if kept:
                store.add_rule(rule['template_id'], rule['field'], rule['label'],
                               kept[0] if len(kept) == 1 else {'kind':'multi','targets':kept},
                               confidence=rule['confidence'],decided_by=rule['decided_by'],batch_no=batch)
    record['slot_migration_version'] = SLOT_DISCOVERY_VERSION


def auto_map(store, path, record, facts, batch):
    """Map unique source facts in the position's role; expose actual ambiguities."""
    inspect(store, path, record, facts)
    _retire_semantic_mismatch(store, path, record, batch)
    _retire_unbound_additional_party_auto_rules(store, path, record, batch)
    _retire_obsolete_slots(store, path, record, batch)
    _refresh_value_cells(store, path, record, facts, batch)
    missing = inspect(store, path, record, facts)
    added = 0
    for slot in missing:
        candidates = scoped_candidates(store, path, slot, facts)
        slot['candidates'] = candidates[:4]
        chosen = _best(candidates)
        if chosen is None:
            continue
        field = chosen['field']
        item = {'field': field, 'slot_id': slot['id'], 'target': {**slot['target'], 'slot_id': slot['id']}}
        retire_overlap(store, path, item['target'], field, batch)
        target = combine(store, path, item)
        validate_target(path, target)
        tid = store.register_template(path, batch)
        store.add_rule(tid, field, slot.get('label') or field, target,
                       confidence=candidates[0]['score'], decided_by='auto', batch_no=batch)
        added += 1
    return added


def compile_updates(store, work, task, updates):
    """Validate the entire proposal before any shared rule is changed."""
    store.batch_scope = task['batch']
    facts = fact_catalog(store)
    checked, blanks = [], []
    seen = set()
    for item in updates:
        name = item['template']
        path = (work / name).resolve(strict=True)
        name = path.relative_to(work).as_posix()
        if name not in task['targets']:
            raise ValueError('只能修改当前任务的目标模板')
        field = item.get('field')
        reference = item.get('template_reference')
        if reference and (field or item.get('leave_blank')):
            raise ValueError('template_reference 不能与 field 或 leave_blank 同时提交')
        if reference and not item.get('leave_blank'):
            if (not isinstance(reference, dict) or set(reference) != {'sheet', 'cell'} or
                    not isinstance(reference.get('sheet'), str) or not isinstance(reference.get('cell'), str) or
                    not item.get('slot_id')):
                raise ValueError('template_reference 需要同模板的 sheet 与 cell')
            field = '模板引用:' + hashlib.sha256((reference['sheet'] + '!' + reference['cell']).encode()).hexdigest()[:16]
        if not reference and not item.get('leave_blank') and field not in facts:
            raise MappingUpdateError(f'来源中没有字段 {field!r}，请从 available_fields 中选择或明确留空', item)
        if task['documents'][name].get('slot_hash') != sha256_file(path):
            raise ValueError('目标文件已变化，请先继续原任务取得最新填写位置；此次没有修改规则')
        if not item.get('slot_id'):
            validate_target(path, item['target'])
            checked.append((path, item))
            continue
        slots = {s['id']: s for s in task['documents'][name].get('slots', [])}
        slot = slots.get(item['slot_id'])
        if slot is None:
            raise MappingUpdateError(f'{name} 未找到提交的 slot_id={item["slot_id"]}；请逐字使用 positions 返回的 id。不能据此判断其它位置已变化。', item)
        identity = (name, slot['id'])
        if identity in seen:
            raise MappingUpdateError(f'{name} 的同一空位重复提交，未写入规则', item, slot)
        seen.add(identity)
        if item.get('leave_blank'):
            reason = str(item.get('reason') or '').strip()
            if not reason:
                raise ValueError('留空需要说明缺失、不确定或不适用的原因')
            blanks.append((name, slot['id'], reason))
            continue
        if slot.get('protected'):
            raise MappingUpdateError('签字、签核位置不能通过填表映射自动代填', item, slot)
        if not reference and not semantic_compatible(field, slot):
            raise MappingUpdateError(f'位置 {slot["id"]}（{slot.get("label") or "未标注"}）与字段 {field} 语义不兼容；合同编号不能填写到合同品种或份数位置', item, slot)
        target = dict(slot['target'])
        if reference:
            from office_kit.template_reference import resolve
            resolve(store, path, target, reference, slot.get('label') or '')
            target['template_reference'] = {'sheet': reference['sheet'], 'cell': reference['cell']}
        target['slot_id'] = slot['id']
        validate_target(path, target)
        checked.append((path, {**item, 'field': field, 'target': target,
                               'label': slot.get('label') or field}))
    return checked, blanks


def combine(store, path, item):
    """Add a precise position; retain other positions for the same source field."""
    field, target = item['field'], item['target']
    old = next((r for r in store.rules_for(path) if r['field'] == field), None)
    if not item.get('slot_id'):
        if old and old['target'] != target and item.get('expected_target') != old['target']:
            raise ValueError('现有位置与准备修改的位置不一致；请使用返回的 current_positions，不按旧编号修改')
        return target
    if not old:
        return target
    engine = read_word_template(path) if path.suffix.lower() == '.docx' else None
    current = physical(path, target, engine)
    retained = []
    for child in old['target'].get('targets', [old['target']]):
        try:
            old_positions = physical(path, child, engine)
            if intersects(old_positions, current):
                continue
            validate_target(path, child)
            retained.append(child)
        except (OfficeKitError, ValueError, KeyError, IndexError, TypeError):
            continue
    retained.append(target)
    return {'kind': 'multi', 'targets': retained} if len(retained) > 1 else target


def retire_overlap(store, path, target, keep_field, batch):
    """Replace this exact location's prior mapping, never delete rules by IDs."""
    engine = read_word_template(path) if path.suffix.lower() == '.docx' else None
    wanted = physical(path, target, engine)
    for rule in list(store.rules_for(path)):
        if rule['field'] == keep_field:
            continue
        kept, removed = [], False
        for child in rule['target'].get('targets', [rule['target']]):
            try:
                overlaps = intersects(physical(path, child, engine), wanted)
            except (OfficeKitError, KeyError, IndexError, ValueError, TypeError):
                overlaps = False
            if overlaps:
                removed = True
            else:
                kept.append(child)
        if removed:
            store.disable_rule(path, rule['field'], rule['target'], '同一明确空位已重新关联字段，保留旧规则记录')
            if kept:
                replacement = kept[0] if len(kept) == 1 else {'kind': 'multi', 'targets': kept}
                store.add_rule(rule['template_id'], rule['field'], rule['label'], replacement,
                               confidence=rule['confidence'], decided_by=rule['decided_by'], batch_no=batch)
